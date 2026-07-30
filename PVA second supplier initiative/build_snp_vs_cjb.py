#!/usr/bin/env python3
"""PVA_CJB_vs_SNP.xlsx — SNP vs CJB, real numbers only, + sourced market reference."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

ARIAL="Arial"
def f(bold=False,size=10,color="000000",italic=False):
    return Font(name=ARIAL,bold=bold,size=size,color=color,italic=italic)
BLUE="0000FF"; GREEN="008000"
YEL=PatternFill("solid",fgColor="FFFF00")
HDR=PatternFill("solid",fgColor="1F3864")
SUB=PatternFill("solid",fgColor="D6DCE5")
GREY=PatternFill("solid",fgColor="F2F2F2")
GOOD=PatternFill("solid",fgColor="C6EFCE")
BAD=PatternFill("solid",fgColor="FFC7CE")
AMB=PatternFill("solid",fgColor="FFEB9C")
thin=Side(style="thin",color="BFBFBF")
BORD=Border(left=thin,right=thin,top=thin,bottom=thin)
CUR='$#,##0.00'; CUR0='$#,##0'; PCT='0.0%'; MULT='0.0"x"'

wb=openpyxl.Workbook()

def band(ws,row,text,span,fill=HDR,color="FFFFFF"):
    c=ws.cell(row=row,column=2,value=text); c.font=f(bold=True,size=11,color=color)
    for col in range(2,2+span): ws.cell(row=row,column=col).fill=fill
    return c

def inp(ws,row,label,value,note,fmt=None,key=True,src=""):
    ws.cell(row=row,column=2,value=label).font=f()
    c=ws.cell(row=row,column=3,value=value); c.font=f(bold=True,color=BLUE); c.alignment=Alignment(horizontal="center")
    if key: c.fill=YEL
    c.border=BORD
    if fmt: c.number_format=fmt
    ws.cell(row=row,column=4,value=note).font=f(size=9,color="595959")
    if src: ws.cell(row=row,column=5,value=src).font=f(size=9,italic=True,color="2E7D32")
    return c

def calc(ws,row,label,formula,note,fmt=None,fill=GREY,bold=True):
    ws.cell(row=row,column=2,value=label).font=f(bold=bold)
    c=ws.cell(row=row,column=3,value=formula); c.font=f(bold=bold,color="000000"); c.alignment=Alignment(horizontal="center")
    c.fill=fill; c.border=BORD
    if fmt: c.number_format=fmt
    ws.cell(row=row,column=4,value=note).font=f(size=9,color="595959")
    return c

# =====================================================================
# TAB 1 — SNP vs CJB
# =====================================================================
ws=wb.active; ws.title="SNP vs CJB"
ws.sheet_view.showGridLines=False
for k,v in {'A':2,'B':44,'C':14,'D':50,'E':34}.items(): ws.column_dimensions[k].width=v

ws['B2']="PVA — SNP vs CJB  (real numbers only)"
ws['B2'].font=f(bold=True,size=15,color="1F3864")
ws['B3']=("Only figures we actually hold: SNP's real quote and CJB's real quote. No market column, no assumed unit basis. "
          "Yellow = given/real or 'enter actual'. Grey = calculated.")
ws['B3'].font=f(italic=True,size=9,color="595959")

band(ws,5,"REAL INPUTS  (given quotes / confirmed facts)",4)
inp(ws,6,"Drum net weight (lb)",450,"Confirmed. SNP quote '1/450 lb. drum'.",None,True,"SNP Est 012726-1")
inp(ws,7,"Total demand (lb finished / week)",1300,"Confirmed by Oscar 2026-06-26.",None,True,"Oscar 06-26")
inp(ws,8,"Weeks per year",52,"")
inp(ws,9,"SNP price ($/lb, delivered all-in)",0.75,"Real quote. Includes materials + mixing + freight.",CUR,True,"SNP Est 012726-1")
inp(ws,10,"CJB toll processing $/lb - LOW",8.00,"Real quote, EXCLUDING raw materials. Basis unconfirmed (see note).",CUR,True,"CJB 'RE: mix test'")
inp(ws,11,"CJB toll processing $/lb - HIGH",8.50,"Real quote, EXCLUDING raw materials.",CUR,True,"CJB 'RE: mix test'")
inp(ws,12,"CJB qualification fee ($)",4900,"Real. Applies win or lose.",CUR0,True,"CJB")
inp(ws,13,"CJB credit per batch ($)",700,"Real. Only if trial succeeds + 1-yr deal.",CUR0,True,"CJB")
inp(ws,14,"CJB # batches credited",7,"$700 x 7 = full qual credited back.",None,True,"CJB")
inp(ws,15,"Raw materials $/drum (we supply)",0,"ENTER ACTUAL. Not yet quantified with real numbers. Kept 0 until you have it.",CUR0,True,"<-- enter real")
inp(ws,16,"Finished freight $/drum (GA -> us)",0,"ENTER ACTUAL / confirm with CJB. Kept 0 until quoted.",CUR0,True,"<-- confirm")

ws['C10'].comment=Comment("OPEN QUESTION to CJB: $8-8.50 per pound of WHAT? Modeled here literally: per lb of the finished 450-lb drum "
                          "(the same basis SNP bills on). If CJB means something else, they must state it.","model")

band(ws,18,"PER-DRUM COMPARISON  (a drum = a batch = 450 lb)",4)
calc(ws,19,"SNP - $/drum, delivered all-in","=C6*C9","= 450 lb x $0.75.",CUR,GREY)
calc(ws,20,"CJB toll $/drum @ $8.00/lb","=C6*C10","= 450 lb x $8.00. Toll only.",CUR0,GREY)
calc(ws,21,"CJB toll $/drum @ $8.50/lb","=C6*C11","= 450 lb x $8.50. Toll only.",CUR0,GREY)
calc(ws,22,"CJB toll $/drum (midpoint)","=(C20+C21)/2","CJB service fee only.",CUR0,AMB)
calc(ws,23,"  + raw materials (we supply)","=C15","On top of toll.",CUR0,GREY,bold=False)
calc(ws,24,"  + finished freight","=C16","On top of toll.",CUR0,GREY,bold=False)
calc(ws,25,"CJB LANDED $/drum (midpoint)","=C22+C23+C24","Full cost to us per drum.",CUR0,AMB)
calc(ws,26,"CJB net $/drum, first 7 batches (- $700)","=C22-C13","Credit applied to the toll midpoint.",CUR0,GREY)
calc(ws,27,"CJB toll vs SNP all-in (x)","=C22/C19","How many times SNP's entire delivered price.",MULT,BAD)
calc(ws,28,"CJB net (first 7) vs SNP all-in (x)","=C26/C19","Even after the $700 credit.",MULT,BAD)

band(ws,30,"ANNUAL VIEW  (at the previously-decided 25% second-source split)",4)
inp(ws,31,"Second-source split %",0.25,"Decided range 20-25% (from main model). Real decision, adjust to test.",PCT,True,"model decision")
calc(ws,32,"Total drums / year","=C7*C8/C6","",'#,##0.0',GREY,bold=False)
calc(ws,33,"2nd-source drums / year","=C32*C31","Drums that would go to CJB.",'#,##0.0',GREY,bold=False)
calc(ws,34,"If those drums stayed at SNP ($/yr)","=C33*C19","",CUR0,GREY)
calc(ws,35,"Same drums via CJB toll only ($/yr)","=C33*C22","Excl. materials, freight, qual.",CUR0,BAD)
calc(ws,36,"All-SNP baseline, whole account ($/yr)","=C32*C19","100% SNP, for reference.",CUR0,GREY)

ws['B38']="Bottom line"; ws['B38'].font=f(bold=True,color="1F3864")
bl=["Taken literally (per lb of the finished 450-lb drum, SNP's own basis), CJB's toll ALONE is ~10-11x SNP's entire delivered price.",
    "The $700/batch credit only reduces the first 7 barrels to ~$2,900-3,125 - still ~9x SNP. It barely moves the needle.",
    "Materials + freight are extra and still unquantified (kept at $0 above until real numbers land).",
    "The one thing that could change this: CJB confirming their '$8/lb' means something other than per finished pound. Force that answer."]
for i,t in enumerate(bl): ws.cell(row=39+i,column=2,value="- "+t).font=f(size=9,color="404040")

# =====================================================================
# TAB 2 — MARKET REFERENCE (sourced)
# =====================================================================
mk=wb.create_sheet("Market Reference (sourced)")
mk.sheet_view.showGridLines=False
for k,v in {'A':2,'B':46,'C':16,'D':52,'E':30}.items(): mk.column_dimensions[k].width=v

mk['B2']="Market Reference — what mixing PVA at this scale really costs"
mk['B2'].font=f(bold=True,size=14,color="1F3864")
mk['B3']=("Every number here is either a cited source or a transparent calc FROM a cited source. "
          "Green cell = sourced. Grey = derived (formula shown). Nothing here is a guess.")
mk['B3'].font=f(italic=True,size=9,color="595959")

band(mk,5,"1) RAW MATERIAL FLOOR",4)
g=mk.cell(row=6,column=2,value="Raw PVA resin - commodity price, N. America ($/lb)"); g.font=f()
c=mk.cell(row=6,column=3,value=1.35); c.fill=GOOD; c.font=f(bold=True,color=GREEN); c.number_format=CUR; c.alignment=Alignment(horizontal="center"); c.border=BORD
mk.cell(row=6,column=4,value="~$3.10/kg (Dec-2025) = ~$1.41/lb; ~$2,778/MT (Mar-2026) = ~$1.26/lb. Using $1.35 mid.").font=f(size=9,color="595959")
mk.cell(row=6,column=5,value="ChemAnalyst; IMARC").font=f(size=9,italic=True,color="2E7D32")
inp(mk,7,"Solids content of finished solution (%)",0.11,"Product spec: 10-12% solids in water. Midpoint.",PCT,True,"process spec")
calc(mk,8,"Material cost per lb of finished solution","=C6*C7","= resin $/lb x solids %. The material floor.",CUR,GREY)
calc(mk,9,"Material cost per 450-lb drum","=C8*450","Just the PVA in a drum (excl. minor NaCl/pigment).",CUR0,GREY)

band(mk,11,"2) REAL ALL-IN MARKET PRICE WE ALREADY HOLD",4)
g=mk.cell(row=12,column=2,value="SNP delivered price ($/lb finished)"); g.font=f()
c=mk.cell(row=12,column=3,value=0.75); c.fill=GOOD; c.font=f(bold=True,color=GREEN); c.number_format=CUR; c.alignment=Alignment(horizontal="center"); c.border=BORD
mk.cell(row=12,column=4,value="A real competitive quote. This IS a market data point for making a drum of this solution.").font=f(size=9,color="595959")
mk.cell(row=12,column=5,value="SNP Est 012726-1").font=f(size=9,italic=True,color="2E7D32")
calc(mk,13,"Implied conversion + pack + freight + margin","=C12-C8","What SNP's price covers beyond raw material.",CUR,GREY)
calc(mk,14,"That as a share of SNP's price","=C13/C12","Most of the price is the operation, not the PVA.",PCT,GREY)

band(mk,16,"3) TOLL CONVERSION RATES",4)
notes3=[
 "No public $/lb toll rate exists. Reputable sources agree tolling is quoted case-by-case (per-lb, per-batch, or flat fee).",
 "Small batches cost MORE per unit: fixed setup, QC, and cleaning are spread over few units (ITC). This is the real source of a 2nd-source premium.",
 "So the 'premium' cannot be looked up - it must come from actual competing quotes. There is no reputable published number to cite.",
]
for i,t in enumerate(notes3): mk.cell(row=17+i,column=2,value="- "+t).font=f(size=9,color="404040")

band(mk,21,"4) HOW CJB'S QUOTE COMPARES TO THESE ANCHORS",4)
calc(mk,22,"CJB toll (midpoint) $/lb finished","=('SNP vs CJB'!C10+'SNP vs CJB'!C11)/2","From the quote, literal basis.",CUR,AMB)
calc(mk,23,"CJB toll vs material floor (x)","=C22/C8","",MULT,BAD)
calc(mk,24,"CJB toll vs SNP all-in price (x)","=C22/C12","",MULT,BAD)

mk['B26']="Reading"; mk['B26'].font=f(bold=True,color="1F3864")
rd=["The real cost of this operation is bracketed by data we can cite: material ~$0.13-0.17/lb, full all-in market price ~$0.75/lb (SNP).",
    "CJB's toll-only $8-8.50/lb sits ~50x the material floor and ~10-11x the entire all-in market price. That is not a premium; it is an outlier on any reading.",
    "The only legitimate way to get a SECOND market price is competing quotes from APV / Piedmont / Columbus. That is the real fix for the blind spot."]
for i,t in enumerate(rd): mk.cell(row=27+i,column=2,value="- "+t).font=f(size=9,color="404040")

band(mk,31,"SOURCES",4)
srcs=[
 "PVA resin commodity price: chemanalyst.com/Pricing-data/polyvinyl-alcohol-1108 ; imarcgroup.com/polyvinyl-alcohol-pricing-report (N. America, Dec-2025 / Mar-2026).",
 "Toll blending is quote-driven, small batches cost more per unit: itctollblenders.com/2025/11/what-are-the-main-toll-blending-cost-factors/ ; royalchemical.com/blog/all-about-chemical-blending-packaging-transportation.",
 "SNP price: SNP Estimate 012726-1, item 5TT-11A ($0.75/lb delivered).",
 "CJB terms: email thread 'RE: CJB PVA mix test' (qual $4,900; $700 x 7 credit; toll $8.00-8.50/lb excl. materials).",
 "Solids spec (10-12%) & 450-lb drum: project process spec / CJB call agenda; demand confirmed by Oscar 2026-06-26.",
]
for i,s in enumerate(srcs): mk.cell(row=32+i,column=2,value="- "+s).font=f(size=8,color="404040")

try: wb.calculation.fullCalcOnLoad=True
except Exception: pass
wb.properties.calcMode="auto"
out="/home/user/Claude-Works/PVA second supplier initiative/PVA_CJB_vs_SNP.xlsx"
wb.save(out)
print("saved",out)
