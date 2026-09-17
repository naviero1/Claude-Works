#!/usr/bin/env python3
"""It-07 exemplars: slides 34 and 39 of the 103-slide deck, visual-pass style.
Works on a scratch copy; nothing in the repo is touched."""
import shutil
import sys

sys.path.insert(0, '/home/user/Claude-Works/ai-training-prompt-engineering/src')
import rebuild_slides_33_44 as R
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_ANCHOR

SP = '/tmp/claude-0/-home-user-Claude-Works/a43332a6-06a6-532a-af4a-fdd550eb8be6/scratchpad'
SRC = '/home/user/Claude-Works/ai-training-prompt-engineering/deliverables/From_Prompts_to_Agents_Facilitated_60_Minute.pptx'
OUT = SP + '/exemplar_deck.pptx'
OFFWHITE = 'F7F8F6'
GRAY = 'A7B0B5'

shutil.copy(SRC, OUT)
prs = Presentation(OUT)
slides = list(prs.slides)
assert len(slides) == 103


def visual_panel(slide, img, px, py, pw, caption):
    from PIL import Image
    iw, ih = Image.open(img).size
    disp_h = pw * ih / iw
    # white card behind
    R.add_box(slide, 'visual-card', px - 0.12, py - 0.12, pw + 0.24, disp_h + 0.72,
              [], fill='FFFFFF', line=GRAY, line_w=0.75)
    slide.shapes.add_picture(img, Inches(px), Inches(py), width=Inches(pw))
    R.add_box(slide, 'visual-caption', px, py + disp_h + 0.06, pw, 0.32,
              [(caption, 9.5, False, False, R.MUTED)], margins=(0, 0.02, 0, 0.02))
    return disp_h


# ---------------- slide 34: upload + inspect, workbook crop -----------------
sl = slides[33]
R.clear_slide(sl)
R.build_chrome(sl, 34, 'DO · ANCHOR', '6:00',
               'Upload the Workbook and Inspect the Data',
               'Attempt first — confirm the scope after prompt 1, then analyze.',
               'From Prompts to Agents',
               R.AMBER_BAR, R.AMBER_CHIP, R.AMBER_BAND, title_size=33.0)
R.set_bg(sl, OFFWHITE)
R.add_box(sl, 'lead', 0.667, 1.98, 6.9, 0.7,
          [('Upload Supplier_Data_Exercise.xlsx —', 21.0, True, False, R.INK),
           ('attach nothing else.', 21.0, True, False, R.INK)])
R.add_box(sl, 'card-1', 0.667, 2.72, 6.9, 1.45,
          [('Prompt 1 · Inspect', 16.0, True, False, R.AMBER_CHIP),
           ('“Inspect the workbook. List the fields, row count, date range, '
            'suppliers and sites, any total row, and missing values. Then wait '
            'for my scope confirmation.”', 18.0, False, False, R.INK)])
R.add_box(sl, 'card-2', 0.667, 4.28, 6.9, 1.9,
          [('Prompt 2 · Analyze', 16.0, True, False, R.AMBER_CHIP),
           ('“Using the confirmed detail rows, compare supplier return rate: '
            'total units returned ÷ total units shipped. Keep defect rate '
            'separate. Show totals, the missing-data rule, and the '
            'reconciliation. Do not infer causes.”', 18.0, False, False, R.INK)])
visual_panel(sl, SP + '/visuals/wb_crop.png', 7.95, 2.45, 4.55,
             'Supplier_Data_Exercise.xlsx — the single intake file (crop).')
R.set_notes(sl, 'MODE: DO\nTIME: 6:00\n(exemplar preview — full notes carry over in the real pass)')

# ---------------- slide 39: presentation storyboard -------------------------
sl = slides[38]
R.clear_slide(sl)
R.build_chrome(sl, 39, 'DO', '2:30',
               ('Create a Five-Slide', 'Management Presentation'),
               'One message per slide · findings separated from recommendations '
               '· no invented causes or commitments.',
               'From Prompts to Agents',
               R.AMBER_BAR, R.AMBER_CHIP, R.AMBER_BAND, title_size=33.0)
R.set_bg(sl, OFFWHITE)
R.add_box(sl, 'lead', 0.667, 2.15, 7.2, 0.85,
          [('Audience: Operations leadership', 21.0, True, False, R.INK),
           ('Decision: Where should management focus follow-up?', 21.0, True, False, R.INK)])
R.add_box(sl, 'main-prompt', 0.667, 3.3, 7.2, 2.75,
          [('1  Business question and scope', 21.0, False, False, R.INK),
           ('2  Supplier comparison', 21.0, False, False, R.INK),
           ('3  Important trend', 21.0, False, False, R.INK),
           ('4  Recommended follow-up and evidence', 21.0, False, False, R.INK),
           ('5  Limitations and next steps', 21.0, False, False, R.INK)])
visual_panel(sl, SP + '/visuals/storyboard.png', 8.35, 2.15, 3.95,
             'Storyboard: Supplier_Quality_Mock_Presentation.pptx (prepared output).')
R.set_notes(sl, 'MODE: DO\nTIME: 2:30\n(exemplar preview — full notes carry over in the real pass)')

prs.save(OUT)
print('exemplar deck saved')
