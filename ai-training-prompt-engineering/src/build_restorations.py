#!/usr/bin/env python3
"""Restoration wave 1 — owner-directed (in session, 2026-09-17).

Restores two of the earlier slides Oscar named "valuable as they are",
appended to the reference layer as slides 104-105 so no existing slide,
number, or cross-reference moves:

  104  Chat, workflow, or agent - the path decides   (+ courier work-order image)
  105  Four context failures - name it, then cure it (+ instruction-layer stack)

Both are fact-stable (no dated claims). Two coherence adaptations to the
current deck are made and documented: the old "Part 3" cross-references
become structure-neutral phrasing. Built in the appendix idiom measured
from slide 73 (Calibri 11 bold teal kicker, Cambria 30 bold ink title,
F2F5F6 cards, REFERENCE-STUDY chrome).

Input: pilot v02.  Output: From_Prompts_to_Agents_Visual_Pilot_v03.pptx
"""
import os

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.dml.color import RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "notes", "visual-polish", "pilot", "From_Prompts_to_Agents_Visual_Pilot_v02.pptx")
OUT = os.path.join(ROOT, "notes", "visual-polish", "pilot", "From_Prompts_to_Agents_Visual_Pilot_v03.pptx")
IMG = os.path.join(ROOT, "src", "assets", "images")

TEAL = "0E7C7B"
INK = "232A31"
SLATE = "3D4B50"
CHROME = "526267"
GRAY = "657278"
CARD = "F2F5F6"
WHITE = "FFFFFF"
AMBER_DARK = "875012"
REF_BAR = "A7B0B5"


def box(slide, name, x, y, w, h, runs_paras, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, ml=0.08, mt=0.03):
    """runs_paras: list of paragraphs, each list of run dicts
    {t, pt, bold, italic, color, font(optional, default Calibri)}."""
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.name = name
    sp.shadow.inherit = False
    sp.fill.background()
    sp.line.fill.background()
    tf = sp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(ml)
    tf.margin_right = Inches(ml)
    tf.margin_top = Inches(mt)
    tf.margin_bottom = Inches(mt)
    for i, runs in enumerate(runs_paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for rd in runs:
            r = p.add_run()
            r.text = rd["t"]
            r.font.size = Pt(rd["pt"])
            r.font.bold = bool(rd.get("bold"))
            r.font.italic = bool(rd.get("italic"))
            r.font.name = rd.get("font", "Calibri")
            r.font.color.rgb = RGBColor.from_string(rd.get("color", INK))
    return sp


def card(slide, name, x, y, w, h, fill=CARD, line=None, weight=1.0):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.name = name
    sp.shadow.inherit = False
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(weight)
    return sp


def chrome(slide, kicker, title, page, detail):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                                 Inches(13.333), Inches(0.042))
    bar.name = "FACILITATION_MODE_BAR"
    bar.shadow.inherit = False
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string(REF_BAR)
    bar.line.fill.background()
    box(slide, "FACILITATION_MODE_LABEL", 10.417, 0.177, 2.552, 0.312,
        [[{"t": "REFERENCE  STUDY", "pt": 10.5, "bold": True, "color": CHROME}]],
        align=PP_ALIGN.RIGHT, ml=0.02, mt=0.01)
    box(slide, "kicker", 0.55, 0.30, 12.2, 0.30,
        [[{"t": kicker, "pt": 11, "bold": True, "color": TEAL}]], ml=0.0, mt=0.0)
    box(slide, "title", 0.55, 0.58, 12.2, 0.75,
        [[{"t": title, "pt": 30, "bold": True, "color": INK, "font": "Cambria"}]],
        ml=0.0, mt=0.0)
    box(slide, "detail-link", 0.552, 7.146, 11.25, 0.281,
        [[{"t": detail, "pt": 10.5, "bold": False, "color": CHROME}]], ml=0.0, mt=0.0)
    box(slide, "page-number", 12.448, 7.146, 0.521, 0.281,
        [[{"t": str(page), "pt": 10.5, "bold": False, "color": GRAY}]],
        align=PP_ALIGN.RIGHT, ml=0.0, mt=0.0)


def build_vehicle(slide):
    chrome(slide, "PICK THE VEHICLE", "Chat, workflow, or agent — the path decides",
           104, "DETAILED REFERENCE · Related live slides 46–48")
    cols = [
        ("CHAT", "You are the loop.",
         "Exploring, drafting, judging as you go. You read every answer and steer "
         "every turn — the prompting craft, live."),
        ("WORKFLOW", "The loop is frozen.",
         "Known path, few tools, same steps every time — script it (n8n, Zapier, "
         "Power Automate). Cheaper, predictable, auditable. No judgment needed mid-run."),
        ("AGENT", "The loop runs inside.",
         "Unknown path — the next step depends on what it finds. Brief it, with "
         "gates where the judgment stays yours. Your control lives in the brief, "
         "not the conversation."),
    ]
    for i, (head, tag, body) in enumerate(cols):
        x = 0.55 + i * 4.13
        card(slide, f"veh-card-{i}", x, 1.62, 3.95, 2.52)
        box(slide, f"veh-head-{i}", x + 0.14, 1.76, 3.6, 0.34,
            [[{"t": head, "pt": 13, "bold": True, "color": TEAL}]])
        box(slide, f"veh-tag-{i}", x + 0.14, 2.14, 3.6, 0.32,
            [[{"t": tag, "pt": 11, "bold": True, "italic": True, "color": INK}]])
        box(slide, f"veh-body-{i}", x + 0.14, 2.52, 3.67, 1.55,
            [[{"t": body, "pt": 11, "color": SLATE}]])

    # courier work-order image card, bottom left
    card(slide, "veh-img-card", 0.55, 4.38, 3.1, 2.55, fill=WHITE, line=REF_BAR, weight=1.0)
    im = Image.open(os.path.join(IMG, "courier_work_order.jpg"))
    aspect = im.size[0] / im.size[1]
    iw = 2.82
    ih = iw / aspect
    if ih > 1.95:
        ih = 1.95
        iw = ih * aspect
    slide.shapes.add_picture(os.path.join(IMG, "courier_work_order.jpg"),
                             Inches(0.55 + (3.1 - iw) / 2), Inches(4.50), Inches(iw), Inches(ih))
    box(slide, "veh-img-cap", 0.62, 4.50 + 1.98, 2.96, 0.42,
        [[{"t": "Don’t chat with an agent — hand it a work order.",
           "pt": 9.5, "italic": True, "color": GRAY}]],
        align=PP_ALIGN.CENTER)

    card(slide, "veh-rule", 3.85, 4.38, 8.9, 1.08)
    box(slide, "veh-rule-t", 3.99, 4.48, 8.62, 0.9,
        [[{"t": "The decision rule: ", "pt": 11, "bold": True, "color": INK},
          {"t": "known path + few tools → workflow. Unknown path, real judgment "
                "→ agent, with gates. Still thinking it through → chat.",
           "pt": 11, "color": SLATE}]])
    card(slide, "veh-cost", 3.85, 5.62, 8.9, 1.08)
    box(slide, "veh-cost-t", 3.99, 5.72, 8.62, 0.9,
        [[{"t": "The wrong vehicle costs: ", "pt": 11, "bold": True, "color": INK},
          {"t": "an agent on a known path pays judgment prices for clerk work; a "
                "workflow on an unknown path is a script that breaks on the first "
                "surprise.", "pt": 11, "color": SLATE}]])


def build_context_failures(slide):
    chrome(slide, "WHEN THE RUN GOES WRONG",
           "Four context failures — name it, then cure it",
           105, "DETAILED REFERENCE · Related live slide 54")
    rows = [
        ("POISONING",
         "One bad “fact” entered the log early — every later step politely builds on it.",
         "fresh session; re-load only what you trust — the standing brief + verified files."),
        ("DRIFT",
         "A long run wanders off the goal; the middle of a huge context gets skimmed.",
         "re-inject the mission; compact (the HANDOFF move); stop conditions in the brief."),
        ("CONFUSION",
         "Too many tools on the desk — it picks the wrong one, or dithers between them.",
         "fewer tools per phase; the <plan> names the methods allowed."),
        ("CLASH",
         "Two sources disagree and the run silently picks one of them.",
         "one source of truth per fact, named in <inputs>; an explicit override order."),
    ]
    for i, (label, desc, cure) in enumerate(rows):
        y = 1.62 + i * 1.30
        card(slide, f"cf-card-{i}", 0.55, y, 7.05, 1.18)
        box(slide, f"cf-label-{i}", 0.69, y + 0.10, 1.65, 0.98,
            [[{"t": label, "pt": 12, "bold": True, "color": INK}]])
        box(slide, f"cf-body-{i}", 2.40, y + 0.08, 5.10, 1.04,
            [[{"t": desc, "pt": 10.5, "color": SLATE}],
             [{"t": "Cure: ", "pt": 10.5, "bold": True, "color": AMBER_DARK},
              {"t": cure, "pt": 10.5, "color": SLATE}]], mt=0.02)

    # right panel: the instruction-layer stack
    card(slide, "cf-panel", 7.80, 1.62, 4.95, 5.36)
    box(slide, "cf-panel-head", 7.96, 1.76, 4.6, 0.34,
        [[{"t": "“Why did it ignore me?”", "pt": 13, "bold": True, "color": INK}]])
    layers = [
        ("SYSTEM PROMPT", "the vendor’s standing orders"),
        ("PROJECT FILE", "CLAUDE.md / AGENTS.md"),
        ("SKILL", "loads when the task matches"),
        ("YOUR MESSAGE", "today’s ask — easiest to lose"),
    ]
    for i, (lead, rest) in enumerate(layers):
        y = 2.22 + i * 0.52
        card(slide, f"cf-layer-{i}", 7.96, y, 4.62, 0.42, fill=WHITE, line=TEAL, weight=1.0)
        box(slide, f"cf-layer-t-{i}", 8.04, y + 0.04, 4.48, 0.34,
            [[{"t": lead + "  ", "pt": 9.5, "bold": True, "color": INK},
              {"t": rest, "pt": 9.5, "color": SLATE}]], ml=0.03, mt=0.01)
    box(slide, "cf-expl", 7.96, 4.42, 4.62, 1.55,
        [[{"t": "Instructions stack in layers; conflicts resolve by precedence — "
                "the layer you typed is the easiest to lose. Find which layer said "
                "what before you blame the model. One-job rules belong in a SKILL, "
                "not another standing paragraph.", "pt": 10, "color": SLATE}]])
    box(slide, "cf-note", 7.96, 6.10, 4.62, 0.72,
        [[{"t": "Tool descriptions are prompts too — the agent reads them like a "
                "junior reads an API doc.", "pt": 9.5, "italic": True, "color": GRAY}]])
    box(slide, "cf-foot", 0.55, 6.86, 7.0, 0.28,
        [[{"t": "The requirements grid diagnosed the PROMPT; this grid diagnoses the RUN.",
           "pt": 10, "italic": True, "color": GRAY}]], ml=0.0, mt=0.0)


def main():
    prs = Presentation(SRC)
    assert len(prs.slides._sldIdLst) == 103
    layout = prs.slides[72].slide_layout  # appendix DEFAULT layout

    s104 = prs.slides.add_slide(layout)
    build_vehicle(s104)
    s105 = prs.slides.add_slide(layout)
    build_context_failures(s105)

    assert len(prs.slides._sldIdLst) == 105
    prs.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
