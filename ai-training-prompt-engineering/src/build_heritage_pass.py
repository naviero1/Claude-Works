#!/usr/bin/env python3
"""Heritage pass v02 — owner-directed (in session, 2026-09-17).

Oscar's direction: the illustrations he valued from earlier iterations go
back INTO the deck. This pass builds on the verdict-state pilot v01 and
touches only live slides 8, 9, 10, 15, reattaching the surviving
owner-approved images from src/assets/images/ with all text preserved
verbatim (existing shapes are repositioned, never rebuilt). Slides 18, 73,
74, 75 already carry their full illustrated layouts and are untouched.

Output: notes/visual-polish/pilot/From_Prompts_to_Agents_Visual_Pilot_v02.pptx
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "notes", "visual-polish", "pilot", "From_Prompts_to_Agents_Visual_Pilot_v01.pptx")
OUT = os.path.join(ROOT, "notes", "visual-polish", "pilot", "From_Prompts_to_Agents_Visual_Pilot_v02.pptx")
IMG = os.path.join(ROOT, "src", "assets", "images")

BG = "F7F8F6"
PANEL = "FFFFFF"
RULE = "A7B0B5"

TOUCHED = {8, 9, 10, 15}


def set_bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)


def shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(name)


def mode_bar_el(slide):
    for sh in slide.shapes:
        if sh.name == "FACILITATION_MODE_BAR":
            return sh._element
    return None


def add_panel(slide, name, x, y, w, h):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.name = name
    sp.shadow.inherit = False
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(PANEL)
    sp.line.color.rgb = RGBColor.from_string(RULE)
    sp.line.width = Pt(1.0)
    anchor = mode_bar_el(slide)
    el = sp._element
    el.getparent().remove(el)
    if anchor is not None:
        anchor.addnext(el)
    else:
        slide.shapes._spTree.insert(2, el)
    return sp


def three_row_with_image(slide, image, img_aspect=1.79):
    """Slides 8/9/10: narrow the three label/body rows, respace them, and
    place the heritage illustration on a white card at the right."""
    set_bg(slide)
    tops = [2.115, 3.445, 4.775]
    for i in range(3):
        lab = shape(slide, f"row-{i}-label")
        bod = shape(slide, f"row-{i}-body")
        lab.left, lab.width = Inches(0.667), Inches(2.75)
        lab.top, lab.height = Inches(tops[i]), Inches(1.25)
        bod.left, bod.width = Inches(3.55), Inches(5.85)
        bod.top, bod.height = Inches(tops[i]), Inches(1.25)
    # image card, vertically centered on the row block
    iw = 2.88
    ih = iw / img_aspect
    px, pw = 9.55, 3.12
    ph = ih + 0.28
    py = 2.115 + (3.91 - ph) / 2.0
    add_panel(slide, "heritage-panel", px, py, pw, ph)
    slide.shapes.add_picture(os.path.join(IMG, image),
                             Inches(px + (pw - iw) / 2.0), Inches(py + 0.14),
                             Inches(iw), Inches(ih))


def four_row_with_emblems(slide):
    """Slide 15: one hallucination emblem beside each failure row."""
    set_bg(slide)
    emblems = ["emblem_cue_card.jpg",        # Invented fact - the confident guess
               "emblem_fake_receipt.jpg",    # Invented citation - the fake receipt
               "emblem_rubber_chicken.jpg",  # Wrong document/revision - the joke taken seriously
               "emblem_gilded_frame.jpg"]    # Lost condition - garbage in, gospel out
    rows = [2.115, 3.062, 4.01, 4.958]
    for i in range(4):
        lab = shape(slide, f"row-{i}-label")
        bod = shape(slide, f"row-{i}-body")
        lab.left, lab.width = Inches(1.48), Inches(3.15)
        lab.text_frame.margin_left = Inches(0.03)
        lab.text_frame.margin_right = Inches(0.02)
        bod.left, bod.width = Inches(4.75), Inches(7.92)
        e = 0.68
        slide.shapes.add_picture(os.path.join(IMG, emblems[i]),
                                 Inches(0.667), Inches(rows[i] + (0.885 - e) / 2.0),
                                 Inches(e), Inches(e))


def main():
    prs = Presentation(SRC)
    assert len(prs.slides._sldIdLst) == 103

    three_row_with_image(prs.slides[7], "lego_tokens.jpg")    # 8  Tokens
    three_row_with_image(prs.slides[8], "desk_cabinet.jpg")   # 9  Context
    three_row_with_image(prs.slides[9], "open_book.jpg")      # 10 Grounded answers
    four_row_with_emblems(prs.slides[14])                     # 15 Failure patterns

    prs.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
