#!/usr/bin/env python3
"""Curriculum revision v07 — owner-directed (in session, 2026-09-18).

Oscar's revision list (his numbers = baseline/live numbering, verified
against titles):

  - Slides 11, 14, 17 are "useless ... the crap you both erased and added":
    the summarized twins. DELETE (v06 positions 11, 14, 20). The old
    six-assistant tools slide he wants "here" is already restored (v06 21)
    and now sits where the summary was.
  - Slide 19's topic (data workspaces/residency) is out of scope — "just
    teaching prompt engineering". v06 25 RELOCATES to self-study.
  - Before slide 27 (menu of requirements per artifact, v06 33): a NEW
    slide defining artifacts and how they relate to generative vs agentic AI.
  - Slides 29-32 (v06 35-38: improvement-loop example, techniques, rubric,
    checklist) are not needed: RELOCATE to self-study; in their place bring
    the old deck's prompt slides — playbook / keepsake prompt / thirteen
    templates / prompt-is-a-document (old 30-33) — then straight into the
    exercises, mirroring the old deck's prompts->examples structure.
  - Plus the two supplied model-selection slides (insertion brief +
    instruction 14), placed where the deleted compare-summary was:
    landscape -> matrix -> tiers -> pick-by-task -> DeepSeek.

Input: pilot v06 (109).  Output: From_Prompts_to_Agents_Visual_Pilot_v07.pptx
109 - 3 deleted + 2 model + 1 artifacts + 4 ported = 113 slides.
"""
import copy
import os
import re

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "notes", "visual-polish", "pilot", "From_Prompts_to_Agents_Visual_Pilot_v06.pptx")
OLD = os.path.join(ROOT, "references", "From_Prompts_to_Agents_Training.pptx")
MS = os.path.join(ROOT, "notes", "model-selection-addition", "Task_and_Model_Selection_2_Slides.pptx")
OUT = os.path.join(ROOT, "notes", "visual-polish", "pilot", "From_Prompts_to_Agents_Visual_Pilot_v07.pptx")

TEAL = "0E7C7B"
INK = "232A31"
CHROME = "526267"
GRAY = "657278"
CARD = "F2F5F6"
BG = "F7F8F6"

# order over v07 sldIdLst positions AFTER the 7 new slides are appended:
# 110=MS matrix, 111=MS tiers, 112=artifacts definition,
# 113=old30 playbook, 114=old31 keepsake, 115=old32 templates, 116=old33 document
NEW_ORDER = (
    list(range(1, 11))            # 1..10
    + [12, 13]
    + [15, 16, 17, 18, 19]
    + [110, 111]
    + [21, 22, 23, 24]
    + list(range(26, 33))         # 26..32
    + [112, 33, 34]
    + [113, 114, 115, 116]
    + list(range(39, 110))        # 39..109 (live rest + appendix)
    + [25, 35, 36, 37, 38]        # relocated to self-study, at the end
)
DELETED = {11, 14, 20}
NEW_SLIDES = set(range(110, 117))
RELOCATED = [25, 35, 36, 37, 38]


def add_box(slide, name, x, y, w, h, text, pt, bold, color, align=PP_ALIGN.LEFT,
            font="Calibri", italic=False):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.name = name
    sp.shadow.inherit = False
    sp.fill.background()
    sp.line.fill.background()
    tf = sp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Inches(0.0))
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(pt)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    r.font.color.rgb = RGBColor.from_string(color)
    return sp


def add_card(slide, name, x, y, w, h):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.name = name
    sp.shadow.inherit = False
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(CARD)
    sp.line.fill.background()
    sp.text_frame.paragraphs[0].text = ""
    return sp


def set_bg(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                                Inches(13.333), Inches(7.5))
    bg.name = "BG"
    bg.shadow.inherit = False
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor.from_string(BG)
    bg.line.fill.background()
    bg._element.addprevious(bg._element)  # keep first in spTree
    slide.shapes._spTree.remove(bg._element)
    slide.shapes._spTree.insert(2, bg._element)
    return bg


def add_live_chrome(slide):
    add_box(slide, "detail-link", 0.552, 7.146, 11.25, 0.281,
            "From Prompts to Agents", 10.5, False, CHROME)
    add_box(slide, "page-number", 12.448, 7.146, 0.521, 0.281, "0", 10.5, False, GRAY,
            align=PP_ALIGN.RIGHT)


def port_old_slide(dest, old_slide, layout):
    """Port an old-deck slide verbatim with LIVE chrome (no reference-study
    label): owner keeps these 'as they are' in the teaching flow."""
    new = dest.slides.add_slide(layout)
    spTree = new.shapes._spTree
    id_map = {}
    # carry the slide-level background (the old dividers use a dark canvas)
    dark = False
    old_bg = old_slide._element.find(qn("p:cSld")).find(qn("p:bg"))
    if old_bg is not None:
        new_cSld = new._element.find(qn("p:cSld"))
        new_cSld.insert(0, copy.deepcopy(old_bg))
        dark = any(int(c.get("val", "FFFFFF")[:2], 16) < 0x60
                   for c in old_bg.iter(qn("a:srgbClr")))
    for sh in old_slide.shapes:
        spTree.append(copy.deepcopy(sh._element))
    for blip in spTree.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid:
            continue
        if rid in id_map:
            blip.set(qn("r:embed"), id_map[rid])
            continue
        image_part = old_slide.part.rels[rid].target_part
        new_rid = new.part.relate_to(image_part, RT.IMAGE)
        id_map[rid] = new_rid
        blip.set(qn("r:embed"), new_rid)
    from pptx.util import Emu
    for sh in list(new.shapes):
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        old_footer = (sh.name in ("Text 0", "Text 1")
                      and (t.strip().isdigit() or t.startswith("PART ")))
        # old kicker-footers in the footer band, whatever the shape name
        band_footer = (sh.top is not None and Emu(sh.top).inches >= 6.9
                       and ("· PART " in t or "PART " in t or t.strip().isdigit()))
        if old_footer or band_footer:
            sh._element.getparent().remove(sh._element)
    if dark:
        add_box(new, "detail-link", 0.552, 7.146, 11.25, 0.281,
                "From Prompts to Agents", 10.5, False, "A7B0B5")
        add_box(new, "page-number", 12.448, 7.146, 0.521, 0.281, "0", 10.5,
                False, "A7B0B5", align=PP_ALIGN.RIGHT)
    else:
        add_live_chrome(new)
    for sh in new.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size and abs(r.font.size.pt - 30.0) < 0.1:
                        r.font.name = "Cambria"
    return new


def port_ms_slide(dest, ms_slide, layout, scale, facilitation):
    """Port a model-selection slide, scaling geometry, tables and type to the
    destination canvas (source 16.667x9.375in -> 13.333x7.5in, x0.8)."""
    new = dest.slides.add_slide(layout)
    spTree = new.shapes._spTree
    for sh in ms_slide.shapes:
        spTree.append(copy.deepcopy(sh._element))

    def scaled(v):
        return str(int(round(int(v) * scale)))

    for el in spTree.iter():
        tag = el.tag
        if tag in (qn("a:off"), qn("a:ext"), qn("a:chOff"), qn("a:chExt")):
            for a in ("x", "y", "cx", "cy"):
                if el.get(a) is not None:
                    el.set(a, scaled(el.get(a)))
        elif tag == qn("a:gridCol"):
            el.set("w", scaled(el.get("w")))
        elif tag == qn("a:tr"):
            if el.get("h") is not None:
                el.set("h", scaled(el.get("h")))
        elif tag in (qn("a:rPr"), qn("a:defRPr"), qn("a:endParaRPr")):
            sz = el.get("sz")
            if sz is not None:
                el.set("sz", str(max(900, int(round(int(sz) * scale)))))

    # the source's bottom caveat line lands in the footer band after scaling;
    # keep the caveat visible by making it the footer line itself
    caveat = ""
    from pptx.util import Emu
    for sh in list(new.shapes):
        if sh.has_text_frame and Emu(sh.top).inches >= 6.95:
            caveat = sh.text_frame.text.strip()
            sh._element.getparent().remove(sh._element)
    add_box(new, "detail-link", 0.552, 7.146, 11.25, 0.281,
            caveat or "From Prompts to Agents", 10.5, False, CHROME)
    add_box(new, "page-number", 12.448, 7.146, 0.521, 0.281, "0", 10.5, False, GRAY,
            align=PP_ALIGN.RIGHT)

    # keep the source notes verbatim, then append the facilitation field
    src_notes = ms_slide.notes_slide.notes_text_frame.text if ms_slide.has_notes_slide else ""
    ntf = new.notes_slide.notes_text_frame
    ntf.text = src_notes + "\n\n" + facilitation
    return new


def build_artifacts_slide(dest, layout):
    """New definitional slide before the requirements menu (owner request:
    'we first need to define artifacts and how they are related to the type
    of ai, agentic vs generative')."""
    s = dest.slides.add_slide(layout)
    set_bg(s)
    add_box(s, "kicker", 0.552, 0.42, 9.0, 0.32,
            "PART 3 · THE CRAFT — NAME THE DELIVERABLE", 11, True, TEAL)
    add_box(s, "title", 0.552, 0.80, 12.2, 0.62,
            "Artifacts — what the AI hands back", 30, True, INK, font="Cambria")
    add_box(s, "subtitle", 0.552, 1.52, 12.2, 0.40,
            "An artifact is the finished work product — a document, table, deck, or working file — not the conversation that produced it.",
            14, False, GRAY)

    # generative card
    add_card(s, "card-generative", 0.552, 2.25, 6.0, 3.30)
    add_box(s, "gen-head", 0.80, 2.50, 5.5, 0.35, "GENERATIVE — you shape it", 15, True, TEAL)
    add_box(s, "gen-body", 0.80, 3.00, 5.5, 1.55,
            "Chat produces drafts inside the conversation: text, tables, code you refine turn by turn. "
            "The artifact emerges from the exchange — you are the loop, and you carry it out of the chat.",
            13, False, INK)
    add_box(s, "gen-ex", 0.80, 4.75, 5.5, 0.65,
            "Draft email · summary · rewritten section · one-off table", 12, False, GRAY, italic=True)

    # agentic card
    add_card(s, "card-agentic", 6.78, 2.25, 6.0, 3.30)
    add_box(s, "agt-head", 7.03, 2.50, 5.5, 0.35, "AGENTIC — you specify it", 15, True, TEAL)
    add_box(s, "agt-body", 7.03, 3.00, 5.5, 1.55,
            "An agent delivers the finished file against your brief: workbook, slide deck, report, working page. "
            "You are not in the loop while it works — so the artifact is the contract, and the requirements travel in the prompt.",
            13, False, INK)
    add_box(s, "agt-ex", 7.03, 4.75, 5.5, 0.65,
            "Analyzed workbook · five-slide deck · action-items table · research spreadsheet", 12, False, GRAY, italic=True)

    add_box(s, "principle", 0.552, 5.90, 12.2, 0.75,
            "Name the artifact first. Every requirement on the next slide attaches to the artifact — "
            "the more of the work you delegate, the more the artifact's definition has to carry.",
            14, True, INK)

    add_live_chrome(s)
    ntf = s.notes_slide.notes_text_frame
    ntf.text = (
        "TIME: 1:00 (TEACH)\n"
        "Define the term before the requirements menu uses it. An artifact is the deliverable itself — "
        "the file or structured output the AI returns — as opposed to the conversational answer.\n"
        "Tie to the vehicle slide: in chat (generative) you shape the artifact turn by turn, so requirements "
        "can arrive incrementally. In a delegated run (agentic) the artifact must be fully specified up front — "
        "the prompt is the contract.\n"
        "Transition: 'So before we pick requirements, we name the artifact. Here is the menu, per artifact.'"
    )
    return s


def main():
    prs = Presentation(SRC)
    old = Presentation(OLD)
    ms = Presentation(MS)
    layout = prs.slides[72].slide_layout
    assert len(prs.slides._sldIdLst) == 109

    scale = float(prs.slide_width) / float(ms.slide_width)
    # 110, 111: model-selection matrix + tiers
    port_ms_slide(prs, ms.slides[0], layout, scale,
                  "FACILITATION: TEACH · about 1:00. Concept slide — capability, "
                  "reasoning effort and workload volume are separate decisions.")
    port_ms_slide(prs, ms.slides[1], layout, scale,
                  "FACILITATION: REFERENCE · 0:20–0:30. Orient, do not read: bands are "
                  "starting choices, not cross-vendor benchmark rankings.")
    # 112: artifacts definition
    build_artifacts_slide(prs, layout)
    # 113..116: old prompt slides, verbatim with live chrome
    for pos in (30, 31, 32, 33):
        port_old_slide(prs, old.slides[pos - 1], layout)

    sldIdLst = prs.slides._sldIdLst
    ids = sldIdLst.findall(qn("p:sldId"))
    assert len(ids) == 116
    assert len(NEW_ORDER) == 113, len(NEW_ORDER)
    assert set(NEW_ORDER) | DELETED == set(range(1, 117))

    new_pos = {p: i + 1 for i, p in enumerate(NEW_ORDER)}
    # deleted slides map to their natural successor so stale refs still land
    new_pos[11] = new_pos[12]
    new_pos[14] = new_pos[15]
    new_pos[20] = new_pos[21]

    elements = {i + 1: el for i, el in enumerate(ids)}
    for el in ids:
        sldIdLst.remove(el)
    for p in NEW_ORDER:
        sldIdLst.append(elements[p])
    for p in DELETED:
        prs.part.drop_rel(elements[p].get(qn("r:id")))

    # position of each v07 slide's v06/new origin, for skipping remap on new content
    origin = {i + 1: p for i, p in enumerate(NEW_ORDER)}

    # page numbers
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.name == "page-number" and sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text.strip().isdigit():
                            r.text = str(i)

    # relocation note on the five moved-to-self-study slides
    reloc_positions = {new_pos[p] for p in RELOCATED}
    for i, slide in enumerate(prs.slides, 1):
        if i in reloc_positions:
            ntf = slide.notes_slide.notes_text_frame
            ntf.text = ("RELOCATED to self-study per owner direction (2026-09-18): "
                        "outside the first-session teaching scope.\n\n" + ntf.text)

    # remap textual slide references (visible text + notes), skipping the new
    # slides (their text predates this numbering) and frozen historical quotes
    ref_pat = re.compile(r"\b(slides?)\s+([0-9]{1,3}(?:\s*[–—-]\s*[0-9]{1,3})?"
                         r"(?:\s*,\s*[0-9]{1,3}(?:\s*[–—-]\s*[0-9]{1,3})?)*)\b",
                         re.IGNORECASE)
    num_pat = re.compile(r"[0-9]{1,3}")

    def remap_numbers(numstr):
        def sub(m):
            n = int(m.group(0))
            return str(new_pos.get(n, n)) if 1 <= n <= 109 else m.group(0)
        return num_pat.sub(sub, numstr)

    changed = []

    def process_tf(tf, where):
        for p in tf.paragraphs:
            for r in p.runs:
                if "v1.10 CONSOLIDATION" in r.text:
                    continue
                if r.text and "slide" in r.text.lower():
                    new = ref_pat.sub(lambda m: m.group(1) + " " + remap_numbers(m.group(2)), r.text)
                    if new != r.text:
                        changed.append((where, r.text.strip()[:60], new.strip()[:60]))
                        r.text = new

    for i, slide in enumerate(prs.slides, 1):
        if origin[i] in NEW_SLIDES:
            continue
        for sh in slide.shapes:
            if sh.has_text_frame:
                process_tf(sh.text_frame, f"s{i}:{sh.name}")
        if slide.has_notes_slide:
            process_tf(slide.notes_slide.notes_text_frame, f"s{i}:notes")

    print(f"remapped {len(changed)} references")
    for w, a, b in changed:
        print(f"  {w}: '{a}' -> '{b}'")

    # a footer pointing at a slide that moved to self-study must not call it live
    first_reloc = min(new_pos[p] for p in RELOCATED)
    live_fix = re.compile(r"live slide (\d{1,3})")
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.name == "detail-link" and sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        m = live_fix.search(r.text)
                        if m and int(m.group(1)) >= first_reloc:
                            r.text = live_fix.sub(r"slide \1 (self-study)", r.text)
                            print("fixed self-study footer:", r.text[:70])

    assert len(prs.slides._sldIdLst) == 113
    prs.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
