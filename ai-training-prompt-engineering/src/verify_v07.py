#!/usr/bin/env python3
"""Read-only verification of the v07 deck: count, page numbers, footers, quotes."""
import os

from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
P = os.path.join(ROOT, "notes", "visual-polish", "pilot",
                 "From_Prompts_to_Agents_Visual_Pilot_v07.pptx")

prs = Presentation(P)
assert len(prs.slides._sldIdLst) == 113
probs = []
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        if sh.name == "page-number":
            d = [r.text for p in sh.text_frame.paragraphs for r in p.runs
                 if r.text.strip().isdigit()]
            if d and d[0] != str(i):
                probs.append(f"s{i}: page {d[0]}")
        if "Restored classic" in sh.text_frame.text:
            probs.append(f"s{i}: stale footer")
    if s.has_notes_slide:
        nt = s.notes_slide.notes_text_frame.text
        if "v1.10 CONSOLIDATION" in nt and "28, 29 and 33" not in nt:
            probs.append(f"s{i}: quote broken")
print("verify:", probs if probs else "ALL CLEAN - 113 slides")
