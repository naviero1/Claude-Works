#!/usr/bin/env python3
# Owner request (2026-09-17): the course PDF becomes the ALL-IN-ONE edition —
# the full course plus every piece of practice material, the worked examples
# and answer keys, and a full bibliography (the separate files remain).
# This pass edits src/assets/course_content.json:
#   - ch7 "Sources this training is built on" -> "Bibliography" (full entries,
#     grouped; every item verified against notes/research/ or the deck)
#   - appends "Appendix A - The practice pack" (all exercise sources, verbatim,
#     generated FROM the shipped files so they can never drift)
#   - appends "Appendix B - Answer keys" (the instructor keys, verbatim)
# The cheat-sheet page (Appendix C) is merged by build_course_pdf.py.
# One-shot; idempotency guard below.
import json
import os
import subprocess

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, '..')
ED = os.path.join(ROOT, 'references', 'exercise-data')
P = os.path.join(HERE, 'assets', 'course_content.json')
cc = json.load(open(P))
assert not any('bibliography' in ch['title'].lower() or 'practice pack' in ch['title'].lower()
               for ch in cc), 'already applied'


def p(text):
    return {'type': 'p', 'text': text}


def h3(text):
    return {'type': 'h3', 'text': text}


def bullets(items):
    return {'type': 'bullets', 'items': items}


def mono_chunks(text, caption, limit=1500):
    """Split file text on blank lines into mono blocks that each fit a page."""
    groups = [g for g in text.replace('\r', '').split('\n\n') if g.strip()]
    chunks, cur = [], ''
    for g in groups:
        if cur and len(cur) + len(g) + 2 > limit:
            chunks.append(cur)
            cur = g
        else:
            cur = (cur + '\n\n' + g) if cur else g
    if cur:
        chunks.append(cur)
    out = []
    for i, c in enumerate(chunks):
        cap = caption if i == len(chunks) - 1 else None
        b = {'type': 'mono', 'text': c}
        if cap:
            b['caption'] = cap
        out.append(b)
    return out


# ===========================================================================
# 1. Bibliography (replaces the summarized sources section in ch7)
# ===========================================================================
ch7 = cc[6]
idx = [i for i, s in enumerate(ch7['sections'])
       if s['heading'] == 'Sources this training is built on']
assert len(idx) == 1
BIB = {'heading': 'Bibliography', 'blocks': [
    p('Every evidence claim in this course traces to a study, a vendor document, a standard, or a '
      'dated snapshot. This bibliography expands the deck’s sources page into full entries, grouped '
      'by kind. Volatile items (leaderboards, product claims) carry their retrieval date and are '
      'refreshed quarterly; the dated research record behind each entry ships with the package '
      'source (notes/research/, files r1–r28 and t1–t6, each with retrieval dates).'),

    h3('Research papers and studies'),
    bullets([
        'Brown, T., et al. (2020). “Language Models are Few-Shot Learners.” NeurIPS 2020. arXiv:2005.14165. — few-shot prompting.',
        'Wei, J., et al. (2022). “Chain-of-Thought Prompting Elicits Reasoning in Large Language Models.” NeurIPS 2022. arXiv:2201.11903.',
        'Kojima, T., et al. (2022). Zero-shot chain-of-thought (“Let’s think step by step”). NeurIPS 2022. arXiv:2205.11916.',
        'Wang, X., et al. (2023). Self-consistency over sampled reasoning paths. ICLR 2023. arXiv:2203.11171.',
        'Liu, N. F., et al. (2024). “Lost in the Middle” — long-context position effects. TACL. arXiv:2307.03172.',
        'Sharma, M., et al. (2023). “Towards Understanding Sycophancy in Language Models.” Anthropic. arXiv:2310.13548.',
        'Sclar, M., et al. (2024). Prompt-format sensitivity — up to 76 accuracy points from trivial formatting changes. ICLR 2024. arXiv:2310.11324.',
        'Zheng, M., et al. (2024). “When ‘A Helpful Assistant’ Is Not Really Helpful.” EMNLP 2024 Findings. arXiv:2311.10054. — personas and accuracy.',
        'PersonaLLM (2024). NAACL 2024 Findings (ACL Anthology 2024.findings-naacl.229). — personas shape voice, not correctness.',
        'Salinas, A., and Morstatter, F. (2024). Prompt perturbations — “even the smallest of perturbations … can cause the LLM to change its answer.” arXiv:2401.03729.',
        'Wharton Prompting Science Reports 1–4 (2025; Report 4: Dec 7, 2025, SSRN 5879722). — politeness, tips/threats, and the six-model persona replication.',
        'Sprague, Z., et al. (2025). “To CoT or not to CoT?” — meta-analysis of 100+ papers; chain-of-thought helps mainly on math and symbolic tasks. ICLR 2025. arXiv:2409.12183.',
        'Yang et al. (2025; revised 2026). Requirements-aware prompt optimization: models infer unspecified requirements correctly only 41.1% of the time; requirements-aware optimization +4.8% on average. arXiv:2505.13360.',
        'Kalai, A. T., et al. (OpenAI, 2025). “Why Language Models Hallucinate.” arXiv:2509.04664.',
        'Cheng, et al. (2026). Sycophancy across 11 models (+49%). Science. — with the follow-up “social sycophancy” work from the same lab.',
        'Chain-of-Verification (CoVe). ACL Findings 2024. — self-check against named criteria, blinded from generation (FactScore 55.9→71.4).',
        'OPRO (ICLR 2024) and GEPA (ICLR 2026, arXiv:2507.19457). — metaprompting: letting the model draft and improve the prompt.',
        'OpenAI Instruction Hierarchy (2024). arXiv:2404.13208. — the chain of command behind system/developer/user layers.',
    ]),

    h3('Vendor guidance (fetched August 2026 unless dated)'),
    bullets([
        '<b>Anthropic</b> — Prompting best practices; long-context tips; reduce-hallucinations guide; “Building Effective Agents”; context-engineering guidance; Claude Code documentation.',
        '<b>OpenAI</b> — GPT-5 prompting guide; reasoning best practices; “A Practical Guide to Building Agents”; the Model Spec (model-spec.openai.com, 2025-09-12 edition, checked 2026-08-18).',
        '<b>Google</b> — Gemini for Workspace prompting guides (Persona · Task · Context · Format).',
        '<b>Microsoft</b> — Copilot prompting guidance (Goal · Context · Source · Expectations); Copilot agents documentation; Copilot-in-Outlook behavior verified September 2026 (research note r28).',
    ]),

    h3('Standards, governance, and security'),
    bullets([
        'NASA — “Appendix C: How to Write a Good Requirement” (nasa.gov/reference/appendix-c-how-to-write-a-good-requirement/).',
        'INCOSE — Guide to Writing Requirements, version 4 summary sheet (incose.org).',
        'OWASP — LLM Top 10 (2025) and “Top 10 for Agentic Applications for 2026” (genai.owasp.org, fetched 2026-09-10).',
        'UK NCSC — “Thinking carefully before adopting agentic AI” (ncsc.gov.uk) and “Prompt injection is not SQL injection (it may be worse)” (Dec 8, 2025).',
        'Singapore IMDA — responsible-deployment case study; CSA advisory AD-2026-005 (2026-05-28).',
    ]),

    h3('Landscape and benchmarks (dated snapshots — refresh quarterly)'),
    bullets([
        'Stanford HAI — AI Index 2026 (adoption and capability-gap figures, cited as of Mar/Apr 2026).',
        'Artificial Analysis — model leaderboard (artificialanalysis.ai/leaderboards/models), transcribed 2026-09-13 for the speed-vs-capability chart; three independent fetches agreed.',
        'LMArena and SWE-bench Verified — August–September 2026 snapshots (vendor launch figures named as such).',
        'Product dates and pricing — vendor launch posts and business reporting (Reuters, CNBC, TechCrunch, Bloomberg), each dated where cited.',
    ]),

    h3('Books (the recommended-reading shelf)'),
    bullets([
        'Mollick, E. — <b>Co-Intelligence</b> (2024). The working-with-AI mindset book; read first.',
        'Kneusel, R. — <b>How AI Works</b> (2024). The concepts of Part 1, gently and rigorously.',
        'Phoenix, J., and Taylor, M. — <b>Prompt Engineering for Generative AI</b> (O’Reilly, 2024). The five principles applied; closest to Parts 3–4.',
        'Berryman, J., and Ziegler, A. — <b>Prompt Engineering for LLMs</b> (O’Reilly, 2024). How prompts actually meet the model.',
        'Huyen, C. — <b>AI Engineering: Building Applications with Foundation Models</b> (O’Reilly, 2025). RAG, agents, evaluation — the engineering behind Part 5.',
        'Alammar, J., and Grootendorst, M. — <b>Hands-On Large Language Models</b> (2024). Visual internals: tokens, embeddings, transformers.',
    ]),

    {'type': 'callout', 'kind': 'key', 'title': 'The fastest teacher of all',
     'body': 'A real task. Adapt a course example, run it, check the result against its source, and '
             'save the improved prompt. The bibliography is where the evidence lives; the habit is '
             'what the course is for.'},
]}
ch7['sections'][idx[0]] = BIB
print('ch7: Bibliography section written')

# ===========================================================================
# 2. Appendix A - the practice pack (generated from the shipped files)
# ===========================================================================
secA = []

# -- A.1 the supplier dataset (full, from the participant input workbook) ----
wb = openpyxl.load_workbook(os.path.join(ED, 'Supplier_Data_Exercise.xlsx'))
ws = wb['Data']
rows = []
for r in ws.iter_rows(min_row=2, values_only=True):
    fmt = []
    for v in r:
        if v is None:
            fmt.append('')
        elif isinstance(v, float):
            fmt.append(f'{v:g}')
        else:
            fmt.append(str(v))
    rows.append(fmt)
hdr = [str(c.value) for c in ws[1]]
assert len(rows) == 145 and rows[-1][0] == 'TOTAL'
secA.append({'heading': 'A.1 · The supplier dataset (Supplier_Data_Exercise.xlsx)', 'blocks': [
    p('The single input for the data arc, printed in full: 144 detail rows plus the deliberate TOTAL '
      'row, one row per month × site × supplier, period 2025-09 to 2026-08. The two planted quirks '
      'are here to find: the TOTAL row (exclude it from every calculation, reconcile against it), and '
      'one blank Inspection_Hours cell (missing, not zero). All data fictional and seeded — stable '
      'across rebuilds and identical to the shipped file.'),
    {'type': 'table',
     'cols': ['Month', 'Site', 'Supplier', 'Shipped', 'Returned', 'Defects', 'Insp.Hrs', 'Cost USD', 'On-time %'],
     'rows': rows,
     'note': 'Columns as in the file: ' + ', '.join(hdr) + '.'},
]})

# -- A.2 the packaging-change thread ----------------------------------------
thread = open(os.path.join(ED, 'Packaging_Change_Thread.txt')).read()
secA.append({'heading': 'A.2 · The email thread (Packaging_Change_Thread.txt)', 'blocks':
    [p('The fictional ten-message conversation for the email exercise, verbatim — paste it into any '
       'approved assistant with the structured brief from chapter 4, step 6.')]
    + mono_chunks(thread, 'Verbatim copy of Packaging_Change_Thread.txt (references/exercise-data/).')})

# -- A.3 the three quotations ------------------------------------------------
qblocks = [p('The three quotation documents for steps 7–8, transcribed from the shipped PDFs '
             '(Quote_Alpha_Components.pdf, Quote_Bravo_Plastics.pdf, Quote_Cardinal_Metals.pdf). They '
             'are engineered to be non-comparable as quoted: different currencies, delivery terms, '
             'and inclusions.')]
for f in ('Quote_Alpha_Components.pdf', 'Quote_Bravo_Plastics.pdf', 'Quote_Cardinal_Metals.pdf'):
    txt = subprocess.run(['pdftotext', '-layout', os.path.join(ED, f), '-'],
                         capture_output=True, text=True).stdout.strip()
    qblocks += mono_chunks(txt, f'Transcription of {f}.')
secA.append({'heading': 'A.3 · The supplier quotations (Quote_*.pdf)', 'blocks': qblocks})

# -- A.4 the research pack ----------------------------------------------------
rp = os.path.join(ED, 'research-pack')
rblocks = [p('The four reference files for step 9, verbatim. Two of them disagree on review cadence '
             'and one carries no date — deliberately: the exercise is scored on keeping that visible.')]
for f in ('src1_quality_manual_excerpt.md', 'src2_quality_newsletter.md',
          'src3_scorecard_guideline.txt', 'src4_8d_onepager.pdf'):
    path = os.path.join(rp, f)
    if f.endswith('.pdf'):
        txt = subprocess.run(['pdftotext', '-layout', path, '-'],
                             capture_output=True, text=True).stdout.strip()
        cap = f'Transcription of {f}.'
    else:
        txt = open(path).read().strip()
        cap = f'Verbatim copy of {f}.'
    rblocks += mono_chunks(txt, cap)
secA.append({'heading': 'A.4 · The research pack (research-pack/)', 'blocks': rblocks})

# -- A.5 the plain-language pack ---------------------------------------------
pl = os.path.join(ED, 'plain-language')
pblocks = [p('The reference exercise’s four topics — each file holds a SOURCE text to hand the '
             'assistant and a sample output produced with the plain-language prompt (chapter 4). '
             'Printed verbatim; run the prompt on a SOURCE before reading its sample.')]
for f in ('inventory_replenishment.md', 'household_budgeting.md', 'water_cycle.md',
          'internet_message.md'):
    txt = open(os.path.join(pl, f)).read().strip()
    pblocks.append(h3(f))
    pblocks += mono_chunks(txt, f'Verbatim copy of plain-language/{f}.')
secA.append({'heading': 'A.5 · The plain-language pack (plain-language/)', 'blocks': pblocks})

cc.append({
    'title': 'The practice pack: every exercise source, in full',
    'kicker': 'APPENDIX A · PRACTICE MATERIAL',
    'intro': ('Everything the exercises consume, printed verbatim so this book is self-contained: the '
              'supplier dataset, the email thread, the three quotations, the research pack, and the '
              'plain-language sources with their samples. The same content ships as separate files in '
              'references/exercise-data/ — use the files when working, this appendix when reading. '
              'All names, companies, and numbers are fictional.'),
    'sections': secA,
})
print('Appendix A built:', [s['heading'][:14] for s in secA])

# ===========================================================================
# 3. Appendix B - the answer keys
# ===========================================================================
secB = [{'heading': 'B.1 · Supplier analysis — the worked key', 'blocks': [
    p('The full expected results for the data arc, computed from the dataset (they rebuild with it).'),
    {'type': 'table', 'cols': ['Item', 'Expected result'], 'rows': [
        ['Profile', '144 detail rows; one TOTAL row to exclude (and reconcile against); one blank Inspection_Hours value (2025-12, Austin, Alpha Components) — missing, not zero.'],
        ['Totals', '224,902 units shipped · 432 returns (overall 0.192%) · 2,207 defect occurrences (a separate measure, never added to returns).'],
        ['Ranking', 'Bravo Plastics 0.348% (262 / 75,184) — highest, ≈ 3.5× Alpha · Cardinal Metals 0.129% (96 / 74,658) · Alpha Components 0.099% (74 / 75,060).'],
        ['Trend (follow-up)', 'Bravo is elevated and volatile, not steadily worsening: peak ≈0.737% (2026-04), quiet months ≈0.14% (2026-06/07), ≈0.465% (2026-08). Patterns only — any causal claim fails.'],
        ['Denominator lesson', 'No overall on-time delivery rate can be computed: On_Time_Percent has no delivery counts behind it. An averaged percentage must be labeled an unweighted mean of monthly percentages.'],
        ['Dashboard check', 'Berlin × Bravo Plastics, 2026-03..2026-08 → exactly 6 records, 8,575 units, 21 returns, 0.245% — in tiles, chart, AND table.'],
        ['Wrong turns', 'Averaging monthly percentages · adding defects to returns · treating the blank as zero · including the TOTAL row · claiming a cause the data cannot show.'],
    ]},
]}]

for heading, fname, blurb in (
    ('B.2 · Email brief — the expected result and the four traps',
     'Packaging_Change_Expected_Brief.md',
     'The instructor key for the email exercise, verbatim.'),
    ('B.3 · Quotation comparison — the key',
     'Quote_Comparison_Key.md',
     'The expected extraction and normalized values — and the expected outcome: no best quote while the basis is incomplete.'),
    ('B.4 · Research workbook — the key',
     'Research_Workbook_Key.md',
     'What a complete evidence workbook shows, including the planted conflict kept visible.'),
):
    txt = open(os.path.join(ED, 'instructor-keys', fname)).read().strip()
    secB.append({'heading': heading, 'blocks':
        [p(blurb)] + mono_chunks(txt, f'Verbatim copy of instructor-keys/{fname}.')})

cc.append({
    'title': 'Answer keys: check yourself after attempting',
    'kicker': 'APPENDIX B · ANSWER KEYS',
    'intro': ('The worked keys for every exercise — the same material a facilitator holds. In a live '
              'session these stay with the instructor; for self-study, attempt each exercise first, '
              'then grade your result here. Reading a key before attempting removes exactly the '
              'practice the exercise exists to give.'),
    'sections': secB,
})
print('Appendix B built:', len(secB), 'sections')

json.dump(cc, open(P, 'w'), ensure_ascii=False, indent=1)
print(f'saved: {len(cc)} chapters')
