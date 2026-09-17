#!/usr/bin/env python3
"""Build the approved six-slide visual pilot (versions A and B).

Scope: approvals.md + rocksteady-handoff.md at commit 5a83c8a (owner-approved
in session 2026-09-17). Modifies ONLY original/current slides 23, 28, 34, 37,
53, 87 (sids 278/283/291/294/308/342) in a versioned copy of the baseline
deck (blob e83fd87a...). Version B differs from A only on slide 34 (authentic
workbook crop). Font exception: text below 20 pt may be enlarged, wording
unchanged, every change logged. No master/theme/layout changes; background is
a per-slide property.

Outputs (all under notes/visual-polish/):
  pilot/From_Prompts_to_Agents_Visual_Pilot_v01.pptx        (version A)
  pilot/From_Prompts_to_Agents_Visual_Pilot_v01_S34B.pptx   (version B)
  font-change-log.md
  crop-provenance.md
"""
import copy
import hashlib
import json
import os
import subprocess
import sys

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASELINE = os.path.join(ROOT, "deliverables", "From_Prompts_to_Agents_Facilitated_60_Minute.pptx")
BASELINE_SHA256 = "117f60b30c292527d1ac8444ab3efafdbb0ef8d0f3ff159b68149e5804bd24fa"
PILOT_DIR = os.path.join(ROOT, "notes", "visual-polish", "pilot")
OUT_A = os.path.join(PILOT_DIR, "From_Prompts_to_Agents_Visual_Pilot_v01.pptx")
OUT_B = os.path.join(PILOT_DIR, "From_Prompts_to_Agents_Visual_Pilot_v01_S34B.pptx")
XLSX_SRC = os.path.join(ROOT, "references", "exercise-data", "Supplier_Data_Exercise.xlsx")
CROP_PNG = os.path.join(PILOT_DIR, "s34b_workbook_crop.png")
FONT_LOG_MD = os.path.join(ROOT, "notes", "visual-polish", "font-change-log.md")
CROP_MD = os.path.join(ROOT, "notes", "visual-polish", "crop-provenance.md")

TARGETS = {23: 278, 28: 283, 34: 291, 37: 294, 53: 308, 87: 342}

# Palette (established deck colors only)
BG = "F7F8F6"        # approved soft off-white (BACKGROUND decision)
PANEL = "FFFFFF"     # white working panel
RULE = "A7B0B5"      # restrained neutral rule (REFERENCE gray)
TEAL = "0E7C7B"      # TEACH teal
TEAL_DARK = "0A6261"
INK = "232A31"
SLATE = "3D4B50"
GRAY = "657278"
REF_GRAY = "5D6B70"
PALE_TEAL = "EAF3F2"  # established pale mode band -> phrase underlay
AMBER = "C47A1F"      # DO amber (checkpoint marker outline)
CONN = "8A9BA3"       # thin wireframe connector gray

FONT_LOG = []  # dicts: slide, sid, shape, text, orig_pt, new_pt, reason

DEJAVU = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
DEJAVU_B = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
_SCALE = 200.0


def wrap_lines(text, pt, width_in, bold=False):
    """Estimate wrapped line count with the render-substitute font (DejaVu,
    wider than Calibri, so estimates are conservative)."""
    f = ImageFont.truetype(DEJAVU_B if bold else DEJAVU, int(round(pt / 72.0 * _SCALE)))
    maxw = width_in * _SCALE
    lines, cur = 1, ""
    for w in text.split(" "):
        t = (cur + " " + w).strip()
        if f.getlength(t) <= maxw:
            cur = t
        else:
            lines += 1
            cur = w
    return lines


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ---------------------------------------------------------------- shape utils

def set_slide_bg(slide, hexcolor):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(hexcolor)


def _mode_bar_el(slide):
    for sh in slide.shapes:
        if sh.name == "FACILITATION_MODE_BAR":
            return sh._element
    return None


def send_to_back_layer(slide, shape):
    """Place shape just above the mode bar (i.e. beneath all text content)."""
    anchor = _mode_bar_el(slide)
    el = shape._element
    el.getparent().remove(el)
    if anchor is not None:
        anchor.addnext(el)
    else:
        slide.shapes._spTree.insert(2, el)


def add_panel(slide, name, x, y, w, h, fill=PANEL, line=None, weight=1.0,
              shape=MSO_SHAPE.RECTANGLE, back=True):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.name = name
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(weight)
    sp.text_frame.paragraphs[0].text = ""
    if back:
        send_to_back_layer(slide, sp)
    return sp


def add_highlight(run, hexcolor):
    rPr = run._r.get_or_add_rPr()
    hl = rPr.makeelement(qn("a:highlight"), {})
    clr = rPr.makeelement(qn("a:srgbClr"), {"val": hexcolor})
    hl.append(clr)
    latin = rPr.find(qn("a:latin"))
    if latin is not None:
        latin.addprevious(hl)
    else:
        rPr.append(hl)


def add_box(slide, name, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """paras: list of paragraphs; each a list of run dicts
    {t, pt, bold, italic, color, highlight(optional)}."""
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.name = name
    sp.shadow.inherit = False
    sp.fill.background()
    sp.line.fill.background()
    tf = sp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.1)
    tf.margin_right = Inches(0.1)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for rd in runs:
            r = p.add_run()
            r.text = rd["t"]
            r.font.size = Pt(rd["pt"])
            r.font.bold = bool(rd.get("bold"))
            if rd.get("italic") is not None:
                r.font.italic = bool(rd.get("italic"))
            r.font.name = "Calibri"
            r.font.color.rgb = RGBColor.from_string(rd.get("color", INK))
            if rd.get("highlight"):
                add_highlight(r, rd["highlight"])
    return sp


def add_arrow(slide, name, x1, y1, x2, y2, color=TEAL, weight=2.0, head=True):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.name = name
    cn.shadow.inherit = False
    cn.line.color.rgb = RGBColor.from_string(color)
    cn.line.width = Pt(weight)
    if head:
        ln = cn.line._get_or_add_ln()
        tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
    return cn


def delete_shape(shape):
    el = shape._element
    el.getparent().remove(el)


def get_shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(name)


def para_texts(shape):
    return [p.text for p in shape.text_frame.paragraphs]


def log_font(slide_no, shape, text, orig, new, reason):
    FONT_LOG.append({"slide": slide_no, "sid": TARGETS[slide_no], "shape": shape,
                     "text": text if len(text) <= 60 else text[:57] + "...",
                     "orig_pt": orig, "new_pt": new, "reason": reason})


def segmented(text, phrases, pt, color=INK, bold=False, italic=None):
    """Split text into runs; runs matching any phrase get the pale underlay."""
    marks = []
    for ph in phrases:
        idx = text.find(ph)
        if idx >= 0:
            marks.append((idx, idx + len(ph)))
    marks.sort()
    runs, pos = [], 0
    for a, b in marks:
        if a > pos:
            runs.append({"t": text[pos:a], "pt": pt, "bold": bold, "italic": italic, "color": color})
        runs.append({"t": text[a:b], "pt": pt, "bold": bold, "italic": italic, "color": color,
                     "highlight": PALE_TEAL})
        pos = b
    if pos < len(text):
        runs.append({"t": text[pos:], "pt": pt, "bold": bold, "italic": italic, "color": color})
    return runs


# ------------------------------------------------------------------ slide 23

def build_s23(slide):
    """VP-S023: annotate the unchanged prompt with its requirement types."""
    set_slide_bg(slide, BG)
    mp = get_shape(slide, "main-prompt")
    paras = [t for t in para_texts(mp) if t.strip()]
    assert len(paras) == 3, paras
    p1, p2, p3 = paras
    assert p1.startswith("For a manager") and "TOTAL row" in p2 and "ranked table" in p3
    delete_shape(mp)

    groups = [
        # (y, text, phrases-to-underlay, label paragraphs)
        (2.12, p1, ["For a manager choosing supplier follow-up", "supplied workbook"],
         "Audience + source"),
        (3.25, p2, ["total returns ÷ total units shipped", "Exclude the existing TOTAL row"],
         "method"),
        (4.34, p3, ["ranked table", "brief, neutral explanation",
                    "State missing information and reconcile totals"],
         "format + tone + check"),
    ]
    heights = [0.12 + wrap_lines(t, 24, 9.72) * 0.42 for _, t, _, _ in groups]
    # one white specification panel behind all three groups, sized to the
    # measured content so the last bracket stays inside it
    panel_bottom = groups[-1][0] + heights[-1] + 0.14
    add_panel(slide, "pilot-spec-panel", 2.60, 2.02, 10.07, panel_bottom - 2.02,
              fill=PANEL, line=RULE, weight=1.0)

    for i, (y, text, phrases, label) in enumerate(groups):
        gh = heights[i]
        add_box(slide, f"pilot-group-{i}", 2.75, y, 9.92, gh,
                [segmented(text, phrases, 24)])
        # bracket opening toward the text
        br = add_panel(slide, f"pilot-bracket-{i}", 2.56, y + 0.04, 0.11, gh - 0.04,
                       fill=None, line=TEAL, weight=1.5,
                       shape=MSO_SHAPE.LEFT_BRACKET, back=False)
        br.adjustments[0] = 0.18
        add_box(slide, f"pilot-label-{i}", 0.60, y, 1.92, 1.05,
                [[{"t": label, "pt": 18, "bold": True, "color": TEAL_DARK}]])


# ------------------------------------------------------------------ slide 28

def build_s28(slide):
    """VP-S028: full editable Plan-Do-Check-Act-Plan cycle, 2x2 clockwise."""
    set_slide_bg(slide, BG)
    rows = []
    for i in range(4):
        lab = get_shape(slide, f"row-{i}-label")
        bod = get_shape(slide, f"row-{i}-body")
        rows.append((lab.text_frame.text, bod.text_frame.text))
        delete_shape(lab)
        delete_shape(bod)
    labels = [r[0] for r in rows]
    assert labels == ["PLAN", "DO", "CHECK", "ACT"], labels

    CW, CH = 5.45, 1.70  # spec envelope 1.58 adjusted for substitute-font wrap
    pos = {  # 2x2 clockwise: PLAN UL, DO UR, CHECK LR, ACT LL
        "PLAN": (0.67, 2.08), "DO": (7.22, 2.08),
        "CHECK": (7.22, 4.22), "ACT": (0.67, 4.22),
    }
    for label, body in rows:
        cx, cy = pos[label]
        add_panel(slide, f"pilot-card-{label}", cx, cy, CW, CH, fill=PANEL, line=RULE, weight=1.0)
        add_box(slide, f"pilot-card-{label}-label", cx + 0.08, cy + 0.05, 2.4, 0.42,
                [[{"t": label, "pt": 21.75, "bold": True, "color": TEAL_DARK}]])
        add_box(slide, f"pilot-card-{label}-body", cx + 0.06, cy + 0.47, CW - 0.14, CH - 0.52,
                [[{"t": body, "pt": 22.5, "bold": False, "color": INK}]])

    # clockwise arrows; the Act->Plan return is mandatory
    add_arrow(slide, "pilot-arrow-plan-do", 6.20, 2.93, 7.14, 2.93)
    add_arrow(slide, "pilot-arrow-do-check", 9.945, 3.84, 9.945, 4.16)
    add_arrow(slide, "pilot-arrow-check-act", 7.14, 5.07, 6.20, 5.07)
    add_arrow(slide, "pilot-arrow-act-plan", 3.395, 4.16, 3.395, 3.84)


# ------------------------------------------------------------------ slide 34

S34_TEXTS = {}  # captured from baseline for reuse in A and B


def capture_s34(slide):
    mp = get_shape(slide, "main-prompt")
    ps = para_texts(mp)
    assert len(ps) == 5, ps
    S34_TEXTS.update({"file": ps[0], "l1": ps[1], "q1": ps[2], "l2": ps[3], "q2": ps[4]})
    assert S34_TEXTS["q1"].endswith("then wait.”")
    assert S34_TEXTS["l2"] == "PROMPT 2 — ANALYZE"


def build_s34_a(slide):
    """Formatting-only version A: full-width separation of the complete
    prompts; body raised 17->18 pt under the approved standing exception."""
    set_slide_bg(slide, BG)
    mp = get_shape(slide, "main-prompt")
    delete_shape(mp)
    t = S34_TEXTS

    # choose 18 pt if both quotes stay at <=3 conservative lines, else 17
    body_pt = 18 if max(wrap_lines(t["q1"], 18, 11.4), wrap_lines(t["q2"], 18, 11.4)) <= 3 else 17
    if body_pt > 17:
        for key, shp in (("l1", "prompt-1 label"), ("q1", "prompt-1 text"),
                         ("l2", "prompt-2 label"), ("q2", "prompt-2 text")):
            log_font(34, shp, t[key], 17, body_pt,
                     "legibility of dense instructional prompt text (approved <20 pt exception)")

    add_box(slide, "pilot-file-line", 0.667, 2.08, 12.0, 0.44,
            [[{"t": t["file"], "pt": 18, "bold": True, "italic": False, "color": INK}]])

    q1_lines = wrap_lines(t["q1"], body_pt, 11.4)
    p1h = 0.14 + 0.34 + q1_lines * (body_pt / 72.0 * 1.22) + 0.10
    add_panel(slide, "pilot-p1-panel", 0.667, 2.64, 12.0, p1h, fill=PANEL, line=RULE, weight=1.0)
    add_box(slide, "pilot-p1", 0.75, 2.70, 11.84, p1h - 0.10,
            [[{"t": t["l1"], "pt": body_pt, "bold": True, "italic": False, "color": AMBER_DARK()}],
             [{"t": t["q1"], "pt": body_pt, "bold": False, "italic": False, "color": INK}]])

    y2 = 2.64 + p1h + 0.24
    q2_lines = wrap_lines(t["q2"], body_pt, 11.4)
    p2h = 0.14 + 0.34 + q2_lines * (body_pt / 72.0 * 1.22) + 0.10
    add_panel(slide, "pilot-p2-panel", 0.667, y2, 12.0, p2h, fill=PANEL, line=RULE, weight=1.0)
    add_box(slide, "pilot-p2", 0.75, y2 + 0.06, 11.84, p2h - 0.10,
            [[{"t": t["l2"], "pt": body_pt, "bold": True, "italic": False, "color": AMBER_DARK()}],
             [{"t": t["q2"], "pt": body_pt, "bold": False, "italic": False, "color": INK}]])
    bottom = y2 + p2h
    assert bottom <= 6.10, f"slide 34 A overflow: {bottom}"


def AMBER_DARK():
    return "875012"  # DO-mode label brown (matches existing chip text)


def build_s34_b(slide, crop_png, crop_w, crop_h):
    """Version B: same wording, baseline 17 pt, plus authentic workbook crop."""
    for name in ["pilot-file-line", "pilot-p1-panel", "pilot-p1",
                 "pilot-p2-panel", "pilot-p2"]:
        delete_shape(get_shape(slide, name))
    t = S34_TEXTS

    # top row: file line (left) + crop panel (right)
    crop_panel_h = 1.24  # image band + filename caption
    add_box(slide, "pilot-file-line", 0.667, 2.06, 5.10, 1.16,
            [[{"t": t["file"], "pt": 18, "bold": True, "italic": False, "color": INK}]])
    panel_w = 6.85
    add_panel(slide, "pilot-crop-panel", 5.82, 2.02, panel_w, crop_panel_h,
              fill=PANEL, line=RULE, weight=1.0)
    # fit image inside panel preserving aspect ratio
    max_w, max_h = panel_w - 0.24, crop_panel_h - 0.32
    scale = min(max_w / crop_w, max_h / crop_h)
    iw, ih = crop_w * scale, crop_h * scale
    ix = 5.82 + (panel_w - iw) / 2.0
    slide.shapes.add_picture(crop_png, Inches(ix), Inches(2.08), Inches(iw), Inches(ih))
    add_box(slide, "pilot-crop-caption", 5.94, 2.02 + crop_panel_h - 0.28, panel_w - 0.24, 0.26,
            [[{"t": "Supplier_Data_Exercise.xlsx", "pt": 10.5, "bold": False, "color": GRAY}]],
            align=PP_ALIGN.CENTER)

    body_pt = 17  # baseline size retained in B (crop consumes the space)
    line_h = body_pt / 72.0 * 1.20
    y1 = 2.02 + crop_panel_h + 0.10
    q1_lines = wrap_lines(t["q1"], body_pt, 11.5)
    p1h = 0.06 + 0.30 + q1_lines * line_h + 0.04
    add_panel(slide, "pilot-p1-panel", 0.667, y1, 12.0, p1h, fill=PANEL, line=RULE, weight=1.0)
    box1 = add_box(slide, "pilot-p1", 0.75, y1 + 0.03, 11.84, p1h - 0.05,
                   [[{"t": t["l1"], "pt": body_pt, "bold": True, "italic": False, "color": AMBER_DARK()}],
                    [{"t": t["q1"], "pt": body_pt, "bold": False, "italic": False, "color": INK}]])
    box1.text_frame.margin_top = Inches(0.02)
    y2 = y1 + p1h + 0.10
    q2_lines = wrap_lines(t["q2"], body_pt, 11.5)
    p2h = 0.06 + 0.30 + q2_lines * line_h + 0.04
    add_panel(slide, "pilot-p2-panel", 0.667, y2, 12.0, p2h, fill=PANEL, line=RULE, weight=1.0)
    box2 = add_box(slide, "pilot-p2", 0.75, y2 + 0.03, 11.84, p2h - 0.05,
                   [[{"t": t["l2"], "pt": body_pt, "bold": True, "italic": False, "color": AMBER_DARK()}],
                    [{"t": t["q2"], "pt": body_pt, "bold": False, "italic": False, "color": INK}]])
    box2.text_frame.margin_top = Inches(0.02)
    bottom = y2 + p2h
    assert bottom <= 6.18, f"slide 34 B overflow: {bottom}"


# ------------------------------------------------------------------ slide 37

def build_s37(slide):
    """VP-S037: left block reflow + labeled schematic wireframe, no results."""
    set_slide_bg(slide, BG)
    req_label = get_shape(slide, "row-1-label").text_frame.text
    req_body = get_shape(slide, "row-1-body").text_frame.text
    tl_label = get_shape(slide, "row-2-label").text_frame.text
    tl_body = get_shape(slide, "row-2-body").text_frame.text
    assert req_label == "It must include" and tl_label == "The teaching line"
    for n in ["row-1-label", "row-1-body", "row-2-label", "row-2-body"]:
        delete_shape(get_shape(slide, n))

    # reclaim empty masthead space: shrink the title box's EMPTY height only
    # and move the full-width prompt row up (reflow; wording untouched)
    title = get_shape(slide, "title")
    title.height = Inches(0.90)
    p_label = get_shape(slide, "row-0-label")
    p_body = get_shape(slide, "row-0-body")
    p_label.top = Inches(1.75)
    p_body.top = Inches(1.75)
    p_body.height = Inches(1.05)
    for para in p_body.text_frame.paragraphs:
        for r in para.runs:
            r.font.size = Pt(18)
    log_font(37, "prompt quote (row-0-body)", p_body.text_frame.text, 17, 18,
             "legibility of the anchor prompt (approved <20 pt exception)")
    log_font(37, "teaching-line body", tl_body, 17, 18,
             "legibility of the spoken teaching line (approved <20 pt exception)")
    # requirements body: the 18 pt increase did NOT fit beside the wireframe;
    # baseline 17 pt retained (reported per handoff section 4)
    add_box(slide, "pilot-req-label", 0.667, 2.86, 3.2, 0.40,
            [[{"t": req_label, "pt": 21.75, "bold": True, "color": TEAL_DARK}]])
    add_box(slide, "pilot-req-body", 0.667, 3.26, 6.95, 1.85,
            [[{"t": req_body, "pt": 17, "bold": False, "color": INK}]])
    add_box(slide, "pilot-tl-label", 0.667, 5.22, 3.2, 0.40,
            [[{"t": tl_label, "pt": 21.75, "bold": True, "color": TEAL_DARK}]])
    add_box(slide, "pilot-tl-body", 0.667, 5.62, 6.95, 0.60,
            [[{"t": tl_body, "pt": 18, "bold": False, "color": INK}]])

    # ---- wireframe panel (7.82, 3.0, 4.83, 2.96) ----
    WX, WY, WW, WH = 7.82, 3.00, 4.83, 2.96
    add_panel(slide, "pilot-wf-panel", WX, WY, WW, WH, fill=PANEL, line=RULE, weight=1.25)
    add_box(slide, "pilot-wf-schematic", WX + 0.06, WY + 0.01, 2.2, 0.34,
            [[{"t": "Schematic", "pt": 17, "bold": True, "color": REF_GRAY}]])

    def chip(name, x, y, w, h, label, pt=17):
        add_panel(slide, name, x, y, w, h, fill=BG, line=CONN, weight=1.0, back=False)
        if label:
            box = add_box(slide, name + "-t", x + 0.02, y, w - 0.04, h,
                          [[{"t": label, "pt": pt, "bold": False, "color": SLATE}]],
                          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            tf = box.text_frame
            tf.margin_left = Inches(0.02)
            tf.margin_right = Inches(0.02)
            tf.margin_top = Inches(0.0)
            tf.margin_bottom = Inches(0.0)

    inner_x = WX + 0.14
    inner_w = WW - 0.28
    # row 1: date-range selector / metric selector / reset control
    r1y, r1h = WY + 0.36, 0.58
    c1w = (inner_w - 0.16) * 0.40
    c2w = (inner_w - 0.16) * 0.34
    c3w = (inner_w - 0.16) * 0.26
    chip("pilot-wf-date", inner_x, r1y, c1w, r1h, "Date-range selector")
    chip("pilot-wf-metric", inner_x + c1w + 0.08, r1y, c2w, r1h, "metric selector")
    reset_x = inner_x + c1w + c2w + 0.16
    chip("pilot-wf-reset", reset_x, r1y, c3w, r1h, "reset control")
    # reset return indicator above the reset chip (editable arc, no promise)
    add_panel(slide, "pilot-wf-reset-return", reset_x + c3w / 2 - 0.12, WY + 0.05,
              0.24, 0.24, fill=CONN, line=None,
              shape=MSO_SHAPE.CIRCULAR_ARROW, back=False)
    # row 2: one wide selector for supplier and site, two empty selector boxes
    r2y, r2h = r1y + r1h + 0.10, 0.40
    chip("pilot-wf-supplier", inner_x, r2y, inner_w, r2h, None)
    sup = add_box(slide, "pilot-wf-supplier-t", inner_x + 0.06, r2y, inner_w - 1.36, r2h,
                  [[{"t": "supplier and site filters", "pt": 17, "bold": False, "color": SLATE}]],
                  align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
    sup.text_frame.margin_top = Inches(0.0)
    sup.text_frame.margin_bottom = Inches(0.0)
    for i in range(2):
        add_panel(slide, f"pilot-wf-selbox-{i}", inner_x + inner_w - 1.24 + i * 0.64,
                  r2y + 0.08, 0.56, r2h - 0.16, fill=PANEL, line=CONN, weight=1.0,
                  back=False)
    # three parallel drops from the selector row: summary, chart, table
    sy, sh_ = r2y + r2h + 0.16, 0.34
    cy = sy + sh_ + 0.08
    ch_ = WY + WH - cy - 0.12
    add_arrow(slide, "pilot-wf-drop-summary", inner_x + inner_w / 2, r2y + r2h,
              inner_x + inner_w / 2, sy, color=CONN, weight=1.25, head=True)
    add_arrow(slide, "pilot-wf-drop-chart", inner_x + 0.12, r2y + r2h,
              inner_x + 0.12, cy, color=CONN, weight=1.25, head=True)
    add_arrow(slide, "pilot-wf-drop-table", inner_x + inner_w - 0.12, r2y + r2h,
              inner_x + inner_w - 0.12, cy, color=CONN, weight=1.25, head=True)
    # summary strip (indented so the outer drops pass beside it)
    chip("pilot-wf-summary", inner_x + 0.30, sy, inner_w - 0.60, sh_, "summary values")
    # chart | table side by side
    chip("pilot-wf-chart", inner_x, cy, inner_w * 0.48, ch_, "trend chart")
    chip("pilot-wf-table", inner_x + inner_w * 0.48 + 0.08, cy,
         inner_w * 0.52 - 0.08, ch_, "sortable detail table")


# ------------------------------------------------------------------ slide 53

def build_s53(slide):
    """VP-S053: written brief vs configured boundary, checkpoint marked."""
    set_slide_bg(slide, BG)
    rb = get_shape(slide, "right-body")
    ps = para_texts(rb)
    assert ps[1] == "Configure available approvals.", ps

    # left: thin paper outline with a folded corner (top-right via flipV)
    paper = add_panel(slide, "pilot-paper", 0.55, 2.01, 5.83, 3.98,
                      fill=PANEL, line=RULE, weight=1.0,
                      shape=MSO_SHAPE.FOLDED_CORNER)
    paper._element.spPr.xfrm.set("flipV", "1")
    # right: stronger boundary outline
    add_panel(slide, "pilot-boundary", 6.80, 2.01, 6.02, 3.98,
              fill=PANEL, line=SLATE, weight=2.5)

    # checkpoint marker beside "Configure available approvals."
    # right-body top 2.76 + 0.05 inset; 24 pt line ~0.40 in -> para 2 center
    marker_cy = 2.76 + 0.05 + 0.40 + 0.20
    add_panel(slide, "pilot-checkpoint", 6.84, marker_cy - 0.08, 0.16, 0.16,
              fill=PANEL, line=AMBER, weight=1.5,
              shape=MSO_SHAPE.DIAMOND, back=False)


# ------------------------------------------------------------------ slide 87

def build_s87(slide):
    """VP-S087: proportional editable bars from the exact recorded fractions."""
    set_slide_bg(slide, BG)
    tbl = None
    for sh in slide.shapes:
        if sh.has_table:
            tbl = sh.table
    rows = [[c.text for c in r.cells] for r in tbl.rows]
    assert rows[1][2] == "74" and rows[2][2] == "262" and rows[3][2] == "96"
    bars = [  # exact spec geometry: x, y, w, h (inches)
        ("Alpha", 11.12, 3.1015625, 0.27111644018118836, 0.2),
        ("Bravo", 11.12, 3.8828125, 0.9583155990636305, 0.2),
        ("Cardinal", 11.12, 4.6640625, 0.3536124728763161, 0.2),
    ]
    for name, x, y, w, h in bars:
        sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sp.name = f"pilot-bar-{name}"
        sp.shadow.inherit = False
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(TEAL)
        sp.line.fill.background()


# ------------------------------------------------------------------ the crop

def make_crop():
    """Authentic crop of the participant input workbook: LibreOffice renders
    the real file; we crop the header row plus leading source rows."""
    tmp = os.path.join(PILOT_DIR, "_croptmp")
    os.makedirs(tmp, exist_ok=True)
    env = dict(os.environ, HOME="/root")
    # SinglePageSheets keeps every column on one page so the crop shows the
    # complete header row of the real file
    subprocess.run(["soffice", "--headless", "--convert-to",
                    'pdf:calc_pdf_Export:{"SinglePageSheets":{"type":"boolean","value":"true"}}',
                    "--outdir", tmp, XLSX_SRC], check=True, env=env,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pdf = os.path.join(tmp, "Supplier_Data_Exercise.pdf")
    subprocess.run(["pdftoppm", "-r", "200", "-f", "1", "-l", "1", "-png",
                    pdf, os.path.join(tmp, "page")], check=True)
    page = [f for f in sorted(os.listdir(tmp)) if f.startswith("page") and f.endswith(".png")][0]
    img = Image.open(os.path.join(tmp, page)).convert("L")
    px = img.load()
    W, H = img.size
    # find text rows (rows containing dark pixels)
    dark_rows = [y for y in range(H) if any(px[x, y] < 128 for x in range(0, W, 3))]
    # group into lines
    lines = []
    start = None
    prev = None
    for y in dark_rows:
        if start is None:
            start = y
        elif y - prev > 4:
            lines.append((start, prev))
            start = y
        prev = y
    if start is not None:
        lines.append((start, prev))
    n_lines = 9  # header + 8 source rows
    take = lines[:n_lines]
    y0 = max(0, take[0][0] - 12)
    y1 = min(H, take[-1][1] + 12)
    # horizontal extent of content
    dark_cols = [x for x in range(W) if any(px[x, y] < 128 for y in range(y0, y1, 2))]
    x0 = max(0, dark_cols[0] - 12)
    x1 = min(W, dark_cols[-1] + 12)
    crop = Image.open(os.path.join(tmp, page)).convert("RGB").crop((x0, y0, x1, y1))
    crop.save(CROP_PNG)
    n_source_rows = n_lines - 1
    return crop.size, n_source_rows


# ---------------------------------------------------------------------- main

def main():
    assert sha256(BASELINE) == BASELINE_SHA256, "baseline identity changed - reconcile first"
    os.makedirs(PILOT_DIR, exist_ok=True)

    prs = Presentation(BASELINE)
    sldIds = prs.slides._sldIdLst.findall(qn("p:sldId"))
    assert len(sldIds) == 103
    for pos, sid in TARGETS.items():
        assert int(sldIds[pos - 1].get("id")) == sid, f"slide {pos} identity mismatch"

    capture_s34(prs.slides[33])
    build_s23(prs.slides[22])
    build_s28(prs.slides[27])
    build_s34_a(prs.slides[33])
    build_s37(prs.slides[36])
    build_s53(prs.slides[52])
    build_s87(prs.slides[86])
    prs.save(OUT_A)
    print("saved", OUT_A)

    # version B: reopen A, change only slide 34
    (cw_px, ch_px), n_rows = make_crop()
    aspect = cw_px / ch_px
    prsb = Presentation(OUT_A)
    build_s34_b(prsb.slides[33], CROP_PNG, aspect, 1.0)
    prsb.save(OUT_B)
    print("saved", OUT_B)

    # font-change log
    with open(FONT_LOG_MD, "w") as f:
        f.write("# Pilot font-change log\n\n")
        f.write("Approved standing exception: text below 20 pt may be enlarged with wording,\n"
                "families, and emphasis unchanged (approvals.md, FONT). No reductions; nothing\n"
                "at or above 20 pt was changed. Version B keeps slide 34 body at the 17 pt\n"
                "baseline because the workbook crop consumes the vertical space the increase\n"
                "needs; see crop-provenance.md and the change report.\n\n")
        f.write("| Slide | Stable ID | Shape | Text | Original | Result | Reason |\n")
        f.write("|---|---|---|---|---|---|---|\n")
        for e in FONT_LOG:
            f.write(f"| {e['slide']} | sid-{e['sid']} | {e['shape']} | {e['text']} "
                    f"| {e['orig_pt']} pt | {e['new_pt']} pt | {e['reason']} |\n")
        f.write("\n## Increases evaluated but NOT applied (reported per handoff §4)\n\n"
                "- Slide 37 (sid-294), requirements body (\"Date-range selector · …\"): "
                "an 18 pt trial did not fit the left block beside the approved wireframe "
                "without crowding the teaching line; baseline 17 pt retained.\n"
                "- Slide 34 version B, both prompts: baseline 17 pt retained because the "
                "workbook crop consumes the vertical space the 18 pt increase needs; "
                "version A carries the increase.\n\n"
                "All other text on the six pilot slides keeps its baseline size, family, "
                "and emphasis. New annotation/wireframe labels are new objects at their "
                "specified sizes, not size changes to existing text.\n")
    print("wrote", FONT_LOG_MD)

    # crop provenance
    with open(CROP_MD, "w") as f:
        f.write("# Slide 34 version B crop provenance\n\n")
        f.write(f"- Source file: `references/exercise-data/Supplier_Data_Exercise.xlsx`\n")
        f.write(f"- Source SHA-256: `{sha256(XLSX_SRC)}`\n")
        f.write(f"- Displayed range: header row plus the first {n_rows} source detail rows "
                f"of the participant input sheet (all columns), rendered by LibreOffice "
                f"from the unmodified file at 200 DPI and cropped to the text extent; "
                f"aspect ratio preserved.\n")
        f.write(f"- Crop image: `notes/visual-polish/pilot/s34b_workbook_crop.png`, "
                f"SHA-256 `{sha256(CROP_PNG)}`\n")
        f.write("- No analyzed output, calculated finding, instructor key, or answer "
                "highlight appears in the crop; it is the raw exercise input only.\n")
    print("wrote", CROP_MD)

    with open(os.path.join(PILOT_DIR, "font_change_log.json"), "w") as f:
        json.dump(FONT_LOG, f, indent=1)


if __name__ == "__main__":
    main()
