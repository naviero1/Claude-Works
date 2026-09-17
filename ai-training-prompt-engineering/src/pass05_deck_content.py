#!/usr/bin/env python3
# Iteration-05 deck pass, stage 2 of 2: CONTENT.
# Rewrites the reordered application block (positions 33-44) to the nine-step
# Excel-first sequence, retargets the appendix pages it touches, renumbers the
# moved pages, and writes fresh MODE/TIME/PURPOSE/LAND THIS/SAY/DO/BRIDGE/
# IF LATE notes on every rewritten slide. Exercise answers live in notes and
# instructor keys, never on participant prompt slides. New body text >= 17pt.
import copy

from pptx import Presentation
from pptx.util import Emu, Pt

PPTX = __file__.rsplit('/', 1)[0] + '/../deliverables/From_Prompts_to_Agents_Facilitated_60_Minute.pptx'
IN = 914400

p = Presentation(PPTX)
slides = list(p.slides)
assert len(slides) == 103


def tshapes(sl):
    return [sh for sh in sl.shapes if sh.has_text_frame]


def classify(sl):
    """Return dict of roles for the standard body/row slide chrome."""
    roles = {'rows': []}
    for sh in tshapes(sl):
        x, y, w = sh.left / IN, sh.top / IN, sh.width / IN
        t = sh.text_frame.text.strip()
        if y < 0.6 and x > 8:
            roles['chip'] = sh
        elif y < 1.5 and w > 10:
            roles['title'] = sh
        elif 6.0 < y < 6.9 and w > 10:
            roles['band'] = sh
        elif y > 7.0 and x < 2:
            roles['footer'] = sh
        elif y > 7.0 and x > 11.5:
            roles['pageno'] = sh
        elif 1.8 < y < 6.0:
            roles['rows'].append(sh)
    roles['rows'].sort(key=lambda s: (s.top, s.left))
    return roles


def set_tf(tf, paras):
    """paras: list of (text, size_pt, bold, italic). Rebuilds the frame."""
    first = tf.paragraphs[0]
    for para in list(tf.paragraphs)[1:]:
        para._p.getparent().remove(para._p)
    for r in first.runs:
        r._r.getparent().remove(r._r)
    for i, (text, size, bold, italic) in enumerate(paras):
        para = first if i == 0 else tf.add_paragraph()
        run = para.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
    return tf


def one_run_text(sh, text, size=None):
    """Replace a uniform-format frame's text, keeping its first run's style."""
    tf = sh.text_frame
    para = tf.paragraphs[0]
    assert para.runs, sh.shape_id
    para.runs[0].text = text
    if size:
        para.runs[0].font.size = Pt(size)
    for r in para.runs[1:]:
        r.text = ''
    for extra in list(tf.paragraphs)[1:]:
        for r in extra.runs:
            r.text = ''


def copy_chip_style(src_sl, dst_sl, text):
    src = classify(src_sl).get('chip')
    dst = classify(dst_sl).get('chip')
    assert src is not None and dst is not None
    srun = src.text_frame.paragraphs[0].runs[0]
    drun = dst.text_frame.paragraphs[0].runs[0]
    drun.text = text
    drun.font.bold = srun.font.bold
    drun.font.size = srun.font.size
    if srun.font.color and srun.font.color.type is not None:
        drun.font.color.rgb = srun.font.color.rgb
    for r in dst.text_frame.paragraphs[0].runs[1:]:
        r.text = ''


def set_pageno(sl, n):
    r = classify(sl).get('pageno')
    assert r is not None
    one_run_text(r, str(n))


def set_notes(sl, text):
    tf = sl.notes_slide.notes_text_frame
    first = tf.paragraphs[0]
    for para in list(tf.paragraphs)[1:]:
        para._p.getparent().remove(para._p)
    for r in first.runs:
        r._r.getparent().remove(r._r)
    lines = text.strip().split('\n')
    for i, line in enumerate(lines):
        para = first if i == 0 else tf.add_paragraph()
        run = para.add_run()
        run.text = line


CHIP_DO = slides[33]      # position 34 pre-rewrite carries a DO chip
CHIP_TEACH = slides[36]   # position 37 carries TEACH
CHIP_REF = slides[98]     # appendix page carries REFERENCE SELF-STUDY

# =====================  position 33 — opener (divider layout)  ==============
sl = slides[32]
sh = tshapes(sl)
# divider shapes: [1] kicker, [2] big title, [3] subtitle, [4] supporting, [5] chip, [6] pageno
one_run_text(sh[1], 'PART 4 · APPLICATION')
one_run_text(sh[2], 'Data Analytics Exercise')
one_run_text(sh[3], 'One workbook. Questions. Charts. A returned file.')
one_run_text(sh[4], 'Analyze the data in Excel, return the workbook with editable charts, '
                    'then reuse it for the dashboard and presentation mock-up.', 20)
one_run_text(sh[5], 'TEACH  0:40')
one_run_text(sh[6], '33')
set_notes(sl, """MODE: TEACH
TIME: 0:40
PURPOSE: Open the application block with the whole arc in one breath.
LAND THIS: One dataset, uploaded once, becomes an analysis, charts, a returned file, a dashboard, and a presentation.
SAY: Everything in this block runs on one workbook you upload at the start. You will analyze it, question it, get it back with editable charts, feed the returned file to a dashboard, and end with a management mock-up.
SAY: Each added instruction in a prompt exists to resolve a specific ambiguity. A basic ask like "analyze the supplier data and tell me which supplier is worst" makes the model pick the rows, the metric, and "worst" for you, silently. The prompts you are about to run name the file, define the metric from sums, handle the missing value, and demand reconciliation - each line closes one gap. That lesson used to be its own slide; it now travels with every prompt you run today.
DO: Nothing yet - the upload comes on the next slide.
BRIDGE: Move to "Upload the Data File and Run These Two Prompts".
IF LATE: COMPRESS to the arc sentence.""")

# =====================  position 34 — upload + two prompts  =================
sl = slides[33]
r = classify(sl)
one_run_text(r['title'], 'Upload the Data File and Run These Two Prompts', 33)
set_tf(r['rows'][0].text_frame, [
    ('File: Supplier_Data_Exercise.xlsx — upload it, attach nothing else.', 18, True, False),
    ('PROMPT 1 — INSPECT', 17, True, False),
    ('“Before calculating anything, inspect the file. Tell me what one row represents, how many detail rows it contains, whether any values are missing, and whether any total or summary row could be counted twice. Tell me what you would include or exclude, then wait.”', 17, False, False),
    ('PROMPT 2 — ANALYZE', 17, True, False),
    ('“Using detail rows only, calculate each supplier’s return rate as total Units_Returned divided by total Units_Shipped. Rank suppliers from highest to lowest. Show the totals used, reconcile the overall totals to the source file, and state one limitation of the comparison.”', 17, False, False),
])
one_run_text(r['band'], 'Attempt first — the verified key stays with the instructor and the workbook keys.', 18)
one_run_text(r['footer'], 'From Prompts to Agents')
set_pageno(sl, 34)
copy_chip_style(CHIP_DO, sl, 'DO   6:00')
set_notes(sl, """MODE: DO
TIME: 6:00
PURPOSE: The single intake for the whole block - inspect before calculating, then compute checked rates.
LAND THIS: Inspection is a separate prompt that ends with "then wait" - the model must not calculate until the scope is agreed.
SAY: Upload the exercise file, run the inspect prompt, and read what comes back before you let it calculate.
DO: 0:00-1:00 upload and run PROMPT 1; 1:00-2:30 compare inspection findings across the room; 2:30-5:00 run PROMPT 2; 5:00-6:00 check rankings and reconciliation together.
ANSWER KEY (instructor only - not shown to participants before they attempt): 144 detail rows plus one TOTAL row; one missing Inspection_Hours value (missing, not zero); 224,902 units shipped and 432 returns; Bravo Plastics highest at about 0.348 percent (262 / 75,184); Cardinal 0.129 percent; Alpha 0.099 percent; the data do not establish causation. Full key: workbook tab KEY-Analysis.
BRIDGE: Move to "Ask a Follow-Up Question".
IF LATE: KEEP - this is the anchor exercise of the block.""")

# =====================  position 35 — follow-up question  ===================
sl = slides[34]
r = classify(sl)
one_run_text(r['title'], 'Ask a Follow-Up Question')
set_tf(r['rows'][0].text_frame, [
    ('Same chat, same data — one more prompt:', 18, True, False),
    ('“Using the same data, compare return rate by month and site. Identify the clearest increase, decrease, or unusual pattern. Show the values behind it. Do not claim a cause. List the additional data needed to investigate why it happened.”', 18, False, False),
    ('A finding names a pattern and shows its values. A cause needs data this file does not contain — the honest answer lists what to collect next.', 17, False, True),
])
one_run_text(r['band'], 'Findings, not causes — the pattern, its values, and the data you would need next.', 18)
one_run_text(r['footer'], 'From Prompts to Agents')
set_pageno(sl, 35)
copy_chip_style(CHIP_DO, sl, 'DO   2:30')
set_notes(sl, """MODE: DO
TIME: 2:30
PURPOSE: Business questioning of the same data - and the findings-versus-causes discipline.
LAND THIS: The model may offer explanations; the data only support patterns.
SAY: Ask the follow-up, then look for any causal language in the answer - that is the teaching moment.
DO: 1:30 run and read; 1:00 name one causal claim someone's answer made and strip it back to the finding.
BRIDGE: Move to "Add the Charts and Return the Excel File".
IF LATE: COMPRESS to the prompt and one example pattern.""")

# =====================  position 36 — Excel charts (4-row donor)  ===========
sl = slides[35]
r = classify(sl)
one_run_text(r['title'], 'Add the Charts and Return the Excel File', 33)
ROWS36 = [
    ('The prompt', 'Update the uploaded workbook and return an editable Excel file named Supplier_Data_Analyzed.xlsx — original data unchanged, analysis on a new Summary sheet.'),
    ('Summary sheet', 'Supplier comparison table (units shipped, units returned, return rate, defect rate) · editable native Excel column chart of supplier return rates · editable native line chart of return rate by month · period, metric definitions, one limitations note.'),
    ('The rules', 'Formulas, PivotTables, or other traceable Excel logic — never hard-coded chart values · missing values stay missing · summary totals reconcile to the source before the file returns.'),
    ('Prove it', 'Open the returned workbook, inspect the Summary sheet, click one chart — its source range or PivotTable must be editable.'),
]
for (label, desc), pair_i in zip(ROWS36, range(4)):
    lab, dsc = r['rows'][pair_i * 2], r['rows'][pair_i * 2 + 1]
    one_run_text(lab, label)
    one_run_text(dsc, desc, 17)
one_run_text(r['band'], 'Full prompt: workbook tab Excel-Charts · prepared fallback: Supplier_Data_Analyzed.xlsx.', 18)
one_run_text(r['footer'], 'From Prompts to Agents')
set_pageno(sl, 36)
copy_chip_style(CHIP_DO, sl, 'DO   4:00')
set_notes(sl, """MODE: DO
TIME: 4:00
PURPOSE: The assistant returns a real, editable artifact - and participants verify it is real.
LAND THIS: Native charts you can click and edit, formulas you can trace - not a picture of an analysis.
SAY: Run the full prompt from tab Excel-Charts. When the file comes back, we open it together.
DO: 2:00 run the prompt; 2:00 open the returned file, inspect Summary, click the column chart and show its editable source range; if generation stalls, open the prepared Supplier_Data_Analyzed.xlsx instead.
BRIDGE: Move to "Describe the Dashboard You Want" - the returned file is the dashboard's input.
IF LATE: Show the prepared fallback and click one chart; skip the live generation.""")

# =====================  position 37 — describe the dashboard (3-row)  =======
sl = slides[36]
r = classify(sl)
one_run_text(r['title'], 'Describe the Dashboard You Want', 33)
ROWS37 = [
    ('The prompt', '“Using the returned Supplier_Data_Analyzed.xlsx workbook and its verified definitions, create one self-contained HTML dashboard that works offline.”'),
    ('It must include', 'Date-range selector · supplier and site filters · metric selector · summary values · trend chart · sortable detail table · reset control. Every filter updates every view; show the selected period and metric definitions; display a clear message when no rows match.'),
    ('The teaching line', 'You do not need to write the code. You need to describe the behavior.'),
]
for (label, desc), pair_i in zip(ROWS37, range(3)):
    lab, dsc = r['rows'][pair_i * 2], r['rows'][pair_i * 2 + 1]
    one_run_text(lab, label)
    one_run_text(dsc, desc, 17)
one_run_text(r['band'], 'Extended dashboard vocabulary: Course_Workbook.xlsx and the reference layer — not this slide.', 18)
one_run_text(r['footer'], 'From Prompts to Agents')
set_pageno(sl, 37)
copy_chip_style(CHIP_TEACH, sl, 'TEACH   1:00')
set_notes(sl, """MODE: TEACH
TIME: 1:00
PURPOSE: Requirements for interactive behavior, in plain language, before anything renders.
LAND THIS: You do not need to write the code. You need to describe the behavior.
SAY: Read the prompt aloud - every clause is a behavior you could check by clicking.
BRIDGE: Move to "Build and Test the Dashboard".
IF LATE: COMPRESS to the teaching line.""")

# =====================  position 38 — build & test dashboard  ===============
sl = slides[37]
r = classify(sl)
one_run_text(r['title'], 'Build and Test the Dashboard', 33)
set_tf(r['rows'][0].text_frame, [
    ('Input: Supplier_Data_Analyzed.xlsx — the returned workbook: same data, definitions, and period.', 18, True, False),
    ('1  Open the generated HTML file.', 18, False, False),
    ('2  Filter to Berlin, Bravo Plastics, March through August 2026.', 18, False, False),
    ('3  Verify 8,575 units, 21 returns, and a 0.245 percent return rate against the source.', 18, False, False),
    ('Then ask for ONE refinement using the vocabulary — a different date range, grouping, metric, drill-down, sort, or reset behavior.', 17, False, True),
])
one_run_text(r['band'], 'Check: 8,575 units · 21 returns · 0.245% — the same numbers, from the same verified file.', 18)
set_pageno(sl, 38)
copy_chip_style(CHIP_DO, sl, 'DO   4:00')
set_notes(sl, """MODE: DO
TIME: 4:00
PURPOSE: The dashboard consumes the verified workbook - and is trusted only after the filtered check.
LAND THIS: Same data, same definitions, same numbers - the dashboard is the analysis wearing a different interface.
SAY: Three actions only: open, filter, verify. Then you name one refinement.
DO: 1:30 generate or open the prepared references/exercise-data/Supplier_Quality_Dashboard.html; 1:30 run the three actions; 1:00 collect refinement asks from the room.
BRIDGE: Move to "Turn the Findings into a Five-Slide Management Mock-up".
IF LATE: Run the three actions on the prepared file; skip refinements.""")

# ============  position 39 — management mock-up (4-row donor)  ==============
sl = slides[38]
r = classify(sl)
one_run_text(r['title'], 'Turn the Findings into a Five-Slide Management Mock-up', 30)
ROWS39 = [
    ('The prompt', '“Using only the verified findings from this exercise, create an editable five-slide PowerPoint mock-up for [audience] to support [decision].”'),
    ('The five slides', '1 Question and scope · 2 Supplier comparison · 3 The most important trend · 4 Recommended follow-up · 5 Limitations and next steps.'),
    ('The rules', 'One main message per slide · readable charts, concise visible text, speaker notes · findings separated from recommendations · no invented causes or commitments.'),
    ('Why it differs', 'The dashboard supports exploration; the presentation supports a management decision — the final step of the connected data exercise.'),
]
for (label, desc), pair_i in zip(ROWS39, range(4)):
    lab, dsc = r['rows'][pair_i * 2], r['rows'][pair_i * 2 + 1]
    one_run_text(lab, label)
    one_run_text(dsc, desc, 17)
one_run_text(r['band'], 'Full prompt: workbook tab 3-Present-Findings · reference page 99 · fallback: Supplier_Quality_Mock_Presentation.pptx.', 17)
one_run_text(r['footer'], 'From Prompts to Agents')
set_pageno(sl, 39)
copy_chip_style(CHIP_DO, sl, 'DO   2:30')
set_notes(sl, """MODE: DO
TIME: 2:30
PURPOSE: Close the data arc on a decision artifact built only from verified findings.
LAND THIS: Exploration and decision are different jobs - the same verified numbers wear two interfaces.
SAY: Fill the audience and the decision, run it, and open the fallback beside the result.
DO: 1:30 run with a chosen [audience]/[decision]; 1:00 compare one slide against the prepared mock deck's discipline (findings left, judgment right).
BRIDGE: Move to "Organize an Email Thread in Outlook".
IF LATE: Open the prepared mock deck and show slide 4's findings/recommendations split.""")

# ============  position 40 — Outlook choices (4-row donor)  =================
sl = slides[39]
r = classify(sl)
set_tf(r['title'].text_frame, [
    ('Organize an Email Thread in Outlook', 32, True, False),
    ('Choose the result you need', 17, False, True),
])
ROWS40 = [
    ('Catch me up', '“Summarize the current situation in three sentences.”'),
    ('What changed · Who owes what', '“List changed dates or decisions and show the latest version.” · “Create a table with task, owner, due date, and dependency.”'),
    ('What remains open', '“List unanswered questions, approval conditions, and missing attachments.”'),
    ('Prepare my reply', '“Draft a concise response. Use brackets where I still need to decide. Do not invent commitments.”'),
]
for (label, desc), pair_i in zip(ROWS40, range(4)):
    lab, dsc = r['rows'][pair_i * 2], r['rows'][pair_i * 2 + 1]
    one_run_text(lab, label, 19)
    one_run_text(dsc, desc, 17)
one_run_text(r['band'], 'Use only messages and attachments you can access · write “Not stated” instead of guessing · cite the supporting sender and date · do not send anything.', 17)
set_pageno(sl, 40)
copy_chip_style(CHIP_TEACH, sl, 'TEACH   1:00')
set_notes(sl, """MODE: TEACH
TIME: 1:00
PURPOSE: One thread, five results - choosing the result IS the prompt.
LAND THIS: Pick the outcome first; the prompt is one sentence once you have.
SAY: Workflow: run Outlook's built-in Copilot summary first, then use Copilot to refine it with one of these. If that feature is unavailable, paste the supplied fictional thread into an approved assistant - same prompts, same results.
ACCOUNT NOTE (verified Sep 2026, notes/research/r28): the full Copilot-in-Outlook experience depends on the organization's Copilot license and current Outlook version; the pasted-text path always works. Never practice on real confidential threads.
BRIDGE: Move to "Outlook Exercise: Decisions, Actions, and Open Questions".
IF LATE: COMPRESS to Catch-me-up and Prepare-my-reply.""")

# ============  position 41 — Outlook structured exercise  ===================
sl = slides[40]
r = classify(sl)
one_run_text(r['title'], 'Outlook Exercise: Decisions, Actions, and Open Questions', 29)
set_tf(r['rows'][0].text_frame, [
    ('“Review this thread using only the messages and attachments you can access. Give me the current status, confirmed decisions and conditions, action items with owner and due date, changed dates or commitments, unanswered questions, missing attachments, and messages that need a reply. Write ‘Not stated’ when the thread does not say. Cite the supporting message by sender and date.”', 18, False, False),
    ('Run it after the built-in summary, in the Copilot chat pane on the same conversation — or paste Packaging_Change_Thread.txt into an approved assistant.', 17, False, True),
])
one_run_text(r['band'], 'Your brief is checked against the instructor key — the thread has planted traps. Attempt first.', 18)
one_run_text(r['footer'], 'Fictional source thread: slides 88–89 · Packaging_Change_Thread.txt in your pack')
set_pageno(sl, 41)
copy_chip_style(CHIP_DO, sl, 'DO   3:00')
set_notes(sl, """MODE: DO
TIME: 3:00
PURPOSE: The structured extraction on a thread with planted traps.
LAND THIS: "Not stated" and a citation per claim - the two habits that survive contact with a real inbox.
SAY: Run the structured prompt, then we check against the key.
DO: 1:30 run; 1:30 walk the key.
ANSWER KEY (instructor only): change date moved Sep 25 to Oct 2, conditional on quality sign-off · trial date Sep 25 supersedes Sep 18 · the 18,000 cap holds and freight is NOT approved · Luis's revised plan due Sep 16 · freight responsibility unresolved · the drawing is referenced but not accessible. Full key: tab KEY-Email + instructor-keys/Packaging_Change_Expected_Brief.md.
BRIDGE: Move to "Extract Supplier Quotes into Excel".
IF LATE: KEEP - this is the email anchor.""")

# ============  position 42 — quote extraction (body donor copy)  ============
sl = slides[41]
r = classify(sl)
one_run_text(r['title'], 'Extract Supplier Quotes into Excel', 33)
set_tf(r['rows'][0].text_frame, [
    ('“Review the attached supplier quotation files and create an editable Excel workbook. Create a Raw Extraction sheet with one row per quotation: supplier, quoted scope, currency, quoted quantity, unit price, tooling or other one-time charges, freight, taxes, lead time, payment terms, quote validity, exclusions, source file, source page, and missing information.”', 17, False, False),
    ('“Copy values exactly as stated. Keep currencies, quantities, and units in their original form. Write ‘Not stated’ when a field is absent. Do not normalize, rank, or recommend yet. Cite the source file and page for every extracted commercial value.”', 17, False, False),
    ('Step one is faithful extraction and traceability — judgment comes next.', 17, False, True),
])
one_run_text(r['band'], 'Files: Quote_Alpha_Components.pdf · Quote_Bravo_Plastics.pdf · Quote_Cardinal_Metals.pdf — full prompt: workbook tab Quote-Extract.', 17)
one_run_text(r['footer'], 'From Prompts to Agents')
set_pageno(sl, 42)
copy_chip_style(CHIP_DO, sl, 'DO   2:30')
set_notes(sl, """MODE: DO
TIME: 2:30
PURPOSE: Documents-to-spreadsheet, step one: extraction with provenance, no judgment.
LAND THIS: Verbatim values, original currencies, "Not stated" for silence, a citation per value.
SAY: Attach the three PDFs and run it. Anything converted, ranked, or recommended at this step is a miss.
DO: 1:30 run; 1:00 spot-check two extracted values against the PDFs on screen.
BRIDGE: Move to "Normalize and Compare the Supplier Quotes".
IF LATE: SKIP to the prepared Quote_Comparison_Workbook.xlsx in the next slide's debrief.""")

# ============  position 43 — quote comparison (body donor copy)  ============
sl = slides[42]
r = classify(sl)
one_run_text(r['title'], 'Normalize and Compare the Supplier Quotes', 33)
set_tf(r['rows'][0].text_frame, [
    ('“Using the Raw Extraction sheet, create a Normalized Comparison sheet. Normalize quantities, units, or currencies only when the conversion rule or basis is supplied. Keep every original extracted value visible. Flag differences in scope, assumptions, exclusions, and commercial terms. Show formulas for comparable totals.”', 17, False, False),
    ('“Do not identify a best quote until the comparison basis is complete. List the missing information and unresolved questions that prevent a fair recommendation.”', 17, False, False),
    ('The trap: the cheapest-looking quote stops looking cheapest once tooling, freight, currency, and scope are on the table.', 17, False, True),
])
one_run_text(r['band'], 'No winner while the basis is incomplete — prepared example: Quote_Comparison_Workbook.xlsx · tab Quote-Compare.', 17)
one_run_text(r['footer'], 'From Prompts to Agents')
set_pageno(sl, 43)
copy_chip_style(CHIP_DO, sl, 'DO   2:30')
set_notes(sl, """MODE: DO
TIME: 2:30
PURPOSE: Step two: normalize only what has a basis; refuse a premature winner.
LAND THIS: A fair comparison names its blocking gaps before it names a favorite.
SAY: Continue in the same chat from the extraction. Watch whether it converts euros without a rate - that is the discipline test.
DO: 1:30 run; 1:00 open the prepared workbook and read the blocking-gap list aloud.
ANSWER SKETCH (instructor only): at 5,000 units with tooling amortized - Alpha USD 15.60/unit (freight excluded), Bravo EUR 14.70/unit (not convertible without a supplied rate; freight, duties, VAT excluded), Cardinal USD 16.40/unit all-in (DDP, certs, longest warranty). Correct outcome: NO best quote yet. Key: instructor-keys/Quote_Comparison_Key.md.
BRIDGE: Move to "Turn Reference Files into a Research Spreadsheet".
IF LATE: Open the prepared workbook; skip live generation.""")

# ============  position 44 — research organizer (4-row donor)  ==============
sl = slides[43]
r = classify(sl)
one_run_text(r['title'], 'Turn Reference Files into a Research Spreadsheet', 31)
ROWS44 = [
    ('The prompt', '“Review the attached reference files and build a research workbook using only information you can trace to a source.”'),
    ('Evidence sheet', 'One claim per row: ID · topic · claim · source file · author · date · page or section · short excerpt · plain-language interpretation · caveat · relevance to [research question] · confidence · open question.'),
    ('Synthesis + Sources', 'Synthesis orders findings from foundations to implications — agreements, conflicts, gaps. Sources indexes every file. “Not stated” for gaps · conflicting claims stay separate · no invented citations.'),
    ('Delivery', 'Editable Excel workbook — filters, frozen header row, wrapped text; separate CSV tables if the tool cannot produce Excel.'),
]
for (label, desc), pair_i in zip(ROWS44, range(4)):
    lab, dsc = r['rows'][pair_i * 2], r['rows'][pair_i * 2 + 1]
    one_run_text(lab, label)
    one_run_text(dsc, desc, 17)
one_run_text(r['band'], 'Pack: references/exercise-data/research-pack/ · full prompt: workbook tab Research · prepared example: Research_Workbook.xlsx.', 17)
one_run_text(r['footer'], 'From Prompts to Agents')
set_pageno(sl, 44)
copy_chip_style(CHIP_DO, sl, 'DO   3:00')
set_notes(sl, """MODE: DO
TIME: 3:00
PURPOSE: The closing exercise: unstructured references become linear, traceable knowledge.
LAND THIS: Provenance survives - every row cites its file and page, conflicts stay visible, gaps say "Not stated".
SAY: The pack hides one deliberate conflict and several unstated details. A good workbook surfaces them instead of resolving them silently.
DO: 1:30 run on the pack; 1:30 open the prepared example - find the cadence conflict in the Synthesis and one "Not stated".
ANSWER SKETCH (instructor only): the planted conflict is review cadence - the manual (2026-02) says quarterly, the newsletter (2026-06) reports monthly; must be kept as two rows and surfaced, never merged. Key: instructor-keys/Research_Workbook_Key.md.
BRIDGE: Move to "Part 5: agentic work".
IF LATE: Open the prepared Research_Workbook.xlsx and show the Evidence and Synthesis sheets.""")

# ============  appendix moves: 102 (study aid), 103 (quote extension)  ======
for pos, note_tail in ((102, 'Optional self-study; no live allocation. Formerly the live transition into the application block.'),
                       (103, 'Optional self-study; no live allocation. The live course now runs the quotation work as two steps (slides 42-43); this page keeps the original one-step version.')):
    sl = slides[pos - 1]
    r = classify(sl)
    if 'chip' in r:
        one_run_text(r['chip'], 'REFERENCE  SELF-STUDY')
    else:  # divider-style chip position on old-33
        for sh in tshapes(sl):
            if 'REFERENCE' in sh.text_frame.text or 'DIVIDER' in sh.text_frame.text:
                one_run_text(sh, 'REFERENCE  SELF-STUDY')
    # page number: the digit-only run
    done = False
    for sh in tshapes(sl):
        t = sh.text_frame.text.strip()
        if t.isdigit():
            one_run_text(sh, str(pos))
            done = True
    assert done, pos
    txt = sl.notes_slide.notes_text_frame.text
    set_notes(sl, f"""MODE: REFERENCE
TIME: 0:00 ({note_tail})
{txt}""")

# ============  appendix retargets  =========================================
def sub_in_slide(pos, where, old, new, expect=1):
    sl = slides[pos - 1]
    frames = ([sh.text_frame for sh in tshapes(sl)] if where in ('body', 'both') else [])
    if where in ('notes', 'both') and sl.has_notes_slide:
        frames.append(sl.notes_slide.notes_text_frame)
    hits = 0
    for tf in frames:
        for para in tf.paragraphs:
            full = ''.join(x.text for x in para.runs)
            if old in full:
                fmts = set((x.font.bold, x.font.size.pt if x.font.size else None)
                           for x in para.runs if x.text)
                assert len(fmts) <= 1, (pos, old[:40], 'multi-format paragraph')
                para.runs[0].text = full.replace(old, new)
                for x in para.runs[1:]:
                    x.text = ''
                hits += 1
    assert hits == expect, (pos, old[:50], hits)

sub_in_slide(81, 'body', 'Related live slide 35', 'Related live slide 33')
sub_in_slide(82, 'body', 'Related live slide 42', 'Related live slide 40')
sub_in_slide(99, 'body', 'live orientation: slide 41.', 'live orientation: slide 39.')
sub_in_slide(99, 'notes', 'RELATED: slide 41 (live orientation)', 'RELATED: slide 39 (live orientation)')
sub_in_slide(100, 'body', 'Inventory replenishment (the live demonstration) · the water cycle · how an internet message travels',
             'Supplier quality (the live example) · household budgeting · the water cycle · how an internet message travels')
sub_in_slide(100, 'body', 'live demonstration: slide 44.', 'reference exercise (no live slot).')
sub_in_slide(100, 'notes', 'RELATED: slide 44 (live demonstration);', 'RELATED: reference-layer exercise (no live slot);')
sub_in_slide(98, 'body', 'One self-contained offline HTML file from the cleaned supplier records',
             'One self-contained offline HTML file from the returned Supplier_Data_Analyzed.xlsx workbook')
sub_in_slide(98, 'notes', 'Supplier_Data_Clean.csv', 'Supplier_Data_Analyzed.xlsx (the returned workbook)')
sub_in_slide(98, 'notes', 'live demonstration slide 40', 'live demonstration slide 38')

# slide 32 bridges into the new opener
sub_in_slide(32, 'notes', 'Move to “Optional practice: a course-log study aid”.',
             'Move to “Data Analytics Exercise”.')

p.save(PPTX)
print('content stage saved: 12 block slides rewritten, appendix renumbered and retargeted')
