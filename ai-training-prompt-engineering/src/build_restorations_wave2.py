#!/usr/bin/env python3
"""Restoration wave 2 — owner-directed (in session, 2026-09-17).

Ports three owner-valued slides verbatim from the preserved original deck
(references/From_Prompts_to_Agents_Training.pptx) into the working deck as
appendix reference pages, by deep-copying their shape XML and re-embedding
their images so layout, typography, colors, icons, and charts carry exactly:

  106  Fluency is not evidence            (old 14: mechanism + measured chart + four personas)
  107  Real, public, verified - and all avoidable  (old 15: the hall of shame)
  108  The flattery bias - measured, and all over the news  (old 29)

Every dated claim was re-verified against notes/research (r14 hallucination
incidents; r25 sycophancy dossier: Cheng et al. Science 2026 +49% across 11
models; OpenAI Apr 2025 sycophancy rollback; NYT Jun 2025; Microsoft AI CEO
Aug 2025; o4-mini 24/75 vs newer 22/26/52; EY 16/27; KPMG 5/45; ~1,500
court-decision tracker). Old footers are replaced with the appendix chrome;
titles are pinned to Cambria to match the reference-layer idiom.

Input: pilot v03.  Output: From_Prompts_to_Agents_Visual_Pilot_v04.pptx
"""
import copy
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "notes", "visual-polish", "pilot", "From_Prompts_to_Agents_Visual_Pilot_v03.pptx")
OLD = os.path.join(ROOT, "references", "From_Prompts_to_Agents_Training.pptx")
OUT = os.path.join(ROOT, "notes", "visual-polish", "pilot", "From_Prompts_to_Agents_Visual_Pilot_v04.pptx")

TEAL = "0E7C7B"
INK = "232A31"
CHROME = "526267"
GRAY = "657278"
REF_BAR = "A7B0B5"

# (old slide position, new page number, detail-link text)
PORTS = [
    (14, 106, "DETAILED REFERENCE · Restored classic · Related live slides 14–15"),
    (15, 107, "DETAILED REFERENCE · Restored classic · Related live slide 15"),
    (29, 108, "DETAILED REFERENCE · Restored classic · Related live slide 14"),
]


def add_box(slide, name, x, y, w, h, text, pt, bold, color, align=PP_ALIGN.LEFT):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.name = name
    sp.shadow.inherit = False
    sp.fill.background()
    sp.line.fill.background()
    tf = sp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Inches(0.0)
    tf.margin_right = Inches(0.0)
    tf.margin_top = Inches(0.0)
    tf.margin_bottom = Inches(0.0)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(pt)
    r.font.bold = bold
    r.font.name = "Calibri"
    r.font.color.rgb = RGBColor.from_string(color)
    return sp


def port_slide(dest, old_slide, layout, page, detail):
    new = dest.slides.add_slide(layout)
    spTree = new.shapes._spTree
    id_map = {}

    for sh in old_slide.shapes:
        el = copy.deepcopy(sh._element)
        spTree.append(el)

    # re-link every copied picture's blip to an image part in the dest package
    for blip in spTree.iter(qn("a:blip")):
        old_rid = blip.get(qn("r:embed"))
        if not old_rid or old_rid in id_map:
            if old_rid:
                blip.set(qn("r:embed"), id_map[old_rid])
            continue
        image_part = old_slide.part.rels[old_rid].target_part
        new_rid = new.part.relate_to(image_part, RT.IMAGE)
        id_map[old_rid] = new_rid
        blip.set(qn("r:embed"), new_rid)

    # swap the old part/page footers for the appendix chrome
    for sh in list(new.shapes):
        if sh.name in ("Text 0", "Text 1") and sh.has_text_frame:
            t = sh.text_frame.text
            if t.strip().isdigit() or t.startswith("PART "):
                sh._element.getparent().remove(sh._element)
    bar = new.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                               Inches(13.333), Inches(0.042))
    bar.name = "FACILITATION_MODE_BAR"
    bar.shadow.inherit = False
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string(REF_BAR)
    bar.line.fill.background()
    add_box(new, "FACILITATION_MODE_LABEL", 10.417, 0.177, 2.552, 0.312,
            "REFERENCE  STUDY", 10.5, True, CHROME, align=PP_ALIGN.RIGHT)
    add_box(new, "detail-link", 0.552, 7.146, 11.25, 0.281, detail, 10.5, False, CHROME)
    add_box(new, "page-number", 12.448, 7.146, 0.521, 0.281, str(page), 10.5, False, GRAY,
            align=PP_ALIGN.RIGHT)

    # pin the 30pt title to Cambria (the old deck's theme carried the serif;
    # the copied runs inherit the DEST theme otherwise)
    for sh in new.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size and abs(r.font.size.pt - 30.0) < 0.1:
                        r.font.name = "Cambria"
    return new


def main():
    dest = Presentation(SRC)
    old = Presentation(OLD)
    assert len(dest.slides._sldIdLst) == 105
    layout = dest.slides[72].slide_layout

    for old_pos, page, detail in PORTS:
        port_slide(dest, old.slides[old_pos - 1], layout, page, detail)

    assert len(dest.slides._sldIdLst) == 108
    dest.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
