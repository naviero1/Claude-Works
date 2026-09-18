#!/usr/bin/env python3
"""v11 — deck split (owner-directed, in session, 2026-09-18).

Owner: "from part 5 on, put it in a different PPTX". The Part 5 divider sits
at v10 slide 47, so:

  Main training deck  = slides 1-46  (Parts 1-4, all exercises)
  Reference deck      = slides 47-113 (Parts 5-6, appendix, self-study)

Printed page numbers are kept as-is in BOTH files so every cross-reference
("Detailed explanation: slide 75" etc.) still points at the right printed
page in the reference deck.

Input: v10.  Outputs: deliverables main + reference decks.
"""
import os

from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "notes", "visual-polish", "pilot", "From_Prompts_to_Agents_Visual_Pilot_v10.pptx")
MAIN = os.path.join(ROOT, "deliverables", "From_Prompts_to_Agents_Facilitated_60_Minute.pptx")
REF = os.path.join(ROOT, "deliverables", "From_Prompts_to_Agents_Reference_Deck.pptx")
SPLIT = 47  # first slide of the reference deck


def cut(keep_pred, out_path, expect):
    prs = Presentation(SRC)
    sldIdLst = prs.slides._sldIdLst
    ids = sldIdLst.findall(qn("p:sldId"))
    assert len(ids) == 113
    for i, el in enumerate(ids, 1):
        if not keep_pred(i):
            rId = el.get(qn("r:id"))
            sldIdLst.remove(el)
            prs.part.drop_rel(rId)
    assert len(prs.slides._sldIdLst) == expect, len(prs.slides._sldIdLst)
    prs.save(out_path)
    print("saved", out_path, expect, "slides")


def main():
    cut(lambda i: i < SPLIT, MAIN, SPLIT - 1)
    cut(lambda i: i >= SPLIT, REF, 113 - SPLIT + 1)


if __name__ == "__main__":
    main()
