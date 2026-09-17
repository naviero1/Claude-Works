#!/usr/bin/env python3
# Correction pass 01 (Beebop review iteration 01) — deck edits.
#  A. Live order: the presentation orientation (slide43.xml) moves to position 41,
#     directly after the dashboard, so the live run matches the five-task arc
#     analyze -> dashboard -> present -> email -> explain. (Item 6)
#  B. Literal on-slide page numbers renumbered for the three moved positions.
#  C. Bridges/notes updated: 40 -> present -> email teach -> email exercise ->
#     explain; the stale "Pick what we drill next" bridge is retired; the email
#     exercise DO now states exactly where the custom prompt runs after the
#     built-in summary (r28). (Item 6)
#  D. Cross-references on appendix pages 82/97/99 renumbered.
#  E. Bravo return rate: 0.349% -> 0.348% (262/75,184 = 0.348478%), notes'
#     full-precision figure corrected, answer-key table cell fixed. (Item 2)
#  F. 14pt teaching text on slides 35 and 44 enlarged to 16pt. (Item 7)
# Text-run edits only; no layout or timing-total changes (REFERENCE 0:10 for the
# moved slide matches the facilitation plan, which is updated in the same commit).
import re
import shutil
import zipfile

from pptx import Presentation
from pptx.util import Pt

PPTX = __file__.rsplit('/', 1)[0] + '/../deliverables/From_Prompts_to_Agents_Facilitated_60_Minute.pptx'

# ---- A. reorder sldIdLst: move position 43 (index 42) to position 41 (index 40)
# (idempotent: skipped when position 41 already resolves to the moved slide)
def _order(path):
    with zipfile.ZipFile(path) as z:
        pres = z.read('ppt/presentation.xml').decode('utf-8')
        rels = z.read('ppt/_rels/presentation.xml.rels').decode('utf-8')
    rid2file = dict(
        (m.group(1), m.group(2)) for m in re.finditer(
            r'Id="([^"]+)"[^>]*Target="slides/(slide\d+\.xml)"', rels))
    rids = re.findall(r'<p:sldId\b[^>]*r:id="([^"]+)"', pres)
    return [rid2file[r] for r in rids]

if _order(PPTX)[40] == 'slide43.xml' or True:
    # After a prior python-pptx save, parts are renamed to match order; detect
    # the reorder by checking whether the email-teach title sits at position 42.
    pass
from pptx import Presentation as _P
_probe = _P(PPTX)
_titles = ['\n'.join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame) for sl in list(_probe.slides)[40:43]]
already = 'Self-study demonstration' in _titles[0]
if already:
    print('A. slide order: already in the corrected sequence, skipping reorder')
else:
    tmp = PPTX + '.tmp'
    with zipfile.ZipFile(PPTX) as zin:
        pres = zin.read('ppt/presentation.xml').decode('utf-8')
        entries = re.findall(r'<p:sldId\b[^>]*/>', pres)
        assert len(entries) == 101, len(entries)
        moved = entries.pop(42)
        entries.insert(40, moved)
        head, sep, tail = pres.partition('<p:sldIdLst>')
        body, sep2, rest = tail.partition('</p:sldIdLst>')
        pres2 = head + sep + ''.join(entries) + sep2 + rest
        with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = pres2.encode('utf-8') if item.filename == 'ppt/presentation.xml' else zin.read(item.filename)
                zout.writestr(item, data)
    shutil.move(tmp, PPTX)
    print('A. slide order: presentation orientation moved to live position 41')

# ---- B–F. text edits via python-pptx, addressed by slide FILE name ------------
p = Presentation(PPTX)
slides = list(p.slides)  # accessing .slides renames parts to positional names
by_file = {sl.part.partname.rsplit('/', 1)[-1]: sl for sl in slides}
assert 'Self-study demonstration' in '\n'.join(
    sh.text_frame.text for sh in slides[40].shapes if sh.has_text_frame)


def para_replace(tf, old, new, expect=1):
    hits = 0
    for para in tf.paragraphs:
        full = ''.join(r.text for r in para.runs)
        if old in full:
            newfull = full.replace(old, new)
            if para.runs:
                para.runs[0].text = newfull
                for r in para.runs[1:]:
                    r.text = ''
            hits += 1
    return hits


def edit(slide_file, where, old, new, expect=1):
    sl = by_file[slide_file]
    hits = 0
    if where in ('body', 'both'):
        for sh in sl.shapes:
            if sh.has_text_frame:
                hits += para_replace(sh.text_frame, old, new)
    if where in ('notes', 'both'):
        if sl.has_notes_slide:
            hits += para_replace(sl.notes_slide.notes_text_frame, old, new)
    assert hits == expect, (slide_file, where, old[:60], hits)


# B. literal page numbers for the three moved positions
edit('slide41.xml', 'body', '43', '41')   # presentation demo (was 43): its only digit run
edit('slide42.xml', 'body', '41', '42')   # email teach (was 41)
edit('slide43.xml', 'body', '42', '43')   # email exercise (was 42)

# C. bridges + the explicit Copilot step
edit('slide40.xml', 'notes', 'Move to “An email summary has a job to do”',
     'Move to “Self-study demonstration: present the findings”')
edit('slide41.xml', 'notes',
     'Optional self-study. Return to the related live slide or continue the reference material as needed.',
     'Move to “An email summary has a job to do” — the numbers on these five slides are the ones the email exercise now defends.')
edit('slide41.xml', 'notes', 'TIME: 0:00 (optional self-study; no live allocation)',
     'TIME: 0:10 (orientation only — the full exercise is reference page 99)')
edit('slide43.xml', 'notes', 'Move to “Pick what we drill next — the method transfers”',
     'Move to “Demonstration: explain it clearly”')
edit('slide43.xml', 'notes',
     'DO: 0:00–0:30 introduce the thread; 0:30–1:30 run the brief;',
     'DO: 0:00–0:30 introduce the thread; 0:30–1:30 run the brief — native path: open the thread in Outlook, click “Summary by Copilot”, then paste the reusable prompt into the Copilot chat pane scoped to that same conversation (that is where the custom prompt runs after the built-in summary; availability depends on your organization’s Copilot license) — fallback that always works: Packaging_Change_Thread.txt pasted with the same prompt into any approved assistant;')
edit('slide43.xml', 'notes', '(paste it, or run it on the thread in Copilot)',
     '(paste the thread into an approved assistant, or — after the built-in summary — run it in Outlook’s Copilot chat pane on the same conversation)')
edit('slide101.xml', 'notes', 'the supplied-text version works in any approved assistant.',
     'the supplied-text version works in any approved assistant; after the built-in summary, the custom prompt runs in the Copilot chat pane scoped to the same conversation.')

# D. appendix cross-references
edit('slide82.xml', 'body', 'Related live slide 41', 'Related live slide 42')
edit('slide99.xml', 'body', 'live orientation: slide 43.', 'live orientation: slide 41.')
edit('slide99.xml', 'notes', 'RELATED: slide 43 (live orientation)', 'RELATED: slide 41 (live orientation)')
edit('slide97.xml', 'notes',
     'RELOCATED from the live sequence (old slide 43) when the presentation demonstration took its slot.',
     'RELOCATED from the live sequence when the presentation demonstration took this content’s slot (that demonstration now sits at live slide 41).')

# E. the corrected rate (262 / 75,184 = 0.348478% -> displayed 0.348%)
edit('slide37.xml', 'body', 'about 0.349%.', 'about 0.348%.')
edit('slide37.xml', 'notes', '= 0.348479%', '= 0.348478%')
edit('slide87.xml', 'notes', '= 0.348479%', '= 0.348478%')
edit('slide92.xml', 'notes', 'Check displayed 0.349% against the underlying ratio.',
     'Check displayed 0.348% against the underlying ratio.')
tbl_hits = 0
for sh in by_file['slide87.xml'].shapes:
    if sh.has_table:
        for row in sh.table.rows:
            for cell in row.cells:
                for para in cell.text_frame.paragraphs:
                    for run in para.runs:
                        if run.text == '0.349%':
                            run.text = '0.348%'
                            tbl_hits += 1
assert tbl_hits == 1, tbl_hits
print('B-E. renumbering, bridges, Copilot step, cross-references, rate fix applied')

# F. enlarge the 14pt teaching text on slides 35 and 44 (files = positions here)
for fn in ('slide35.xml', 'slide44.xml'):
    grown = 0
    for sh in by_file[fn].shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                if run.font.size == Pt(14):
                    run.font.size = Pt(16)
                    grown += 1
    assert grown >= 4, (fn, grown)
    print(f'F. {fn}: {grown} runs 14pt -> 16pt')

p.save(PPTX)
print('saved', PPTX)
