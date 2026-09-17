#!/usr/bin/env python3
"""Restoration wave 3 — owner-directed (in session, 2026-09-17).

Ports the last three of the owner's cited slides verbatim from the preserved
original deck, completing the set:

  109  Four eras: rules -> learning -> generative -> agentic   (old 4)
  110  The hardware story - chips, Nvidia, data centers        (old 6)
  111  Pick by task, not habit - September 2026                (old 17)

Claims spot-checked against notes/research (r1 history: AlexNet 15.3/26.2,
2,300->208B transistors; r22 hardware: x45 vs x45,000,000, Nvidia ~$5.6T
Sep 2026 with ~92% data-center revenue; r13/r24 assistant cards: GDPval,
CJR/Tow citation-error 37% vs 67%, Sept-2026 model lineup) - all carry
September-2026 date stamps and are current this month.

Input: pilot v04.  Output: From_Prompts_to_Agents_Visual_Pilot_v05.pptx
"""
import os

from pptx import Presentation

from build_restorations_wave2 import port_slide

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "notes", "visual-polish", "pilot", "From_Prompts_to_Agents_Visual_Pilot_v04.pptx")
OLD = os.path.join(ROOT, "references", "From_Prompts_to_Agents_Training.pptx")
OUT = os.path.join(ROOT, "notes", "visual-polish", "pilot", "From_Prompts_to_Agents_Visual_Pilot_v05.pptx")

PORTS = [
    (4, 109, "DETAILED REFERENCE · Restored classic · Related live slides 4–5"),
    (6, 110, "DETAILED REFERENCE · Restored classic · Related live slide 6"),
    (17, 111, "DETAILED REFERENCE · Restored classic · Related live slides 13, 16–17"),
]


def main():
    dest = Presentation(SRC)
    old = Presentation(OLD)
    assert len(dest.slides._sldIdLst) == 108
    layout = dest.slides[72].slide_layout

    for old_pos, page, detail in PORTS:
        port_slide(dest, old.slides[old_pos - 1], layout, page, detail)

    assert len(dest.slides._sldIdLst) == 111
    dest.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
