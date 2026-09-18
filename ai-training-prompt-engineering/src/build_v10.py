#!/usr/bin/env python3
"""v10 — owner-directed (in session, 2026-09-18).

Owner: "erase slide 12 it is too redundant" + "adjust all artifacts
accordingly". Slide 12 (Capabilities to look for in an assistant) leaves the
live flow to the self-study tail — it carries approved PROMPT 6/7 (feature
menu), which stays with it rather than being destroyed. Page numbers and every
textual slide reference are renumbered across the deck.

Input: v09.  Output: From_Prompts_to_Agents_Visual_Pilot_v10.pptx (113).
"""
import os
import re

from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "notes", "visual-polish", "pilot", "From_Prompts_to_Agents_Visual_Pilot_v09.pptx")
OUT = os.path.join(ROOT, "notes", "visual-polish", "pilot", "From_Prompts_to_Agents_Visual_Pilot_v10.pptx")


def main():
    prs = Presentation(SRC)
    assert len(prs.slides._sldIdLst) == 113

    NEW_ORDER = list(range(1, 12)) + list(range(13, 114)) + [12]
    assert len(NEW_ORDER) == 113
    sldIdLst = prs.slides._sldIdLst
    ids = sldIdLst.findall(qn("p:sldId"))
    elements = {i + 1: el for i, el in enumerate(ids)}
    for el in ids:
        sldIdLst.remove(el)
    for p in NEW_ORDER:
        sldIdLst.append(elements[p])
    new_pos = {p: i + 1 for i, p in enumerate(NEW_ORDER)}

    # relocation note (owner quote frozen against the remapper below)
    moved = list(prs.slides)[new_pos[12] - 1]
    ntf = moved.notes_slide.notes_text_frame
    ntf.text = ("RELOCATED to self-study per owner direction (2026-09-18): "
                "“erase slide 12 it is too redundant.” Content preserved; "
                "it carries approved PROMPT 6/7 (feature menu).\n\n" + ntf.text)

    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.name == "page-number" and sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text.strip().isdigit():
                            r.text = str(i)

    ref_pat = re.compile(r"\b(slides?)\s+([0-9]{1,3}(?:\s*[–—-]\s*[0-9]{1,3})?"
                         r"(?:\s*,\s*[0-9]{1,3}(?:\s*[–—-]\s*[0-9]{1,3})?)*)\b",
                         re.IGNORECASE)
    num_pat = re.compile(r"[0-9]{1,3}")

    def remap_numbers(numstr):
        def sub(m):
            n = int(m.group(0))
            return str(new_pos.get(n, n)) if 1 <= n <= 113 else m.group(0)
        return num_pat.sub(sub, numstr)

    changed = []

    def process_tf(tf, where):
        for p in tf.paragraphs:
            for r in p.runs:
                if ("v1.10 CONSOLIDATION" in r.text
                        or "RELOCATED to self-study per owner direction" in r.text):
                    continue
                if r.text and "slide" in r.text.lower():
                    newt = ref_pat.sub(lambda m: m.group(1) + " " + remap_numbers(m.group(2)), r.text)
                    if newt != r.text:
                        changed.append((where, r.text.strip()[:50], newt.strip()[:50]))
                        r.text = newt

    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.has_text_frame:
                process_tf(sh.text_frame, f"s{i}:{sh.name}")
        if slide.has_notes_slide:
            process_tf(slide.notes_slide.notes_text_frame, f"s{i}:notes")

    live_fix = re.compile(r"live slide (\d{1,3})")
    first_reloc = new_pos[12]
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.name == "detail-link" and sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        m = live_fix.search(r.text)
                        if m and int(m.group(1)) >= min(first_reloc, 111):
                            r.text = live_fix.sub(r"slide \1 (self-study)", r.text)

    # owner: Gemini tier cells must be bold like the other vendors
    for slide in prs.slides:
        title_hit = any(sh.has_text_frame and "Model tiers and task examples" in sh.text_frame.text
                        for sh in slide.shapes)
        if not title_hit:
            continue
        for sh in slide.shapes:
            if sh.has_table:
                tbl = sh.table
                last = len(tbl.columns) - 1
                fixed = 0
                for ri, row in enumerate(tbl.rows):
                    if ri == 0:
                        continue
                    cell = row.cells[last]
                    for para in cell.text_frame.paragraphs:
                        for run in para.runs:
                            if run.text.strip():
                                run.font.bold = True
                                fixed += 1
                print(f"gemini cells bolded: {fixed} runs")

    # owner: bigger fonts on the four-context-failures slide
    from pptx.util import Pt
    SKIP_NAMES = {"detail-link", "page-number", "FACILITATION_MODE_LABEL",
                  "FACILITATION_MODE_BAR"}
    for slide in prs.slides:
        if not any(sh.has_text_frame and "Four context failures" in sh.text_frame.text
                   for sh in slide.shapes):
            continue
        bumped = 0
        for sh in slide.shapes:
            if not sh.has_text_frame or sh.name in SKIP_NAMES:
                continue
            if sh.text_frame.text.strip() in ("POISONING", "DRIFT", "CONFUSION", "CLASH"):
                continue  # already large; wider would wrap mid-word
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.size and 9 <= run.font.size.pt < 20:
                        run.font.size = Pt(run.font.size.pt + 2)
                        bumped += 1
        print(f"context-failures fonts bumped: {bumped} runs")
        break

    # tiers slide: subtitle must clear the table top edge
    from pptx.util import Inches
    for slide in prs.slides:
        if not any(sh.has_text_frame and "Model tiers and task examples" in sh.text_frame.text
                   for sh in slide.shapes):
            continue
        for sh in slide.shapes:
            if sh.name == "slide-title":
                sh.top = Inches(0.60)
                sh.height = Inches(0.50)
            if sh.name == "slide-subtitle":
                sh.top = Inches(1.14)
                sh.height = Inches(0.34)
                for para in sh.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(12.5)
        print("tiers title/subtitle repositioned")
        break

    print(f"remapped {len(changed)} references")
    for w, a, b in changed:
        print(f"  {w}: '{a}' -> '{b}'")
    assert len(prs.slides._sldIdLst) == 113
    prs.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
