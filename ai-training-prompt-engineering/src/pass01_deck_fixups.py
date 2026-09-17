#!/usr/bin/env python3
# Correction pass 01 — verifier-round fix-ups (three notes-only edits).
#  1. Slide 37 archived source notes: the legacy defect-rate wording added
#     returns to defects; corrected to the separate-measures definition the
#     rest of the package teaches.
#  2/3. Slides 42/43: the archived source notes carry the LONG-COURSE variant
#     thread's answer key (Oct 6 / Sep 22 / 18k on Dan's approval). Label them
#     so a facilitator cannot mistake them for today's thread key (Sep 25 ->
#     Oct 2, Packaging_Change_Thread.txt).
from pptx import Presentation

PPTX = __file__.rsplit('/', 1)[0] + '/../deliverables/From_Prompts_to_Agents_Facilitated_60_Minute.pptx'
p = Presentation(PPTX)
slides = list(p.slides)


def notes_edit(pos, old, new, expect=1):
    tf = slides[pos - 1].notes_slide.notes_text_frame
    hits = 0
    for para in tf.paragraphs:
        full = ''.join(r.text for r in para.runs)
        if old in full:
            para.runs[0].text = full.replace(old, new)
            for r in para.runs[1:]:
                r.text = ''
            hits += 1
    assert hits == expect, (pos, old[:60], hits)


notes_edit(37, 'highest defect rate — (Defects_Found + Units_Returned) ÷ Units_Shipped',
           'highest defect rate — Defects_Found ÷ Units_Shipped; returns are a separate measure, never added')

CAVEAT = (' The archived key below belongs to the long-course variant thread '
          '(Email_Thread_Packaging_Change.txt: Sep 22 → Oct 6) — NOT today’s '
          'Packaging_Change_Thread.txt (Sep 25 → Oct 2, key on this slide and tab KEY-Email).')
for pos in (42, 43):
    notes_edit(pos, 'ARCHIVED SOURCE NOTES', 'ARCHIVED SOURCE NOTES —' + CAVEAT)

p.save(PPTX)
print('deck fix-ups applied (slide 37 formula; slides 42/43 archived-key caveat)')
