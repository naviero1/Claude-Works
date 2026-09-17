#!/usr/bin/env python3
# Round 18 — path pass on the 60-minute deck.
# The package was restructured: deliverables/ now holds only the five main
# artifacts; exercise sources moved to references/exercise-data/ and the
# workbook up to deliverables/Course_Workbook.xlsx. This script updates the
# only three slides that named concrete paths (36, 44, 100). Text-run edits
# only — no layout, timing, or sequence changes.
import sys
from pptx import Presentation

PPTX = __file__.rsplit('/', 1)[0] + '/../deliverables/From_Prompts_to_Agents_Facilitated_60_Minute.pptx'

# (slide number, where, old, new, expected occurrences)
EDITS = [
    (36, 'notes', 'The workbook lives in deliverables/exercise-data/ and regenerates',
                  'The workbook lives in deliverables/ and regenerates', 1),
    (44, 'notes', '(exercise-data/plain-language/)', '(references/exercise-data/plain-language/)', 1),
    (44, 'notes', 'MATERIALS: exercise-data/plain-language/README.md',
                  'MATERIALS: references/exercise-data/plain-language/README.md', 1),
    (100, 'body', 'sample outputs in exercise-data/plain-language/.',
                  'sample outputs in references/exercise-data/plain-language/.', 1),
    (100, 'notes', 'RELATED: slide 44 (live demonstration); exercise-data/plain-language/.',
                   'RELATED: slide 44 (live demonstration); references/exercise-data/plain-language/.', 1),
]

def runs_of(sl, where):
    if where == 'notes':
        frames = [sl.notes_slide.notes_text_frame] if sl.has_notes_slide else []
    else:
        frames = [sh.text_frame for sh in sl.shapes if sh.has_text_frame]
    for tf in frames:
        for para in tf.paragraphs:
            for run in para.runs:
                yield run

p = Presentation(PPTX)
slides = list(p.slides)
total = 0
for n, where, old, new, expect in EDITS:
    hits = 0
    for run in runs_of(slides[n - 1], where):
        if old in run.text:
            run.text = run.text.replace(old, new)
            hits += 1
    assert hits == expect, f'slide {n} {where}: expected {expect} hit(s) for {old!r}, got {hits}'
    total += hits
p.save(PPTX)
print(f'round18 path pass: {total} edits applied, saved {PPTX}')
