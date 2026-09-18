#!/usr/bin/env python3
"""v12 — First-package main deck (owner-directed, in session, 2026-09-18).

Owner: "in the main presentation, add from the reference deck slides 27-42"
— those are the detailed-explanation appendix (v10 printed 73-88). The main
deck becomes slides 1-46 + 73-88 of v10, renumbered 1-62 continuously, with
"Detailed explanation: slide N" references remapped (73-88 -> 47-62). Numbers
that still point at the Reference Deck (47-72, 89-113) are left untouched —
that deck keeps its printed numbering.

Input: v10.  Output: deliverables main deck (62 slides).
"""
import os
import re

from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "notes", "visual-polish", "pilot", "From_Prompts_to_Agents_Visual_Pilot_v10.pptx")
OUT = os.path.join(ROOT, "deliverables", "From_Prompts_to_Agents_Facilitated_60_Minute.pptx")

KEEP = list(range(1, 47)) + list(range(73, 89))   # 46 + 16 = 62


def main():
    prs = Presentation(SRC)
    sldIdLst = prs.slides._sldIdLst
    ids = sldIdLst.findall(qn("p:sldId"))
    assert len(ids) == 113
    elements = {i + 1: el for i, el in enumerate(ids)}
    for i, el in enumerate(ids, 1):
        if i not in KEEP:
            rId = el.get(qn("r:id"))
            sldIdLst.remove(el)
            prs.part.drop_rel(rId)
    assert len(prs.slides._sldIdLst) == 62

    new_pos = {p: i + 1 for i, p in enumerate(KEEP)}   # 73->47 ... 88->62

    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.name == "page-number" and sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text.strip().isdigit():
                            r.text = str(i)

    ref_pat = re.compile(r"\b(slides?)\s+([0-9]{1,3}(?:\s*[–—-]\s*[0-9]{1,3})?"
                         r"(?:\s*,\s*[0-9]{1,3}(?:\s*[–—-]\s*[0-9]{1,3})?)*)\b",
                         re.IGNORECASE)
    num_pat = re.compile(r"[0-9]{1,3}")

    def remap_numbers(numstr):
        def sub(m):
            n = int(m.group(0))
            # only the copied appendix range moves; everything else keeps the
            # Reference Deck's printed numbering
            return str(n - 26) if 73 <= n <= 88 else m.group(0)
        return num_pat.sub(sub, numstr)

    changed = []

    def process_tf(tf, where):
        for p in tf.paragraphs:
            for r in p.runs:
                if ("v1.10 CONSOLIDATION" in r.text
                        or "RELOCATED to self-study per owner direction" in r.text):
                    continue
                if r.text and "slide" in r.text.lower():
                    newt = ref_pat.sub(lambda m: m.group(1) + " " + remap_numbers(m.group(2)), r.text)
                    if newt != r.text:
                        changed.append((where, r.text.strip()[:55], newt.strip()[:55]))
                        r.text = newt

    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.has_text_frame:
                process_tf(sh.text_frame, f"s{i}:{sh.name}")
        if slide.has_notes_slide:
            process_tf(slide.notes_slide.notes_text_frame, f"s{i}:notes")

    # the appendix divider must describe only what this deck contains; the
    # answer keys and full thread stayed in the Reference Deck
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame and "They preserve the detailed prompting" in sh.text_frame.text:
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        if "answer key" in r.text:
                            r.text = ("They preserve the detailed prompting and delegation "
                                      "guidance. The supplier answer key and the complete "
                                      "fictional email thread live in the Reference Deck.")
                            print("appendix divider description corrected")

    print(f"remapped {len(changed)} references")
    for w, a, b in changed:
        print(f"  {w}: '{a}' -> '{b}'")
    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    main()
