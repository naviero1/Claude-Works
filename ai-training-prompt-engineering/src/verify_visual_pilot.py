#!/usr/bin/env python3
"""Preservation verification for the visual pilot, owner-verdict state.

Owner verdict 2026-09-17 (in session): treatments approved on slides 23, 28,
34 (version A), 53; declined on 37 and 87. Version B was superseded by the
owner choosing A (the committed B file remains as comparison evidence only
and is not verified here).

Checks pilot v01 against the confirmed baseline: slide count/identity/order,
exactly the four approved slide parts differing, notes byte-identical,
theme/master/layout byte-identical, verbatim visible text with only approved
labels added, and zero font-size changes. Exits nonzero on any failure.
"""
import json
import os
import sys
import zipfile
from lxml import etree

from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASE = os.path.join(ROOT, "deliverables", "From_Prompts_to_Agents_Facilitated_60_Minute.pptx")
A = os.path.join(ROOT, "notes", "visual-polish", "pilot", "From_Prompts_to_Agents_Visual_Pilot_v01.pptx")
LOG = os.path.join(ROOT, "notes", "visual-polish", "pilot", "font_change_log.json")

SIDS = {23: 278, 28: 283, 34: 291, 37: 294, 53: 308, 87: 342}
APPROVED = [23, 28, 34, 53]
DECLINED = [37, 87]
FAILS = []


def check(ok, msg):
    print(("PASS " if ok else "FAIL ") + msg)
    if not ok:
        FAILS.append(msg)


def canon(xml_bytes):
    return etree.tostring(etree.fromstring(xml_bytes), method="c14n")


def notes_xml(prs, idx):
    s = prs.slides[idx]
    return s.notes_slide.part.blob if s.has_notes_slide else b""


def visible_texts(slide):
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                txt = "".join(r.text for r in p.runs)
                if txt.strip():
                    out.append(txt)
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    if c.text.strip():
                        out.append(c.text)
    return out


def run_sizes(slide):
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.text.strip():
                        out.append((r.text, r.font.size.pt if r.font.size else None,
                                    r.font.bold))
    return out


def main():
    base = Presentation(BASE)
    va = Presentation(A)

    # 1. count / identity / order
    ids = [int(e.get("id")) for e in va.slides._sldIdLst.findall(qn("p:sldId"))]
    base_ids = [int(e.get("id")) for e in base.slides._sldIdLst.findall(qn("p:sldId"))]
    check(len(ids) == 103, "A: 103 slides")
    check(ids == base_ids, "A: native slide id order identical to baseline")
    for pos, sid in SIDS.items():
        check(ids[pos - 1] == sid, f"A: slide {pos} keeps sid-{sid}")

    # 2. exactly the four approved slide parts differ
    diff_a = [i + 1 for i in range(103)
              if canon(base.slides[i].part.blob) != canon(va.slides[i].part.blob)]
    check(sorted(diff_a) == APPROVED,
          f"A: only the four APPROVED slide parts differ ({diff_a})")
    for pos in DECLINED:
        check(pos not in diff_a, f"A: declined slide {pos} is untouched baseline")

    # 3. notes parts identical everywhere
    bad_notes = [i + 1 for i in range(103) if notes_xml(base, i) != notes_xml(va, i)]
    check(not bad_notes, f"A: all notes parts byte-identical to baseline ({bad_notes})")

    # 4. theme / master / layout parts
    with zipfile.ZipFile(BASE) as zb, zipfile.ZipFile(A) as za:
        def infra(z):
            return {n: z.read(n) for n in z.namelist()
                    if ("theme" in n or "slideMaster" in n or "slideLayout" in n)
                    and n.endswith(".xml")}
        ib, ia = infra(zb), infra(za)
        check(set(ib) == set(ia) and all(ib[k] == ia[k] for k in ib),
              "A: theme/master/layout parts byte-identical to baseline")

    # 5. verbatim visible text on the approved slides; only approved labels added
    approved_new = {
        23: {"Audience + source", "method", "format + tone + check"},
        28: set(),
        34: set(),
        53: set(),
    }
    for pos in APPROVED:
        bt = visible_texts(base.slides[pos - 1])
        at = visible_texts(va.slides[pos - 1])
        missing = [t for t in bt if t not in at]
        check(not missing, f"A slide {pos}: every baseline text present verbatim ({missing})")
        extra = [t for t in at if t not in bt and t not in approved_new[pos]]
        check(not extra, f"A slide {pos}: no unapproved new text ({extra})")

    # 6. zero font-size changes in the approved pilot; log must be empty
    with open(LOG) as f:
        log = json.load(f)
    check(log == [], f"font_change_log.json is empty (no applied increases): {log}")
    for pos in APPROVED:
        base_rs = {t: (s, b) for t, s, b in run_sizes(base.slides[pos - 1])}
        for text, size, bold in run_sizes(va.slides[pos - 1]):
            if text in base_rs:
                bs, bb = base_rs[text]
                if bs is not None and size is not None:
                    check(abs(size - bs) < 0.01,
                          f"A slide {pos}: size preserved for '{text[:40]}' ({bs}->{size})")
                if bb is not None and bold is not None:
                    check(bb == bold, f"A slide {pos}: emphasis preserved for '{text[:40]}'")

    print()
    if FAILS:
        print(f"RESULT: {len(FAILS)} FAILURE(S)")
        sys.exit(1)
    print("RESULT: ALL CHECKS PASSED")


if __name__ == "__main__":
    main()
