from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parent.parent
D=Document(); sec=D.sections[0]; sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 s=D.styles[n];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
 s.font.size=Pt(11 if n=='Normal' else 22 if n=='Title' else 16 if n=='Heading 1' else 12)
 s.paragraph_format.space_after=Pt(7)
D.styles['Normal'].paragraph_format.line_spacing=1.05
sec.header.paragraphs[0].text='668XXX  •  Tissue Supplier Personnel Training and Certification'
sec.header.paragraphs[0].style='Caption'
f=sec.footer.paragraphs[0];f.text='Review draft v2  •  17 September 2026                                         Page '
a=OxmlElement('w:fldSimple');a.set(qn('w:instr'),'PAGE');f._p.append(a)
md=[]
def h(t,l=1): D.add_heading(t,l);md.append('#'*(l+1)+' '+t+'\n')
def p(t): D.add_paragraph(t);md.append(t+'\n')
def page(t): D.add_page_break();h(t)
def table(headers,rows,widths=None):
 t=D.add_table(rows=1,cols=len(headers));t.autofit=False
 if widths:
  for col,w in zip(t.columns,widths):col.width=Inches(w)
 for i,x in enumerate(headers):t.rows[0].cells[i].text=x
 for r in rows:
  for c,x in zip(t.add_row().cells,r):c.text=x
 for j,r in enumerate(t.rows):
  pr=r._tr.get_or_add_trPr();el=OxmlElement('w:cantSplit');pr.append(el)
  if j==0:pr.append(OxmlElement('w:tblHeader'))
  for i,c in enumerate(r.cells):
   if widths:c.width=Inches(widths[i])
   cp=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE9F5' if j==0 else ('F5F7FA' if j%2 else 'FFFFFF'));cp.append(shade)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   cp.append(borders)
   mar=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'80');el.set(qn('w:type'),'dxa');mar.append(el)
   cp.append(mar)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(3)
    for rr in pp.runs:rr.font.size=Pt(10);rr.bold=j==0
 md.append('| '+' | '.join(headers)+' |\n|'+'---|'*len(headers))
 md.extend('| '+' | '.join(r)+' |' for r in rows);md.append('')
D.add_paragraph('Tissue Supplier Personnel Training and Certification',style='Title')
p('Advanced Tissue Models • 668XXX • Standard Operating Procedure • Review draft v2')
p('Owner: Supplier Engineering. Approver: Quality. Technical standard owner: Design Engineer. Effective date and controlled revision: assigned upon approval.')
h('1 Purpose and scope')
p('This procedure defines training, monitoring and certification for supplier personnel performing tissue harvest, processing, packaging or build activities. It applies to every tissue family listed in Appendix A. Personnel certification precedes tissue component qualification and does not replace supplier approval or product acceptance.')
p('Entry triggers are a new supplier providing new or existing tissue, an existing supplier providing new tissue, new or reassigned personnel, and retraining caused by changes, technique-related nonconformance or recurring low yield. Non-tissue parts follow the separate Component Qualification procedure. Internal manufacturing training remains under 668729.')
h('2 Responsibilities')
table(['Role','Responsibility'],[
('Supplier Engineering','Own procedure and training log; evaluate individual evidence, certify scope, review trends and coordinate qualification.'),
('Quality','Approve procedure and deviations; certify inspectors with the Design Engineer; govern containment and return to routine work.'),
('Design Engineer','Own the acceptance standard and process instruction content; resolve borderline criteria; approve defect examples and reference judgments.'),
('Tissue expert and site inspector','Expert delivers technique training. Inspector checks output and records attribution. Both sign readiness for Monitoring.'),
('Operations Supervisor and supplier site lead','Support training and staffing; enforce personnel status and controls; notify Supplier Engineering before reassignment.'),
('Training Department and trainee','Training Department retains training records under 668005, subject to release confirmation. Trainee follows the instruction and participates in monitoring.')],[2.0,5.0])
page('3 References and definitions')
p('All existing procedures remain separate. Use their current controlled revisions. The identifiers 668XXX are placeholders requiring document-control assignment.')
table(['Identifier','Document or role in the system'],[
('668008','Supplier Controls, Advanced Tissue Models'),('668009','Supplier Selection, Approval and Monitoring, Advanced Tissue Models'),('668504 / 668529','Supplier Auditing / Approved Supplier List Use and Maintenance'),('668007','Purchasing and Planning'),('668005 / 668729','Training / Manufacturing Training and Certification'),('800082','Training Record'),('668XXX','Component Qualification and its checklist, separate from this procedure'),('Appendix A','Approved tissue matrix, process instruction, specification and named personnel'),('To confirm before release','Nonconformance, corrective action and document-change-control procedure identifiers')],[1.45,5.55])
h('Definitions',2)
p('SH (Supplier Site) means any tissue supplier location, including a slaughterhouse or processor. IQC (Incoming Quality Control) at the supplier site inspects tissue when harvested or built. ATM (Advanced Tissue Models) incoming inspection independently determines acceptance before assembly. Full role names are used elsewhere for clarity.')
p('A tissue family groups parts sharing an approved process instruction and acceptance specification. Certification applies to a named person, supplier site, task, tissue family and applicable document revisions; it is not blanket certification for all tissue tasks.')
p('First-pass yield equals the number of units conforming at their first inspection divided by all units first inspected in the defined window, multiplied by 100. Rework cannot convert a first-pass failure into a first-pass success. Record coached or demonstration units separately.')
p('The Training Phase permits demonstration and direct coaching. The Monitoring Phase evaluates task execution with every unit inspected. There is no standing Mentor or Certified Supplier Trainer role. Quality and the Design Engineer designate the tissue expert for each site and family.')
page('4 Prerequisites and training')
h('4.1 Readiness to begin',2)
p('Supplier Engineering verifies supplier approval under 668008 and 668009 and the Approved Supplier List record under 668529. Before collecting qualification evidence, complete and approve the applicable Appendix A settings, including the acceptance specification, task scope, inspector status and attribution method.')
p('The Design Engineer communicates the acceptance standard to the tissue expert and inspector. They use the controlled instruction, specification and approved defect examples to teach personnel. They must not independently redefine acceptance. Escalate ambiguous criteria to the Design Engineer and record the decision; recurring ambiguity requires a controlled document update.')
p('Set up per-unit attribution before Monitoring. Use a unit identifier, traceable batch log or one-trainee-at-a-time batch method that links every first inspection result to a named person and task. Mixed output without reliable attribution cannot establish individual certification.')
h('4.2 Training Phase',2)
p('The designated tissue expert demonstrates the task and coaches the trainee. The Operations Supervisor assists. Training covers the relevant process steps, acceptable and unacceptable tissue features, defect identification, handling, packaging, labeling, cold-chain handoff and records. Use language and materials the trainee understands.')
p('Record the demonstration count required in Appendix A and each completed demonstration. Include the actual part family and task rather than limiting the procedure to female pelvic harvesting. Direct support and rework guidance are permitted during training and are recorded separately from certification evidence.')
p('The tissue expert and inspector jointly confirm and sign that the trainee can perform the task to the controlled instruction before entering Monitoring. Record the date, person, task, site, tissue family and document revisions. If not ready, continue training.')
h('4.3 Authorization boundaries',2)
p('Personnel perform routine supply work only within their certified scope or the defined Monitoring controls. Instructional work remains supervised and segregated from qualifying evidence. Any work outside these conditions requires the written, time-limited Quality-approved deviation described in Section 7.3. Training status alone never authorizes shipment.')
page('5 Monitoring and personnel certification')
h('5.1 Monitoring execution',2)
p('The trainee performs the task and the qualified supplier-site inspector inspects every unit, records its first result, defect classification, suspected cause and individual attribution. Record observed facts separately from inferred causes. Supplier Engineering reviews the complete sequence to distinguish skill, material, process and inspection issues.')
p('Questions to trainers or certified peers are permitted without penalty. If a unit requires direct coaching or hands-on correction, identify it as coached, retain its result and return to training as appropriate. Coached units do not count as unassisted certification evidence. After coached retraining, open a fresh Monitoring window; retain prior results.')
h('5.2 Certification gate',2)
p('For each person, task and tissue family at a site, certification requires at least 95% first-pass yield across at least 30 inspected units over at least two consecutive harvest days. Harvest days are successive days when the in-scope harvest task is performed; non-harvest calendar days do not supply evidence. Record the evaluation window before evaluating the result and retain all units within it; do not select only favorable days or units.')
p('Use the actual inspected denominator and compare the unrounded yield with 95%. With exactly 30 inspected units, 29 passes meet the gate (96.67%); 28 passes do not (93.33%). The minimum is 30 units in total across the window, not 30 per day. This is a personnel screening criterion, not statistical proof of long-term process capability.')
h('5.3 Failure and certification decision',2)
p('A trainee who has not met every gate remains under Monitoring controls or returns to training. The inspector flags every result below 95% for Supplier Engineering and Quality review. Recurring results below 95%, using the approved recurrence count and window, require the inspector to shadow the harvesting technician, reinforce training and inspect the output. Apply Section 7 for containment and recovery.')
p('Supplier Engineering evaluates and signs the evidence, records certification on FORM 800082 and updates the Supplier Training Log with scope and status. No family may use a lower yield threshold merely because 95% has not yet been achieved. Correct the technique or process and reassess under the same threshold.')
page('6 Inspector certification and component qualification')
h('6.1 Inspector competency',2)
p('Advanced Tissue Models Quality and the Design Engineer train and certify supplier-site inspection personnel. Harvest yield does not measure inspection accuracy. An uncertified inspector cannot provide the sole qualifying judgments used to certify harvesting personnel.')
p('Use a Quality-approved attribute-agreement assessment against Design Engineer-confirmed reference judgments. For perishable tissue, the proposed method is concurrent independent inspection of the same live units, supported by an approved image set for defect classification. Preserve blinding until each inspector records a judgment. Include passes, fails and borderline cases.')
p('Before use, Quality and the Design Engineer approve the reference-set size and mix, agreement and defect-classification thresholds, false-accept and false-reject limits, critical-nonconformance rule, repeatability requirement and reassessment triggers. The original proposed 90% agreement threshold is not approved by this draft. Record competency in explaining defects and escalating uncertain causes.')
h('6.2 Separate tissue component gate',2)
p('After personnel certification, execute the approved tissue component qualification plan at the supplier site. Every required qualification lot must achieve at least 95% first-pass yield. Supplier Engineering and Quality must approve the number of lots, lot definition and inspection coverage before execution. The 30-unit personnel minimum does not establish the lot count or lot inspection sample.')
p('Link lot results to certified personnel, site, part, process and specification revisions. Complete the tissue section of the Component Qualification Checklist; obtain Quality approval before recording the qualified supplier-site and part scope under 668529 and handing off to 668007 Purchasing and Planning. A failed lot requires containment, investigation and an approved repeat plan; it cannot be hidden by averaging it with passing lots.')
h('6.3 Independent product acceptance',2)
p('Only conforming units may move downstream under the applicable disposition controls. A 95% yield does not permit release of the failed 5%. Supplier-site inspections support training and qualification; every tissue unit remains subject to the independent incoming/pre-assembly inspection required at Advanced Tissue Models, regardless of personnel certification.')
p('Compare supplier-site and receipt results. A receipt-only deterioration requires investigation of freezing, handling, transport and inspection consistency. Do not attribute it to harvesting personnel or order retraining solely from the yield difference. New personnel alone require individual certification; assess whether any accompanying process change requires part requalification.')
page('7 Ongoing control and records')
h('7.1 Triggers and review',2)
p('Supplier Engineering reviews yield and defect trends by person, task, family and site and shares findings with Quality. Monthly supplier quality and supply-chain calls track effectiveness, staffing and actions. Do not delay containment or low-yield review until the monthly call.')
p('New or reassigned personnel enter training. Specification or process-instruction changes require documented training on the change; Quality and the Design Engineer determine the need for Monitoring and component requalification. Technique-related nonconformance or corrective action requires retraining and Monitoring. Apply the approved inactivity interval to reassess lapsed competency.')
h('7.2 Recurring yield below 95 percent',2)
p('Segregate nonconforming output and review affected lots under the existing disposition process. When the approved recurrence rule is met, the supplier-site incoming quality control inspector shadows the harvesting technician, reinforces the approved technique and inspects output. Record the observations, coached units, defects and suspected causes. Investigate material, method and inspection causes alongside personnel skill.')
p('After corrective training, obtain a fresh unassisted Monitoring window meeting 95% yield, 30 inspected units and two consecutive harvest days. Supplier Engineering reviews recovery evidence and Quality authorizes return to routine work, with the status and authorization recorded. If the gate is not met, maintain controls, repeat training or correct the process and assess component requalification.')
h('7.3 Written exception controls',2)
p('Before work outside certification or defined Monitoring controls, Quality must approve a documented deviation naming the people, site, tissue and tasks, reason, risk controls, supervision, inspection coverage, quantity or lot limits, start and expiration dates, responsible owner and exit criteria. Supplier Engineering records and tracks the authorization. No verbal or expired deviation permits continued work. Renewals require new approval before expiry.')
p('A deviation does not itself certify a person, qualify a component, lower the 95% criterion or release nonconforming tissue. Product disposition and shipment approval remain separate. When no valid personnel route or deviation exists, stop the affected in-scope supply work and escalate to Quality.')
h('7.4 Records',2)
p('Retain FORM 800082 and linked evidence under 668005. Supplier Engineering owns the status log; confirm the Training Department filing responsibility and approved repository before release. Keep unit attribution, first results, yields, training language, readiness signatures, inspector assessments, deviations, revisions and recovery decisions. Link component evidence to its checklist and retain records for audits under 668504. The repository and retention period follow controlled requirements, not an assumed software platform.')
page('8 Training process map')
p('Roles and decisions are shown in the sequence below. The complete text requirements remain controlling. Blue identifies prerequisites; green identifies the training procedure.')
D.add_picture(str(R/'build/procedure-map-1.png'),width=Inches(5.9))
md.append('![Training process map](../source/procedure-map-1.png)\n')
page('9 Low yield response process map')
p('The shadow-training path does not replace containment or independent product acceptance. Recovery requires documented evidence and authorization.')
D.add_picture(str(R/'build/procedure-map-2.png'),width=Inches(5.9))
md.append('![Low yield response process map](../source/procedure-map-2.png)\n')
page('Appendix A Tissue family approval record')
p('Complete one controlled record per tissue family and site. Supplier Engineering coordinates completion; Quality approves the plan with the Design Engineer for technical criteria. Blank required fields prevent execution of the affected qualification activity.')
table(['Field','Required entry or fixed requirement'],[
('Family, site and tasks','____________________________'),('Parts and revisions','____________________________'),('Controlled process instruction','Document and revision: ____________________'),('Acceptance standard and defect catalog','Document and revision; Design Engineer approval: ____________________'),('Designated expert and inspector','Names, competency record, designation date: ____________________'),('Training language and demonstrations','Language, materials and minimum demonstrations: ____________________'),('Attribution method','Approved method demonstrated before Monitoring: ____________________'),('Personnel gate','At least 95% yield; at least 30 inspected units; at least two consecutive harvest days'),('Monitoring window and attempt limit','Window definition and escalation limit: ____________________'),('Recurring low yield','Count of results below 95% and review window: ____________________'),('Qualification lots','Number, definition and inspection coverage: ____________________\nEvery required lot at least 95% first-pass yield'),('Inspector assessment','Reference set, agreement criteria, error limits and repeatability: ____________________'),('Inactivity and receipt gap','Approved reassessment interval and investigation trigger: ____________________'),('Approvals','Supplier Engineering / Quality / Design Engineer with dates: ____________________')],[2.15,4.85])
p('Pelvic block starting references requiring verification: process instruction 1176680; candidate parts 666541, 666540, 666518 and 666506. Confirm whether 668741 is the controlling acceptance specification and confirm its current revision. The original draft labeled 666541 and 666506 as specifications/work instructions; do not treat those identifiers as validated specification references.')
page('Appendix B Evidence and release checklist')
h('Monitoring evidence record',2)
p('Header: supplier site, tissue family, part and revision, trainee, task, trainer, inspector, instruction/specification revisions, window start/end and readiness signatures.')
table(['Per-unit field','Required information'],[('Traceability','Unit identifier, lot, harvest date and person/task attribution'),('First inspection','First pass or fail, defect class, inspector and timestamp'),('Training support','Unassisted, question only or coached; coaching details'),('Disposition','Segregation, rework or disposition reference; original result retained')],[1.8,5.2])
p('Window summary: inspected count, first-pass count, unrounded yield, consecutive harvest days, excluded coached units with reasons, gate decision, Supplier Engineering signature/date and FORM 800082 reference. Retain failed attempts and link subsequent recovery windows.')
h('Controlled release checklist',2)
table(['Item','Approval needed'],[
('Assigned document identifiers and revisions','Document control; maintain 668008 and 668009 separately'),('Tissue matrix and acceptance standard','Quality and Design Engineer'),('Qualification-lot count and inspection plan','Supplier Engineering and Quality'),('Recurrence window, attempt and inactivity limits','Supplier Engineering and Quality'),('Inspector criteria and reference set','Quality and Design Engineer'),('Attribution method and implementation','Supplier Engineering and site inspector'),('Repository, retention and filing responsibility','Training Department and Quality'),('Corporate applicability and change triggers','Component Qualification owner with Quality')],[3.4,3.6])
p('The three green notes cropped from the original proposed map remain unresolved source material. Their missing text must be recovered before it is relied on for qualification or change-trigger requirements. This review draft does not retire any existing document or establish a controlled effective date.')
for st in D.styles:
 if st.type == 1:
  st.font.color.rgb=RGBColor(0,0,0)
  for e in list(st.element.xpath('.//w:pBdr')):e.getparent().remove(e)
for el in list(D.element.xpath('.//w:pBdr')):el.getparent().remove(el)
for para in sec.header.paragraphs:
 for r in para.runs:r.font.color.rgb=RGBColor(0,0,0)
D.core_properties.title='Tissue Supplier Personnel Training and Certification'
D.core_properties.author='Supplier Engineering'
D.save(R/'deliverables/ATM_Tissue_Supplier_Training_Procedure_v2.docx')
(R/'source').mkdir(exist_ok=True)
(R/'source/ATM_Tissue_Supplier_Training_Procedure_v2.md').write_text('# Tissue Supplier Personnel Training and Certification\n\n'+'\n'.join(md))
print('Created procedure')
