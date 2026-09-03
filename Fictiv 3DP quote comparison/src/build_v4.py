import datetime
from copy import copy
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as gcl
from q5_data import Q5, STATED, PROCESS, QUOTE_FEEDBACK

SRC = '../Fictiv_Quote_Comparison_V3.xlsx'
OUT = '../Fictiv_Quote_Comparison_V4.xlsx'

# ---- style palette lifted from the existing workbook ----------------------
F   = 'Arial'
NAVY, MIDBLUE, GREEN_H, GREEN_D, GREY_H = 'FF1F4E78','FF2E75B6','FFC6E0B4','FFE2EFDA','FFEDEDED'
YELLOW, SECT = 'FFFFF2CC','FFD9E1F2'
ORNG_H, ORNG_D = 'FFF8CBAD','FFFCE4D6'          # new: Quote 5 (current quote)
INPUT_F = 'FFFFFF00'                            # editable input cells
CUR  = '\\$#,##0.00'
CURS = '\\$#,##0.00;[RED]"($"#,##0.00\\)'
PCT  = '0.0%;[RED]\\-0.0%'
PCT2 = '0.00%'
CNT  = '0" of 15"'
THIN = Border(*[Side(style='thin')]*4)

def sty(c, size=10, b=False, i=False, color=None, fill=None, fmt=None,
        h=None, v=None, wrap=False, border=True):
    c.font = Font(name=F, sz=size, b=b, i=i, color=color)
    if fill: c.fill = PatternFill('solid', start_color=fill)
    if fmt:  c.number_format = fmt
    if h or v or wrap: c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap)
    if border: c.border = THIN

def put(ws, coord, val, **kw):
    c = ws[coord]; c.value = val; sty(c, **kw); return c

def title(ws, coord, text, last_col):
    r = int(''.join(ch for ch in coord if ch.isdigit()))
    ws.merge_cells(f'A{r}:{last_col}{r}')
    put(ws, coord, text, size=14, b=True, border=False)
    ws.row_dimensions[r].height = 17.35

def subtitle(ws, r, text, last_col):
    ws.merge_cells(f'A{r}:{last_col}{r}')
    put(ws, f'A{r}', text, size=9, i=True, color='FF595959', border=False, h='left', v='top', wrap=True)

def section(ws, r, text, last_col):
    ws.merge_cells(f'A{r}:{last_col}{r}')
    put(ws, f'A{r}', text, size=10, b=True, fill=SECT, border=False)
    ws.row_dimensions[r].height = 15

def hdr(ws, coord, text, fill=NAVY, color='FFFFFFFF'):
    put(ws, coord, text, size=11, b=True, color=color, fill=fill, h='center', v='center')

def note(ws, r, text, last_col):
    ws.merge_cells(f'A{r}:{last_col}{r}')
    put(ws, f'A{r}', text, size=9, i=True, color='FF595959', border=False)
    ws.row_dimensions[r].height = 15

wb = openpyxl.load_workbook(SRC)

# ===========================================================================
# 1. SOURCES — insert Quote 5 at row 12, push the Xometry reference to row 13
# ===========================================================================
ws = wb['Sources']
for col in 'ABCD':                                    # move row 12 -> row 13
    old, new = ws[f'{col}12'], ws[f'{col}13']
    new.value = old.value
    new._style = copy(old._style)
ws.row_dimensions[13].height = ws.row_dimensions[12].height

ws['A3'] = ('Five quotes in the Fictiv series are compared head-to-head. All extended subtotals are '
            'normalized to Quote 2 quantities so the quotes are directly comparable. Quote 5 (Sep 3, 2026) '
            'is quoted in two volume tiers; its Tier 1 unit price (qty 2 per part, 26 business days) is the '
            'one carried into the comparison, to stay like-for-like with the Q2 quantity basis — the Tier 2 '
            'volume pricing and the landed-cost build-up live on the "Quote 5 Detail" tab. Xometry pricing '
            'appears at the right-most columns as a reference benchmark only.')
ws['A12'] = 'Quote 5 — MJF China production (current) ◆'
ws['B12'] = 'Images provided by user: 3-page quote, 15 line items + quote-level feedback'
ws['C12'] = 'September 3, 2026'
ws['D12'] = ('Two-tier 3DP quote. ' + PROCESS + '. Tier 1 = qty 2 per part @ 26 business days; '
             'Tier 2 = qty 30 per part @ 45 business days. China production, DAP shipping, 7.5% sales tax '
             'on parts. Quote-level feedback: ' + '; '.join(QUOTE_FEEDBACK) + '. NOTE: no vendor name appears '
             'on the pages provided — confirm it is Fictiv before citing it as such.')

# ===========================================================================
# 2. COMPARISON — rebuilt as five quotes
#    A #  B PN  C Desc  D Qty | E-I unit Q1..Q5 | J-N ext Q1..Q5
#    | O-R deltas | S-T Xometry
# ===========================================================================
old = wb['Four-Quote Comparison']
prev = {c.coordinate: c.value for row in old.iter_rows() for c in row if c.value is not None}
idx  = wb.sheetnames.index('Four-Quote Comparison')
wb.remove(old)
cmp_name = 'Five-Quote Comparison'
ws = wb.create_sheet(cmp_name, idx)
LAST = 'T'

for col, w in [('A',4),('B',11),('C',24),('D',6),('E',12),('F',11),('G',11),('H',14),('I',13),
               ('J',11),('K',11),('L',11),('M',12),('N',12),('O',10),('P',12),('Q',13),('R',12),
               ('S',11),('T',11)]:
    ws.column_dimensions[col].width = w

title(ws, 'A1', 'Five-Quote Comparison — Q1 (AI orig) | Q2 (AI re-quote) | Q3 (MP w/ error) | '
                'Q4 (MP corrected) ★ | Q5 (MJF China, Sep 3) ◆', LAST)
subtitle(ws, 2, 'All extended subtotals normalized to Quote 2 quantities. Q5 uses its Tier 1 unit price '
                '(qty 2 / 26 business days) for like-for-like comparison. Xometry shown at right as reference only.', LAST)
ws.row_dimensions[2].height = 15

# group header row
ws.row_dimensions[4].height = 21.75
for rng, txt, fill, col in [('A4:C4','Part',NAVY,'FFFFFFFF'), ('D4:D4','Qty',NAVY,'FFFFFFFF'),
                            ('E4:I4','Unit Price by Quote ($)',MIDBLUE,'FFFFFFFF'),
                            ('J4:N4','Extended Subtotal @ Q2 Qty ($)',MIDBLUE,'FFFFFFFF'),
                            ('O4:R4','Quote-to-Quote Δ (%)',MIDBLUE,'FFFFFFFF'),
                            ('S4:T4','Reference: Xometry',GREY_H,'FF000000')]:
    a, b = rng.split(':')
    if a != b: ws.merge_cells(rng)
    hdr(ws, a, txt, fill, col)
    for c in ws[rng][0]: sty(c, size=11, b=True, color=col, fill=fill, h='center', v='center')
    ws[a].value = txt

# column header row
ws.row_dimensions[5].height = 34.5
H5 = {'A':'#','B':'Part No.','C':'Description','D':'Qty','E':'Q1 (AI orig)','F':'Q2 (Apr 23)',
      'G':'Q3 (May 1)','H':'Q4 (May 11) ★','I':'Q5 (Sep 3) ◆','J':'Q1 Ext','K':'Q2 Ext',
      'L':'Q3 Ext','M':'Q4 Ext ★','N':'Q5 Ext ◆','O':'Δ Q2→Q3','P':'Δ Q3→Q4 ★',
      'Q':'Δ Q2→Q4 (net)','R':'Δ Q4→Q5 ◆','S':'Xom. Unit','T':'Xom. Ext'}
for col, txt in H5.items():
    if col in ('H','M','P'):   hdr(ws, f'{col}5', txt, GREEN_H, 'FF000000')
    elif col in ('I','N','R'): hdr(ws, f'{col}5', txt, ORNG_H,  'FF000000')
    elif col in ('S','T'):     hdr(ws, f'{col}5', txt, GREY_H,  'FF000000')
    else:                      hdr(ws, f'{col}5', txt)

# body: preserve every Q1-Q4 / Xometry number from V3, keyed by part number
byrow = {prev[f'B{r}']: r for r in range(6, 21)}
q5row = {r[0]: 19 + i for i, r in enumerate(Q5)}   # part no. -> its row on the detail tab
det   = 'Quote 5 Detail'
for n, r in enumerate(range(6, 21)):
    pn, o = prev[f'B{r}'], r                          # same row numbers as V3
    ws.row_dimensions[r].height = 15
    put(ws, f'A{r}', prev[f'A{o}'])
    put(ws, f'B{r}', pn)
    put(ws, f'C{r}', prev[f'C{o}'])
    put(ws, f'D{r}', prev[f'D{o}'], h='center', v='center')
    for col, src in [('E','E'),('F','F'),('G','G')]:
        put(ws, f'{col}{r}', prev[f'{src}{o}'], fmt=CUR)
    put(ws, f'H{r}', prev[f'H{o}'], fmt=CUR, fill=GREEN_H)
    # Q5 unit price links to the detail tab: one place to edit
    put(ws, f'I{r}', f"='{det}'!E{q5row[pn]}", fmt=CUR, fill=ORNG_D)
    for col, up in [('J','E'),('K','F'),('L','G')]:
        put(ws, f'{col}{r}', f'=$D{r}*{up}{r}', fmt=CUR)
    put(ws, f'M{r}', f'=$D{r}*H{r}', fmt=CUR, fill=GREEN_H)
    put(ws, f'N{r}', f'=$D{r}*I{r}', fmt=CUR, fill=ORNG_D)
    put(ws, f'O{r}', f'=(G{r}-F{r})/F{r}', fmt=PCT)
    put(ws, f'P{r}', f'=(H{r}-G{r})/G{r}', fmt=PCT, fill=GREEN_D)
    put(ws, f'Q{r}', f'=(H{r}-F{r})/F{r}', fmt=PCT)
    put(ws, f'R{r}', f'=(I{r}-H{r})/H{r}', fmt=PCT, fill=ORNG_D)
    put(ws, f'S{r}', prev[f'S{o}'] if f'S{o}' in prev else prev[f'P{o}'], fmt=CUR, fill=GREY_H)
    put(ws, f'T{r}', f'=$D{r}*S{r}', fmt=CUR, fill=GREY_H)

# total row
ws.row_dimensions[21].height = 15
for col in 'AB': put(ws, f'{col}21', None, fill=YELLOW, b=True)
put(ws, 'C21', 'TOTAL (parts only, @ Q2 qty)', b=True, fill=YELLOW)
put(ws, 'D21', None, b=True, fill=YELLOW)
for col in 'JKL': put(ws, f'{col}21', f'=SUM({col}6:{col}20)', b=True, fill=YELLOW, fmt=CUR)
put(ws, 'M21', '=SUM(M6:M20)', b=True, fill=GREEN_H, fmt=CUR)
put(ws, 'N21', '=SUM(N6:N20)', b=True, fill=ORNG_H, fmt=CUR)
put(ws, 'O21', '=(L21-K21)/K21', b=True, fill=YELLOW,  fmt=PCT)
put(ws, 'P21', '=(M21-L21)/L21', b=True, fill=GREEN_D, fmt=PCT)
put(ws, 'Q21', '=(M21-K21)/K21', b=True, fill=GREEN_H, fmt=PCT)
put(ws, 'R21', '=(N21-M21)/M21', b=True, fill=ORNG_H,  fmt=PCT)
put(ws, 'S21', None, b=True, fill=GREY_H)
put(ws, 'T21', '=SUM(T6:T20)', b=True, fill=GREY_H, fmt=CUR)

for i, txt in enumerate([
    '★ = Q4 (corrected MP quote).  ◆ = Q5, the current quote (Sep 3, 2026).',
    '• Δ Q2→Q3: the +10.34% systematic markup that Fictiv later acknowledged as a mistake.',
    '• Δ Q3→Q4: the correction step — Q4 rolls back 12 parts to Q2 prices; keeps 3 parts at Q3 discount.',
    '• Δ Q2→Q4: the net MP value-add after correction. Equals exactly the genuine MP-discount savings on 670832, 670833, 670845.',
    '• Δ Q4→Q5: the move in the current quote. Q5 is a two-tier quote; this column compares its Tier 1 (qty 2 / 26-day) unit price against Q4.',
    '• Q5 unit prices are pulled from the "Quote 5 Detail" tab (cells E19:E33) so there is a single place to correct them.',
    '• Q4 PDF lists Head Cap (670836) at qty 4; normalized to qty 2 here to match Q2 basis (established convention).',
    '• Neither the Fictiv quotes nor Q5 include FAI / inspection — Q5 states explicitly that 3DP parts exclude the inspection report.',
    '• CAVEAT: Q5 is explicitly MJF / Nylon 12 / vapor-smoothed. The V3 workbook never recorded the process for Q1–Q4, so confirm they were quoted to the same finish spec before treating Δ Q4→Q5 as a pure price move.',
]):
    note(ws, 23 + i, txt, LAST)

# ===========================================================================
# 3. QUOTE 5 DETAIL — as-quoted line items, both tiers, three audit blocks
# ===========================================================================
d = wb.create_sheet(det, idx + 1)
DLAST = 'L'
for col, w in [('A',4),('B',11),('C',24),('D',46),('E',12),('F',13),('G',12),('H',14),
               ('I',10),('J',15),('K',20),('L',44)]:
    d.column_dimensions[col].width = w

title(d, 'A1', 'Quote 5 — As-Quoted Detail (quote date: September 3, 2026)', DLAST)
subtitle(d, 2, 'Transcribed from the 3-page quote images supplied by the user on Sep 3, 2026. Both stated '
               'totals reconcile exactly from these line items (see the Tier structure check column), which is '
               'the transcription\'s proof of accuracy. Yellow cells are the as-quoted inputs — edit those; '
               'every other cell is a formula.', DLAST)
d.row_dimensions[2].height = 40

section(d, 4, 'Quote inputs & assumptions  —  edit the yellow cells only', DLAST)
put(d, 'A5', 'Quote date');                     put(d, 'B5', datetime.date(2026,9,3), fill=INPUT_F, fmt='mmmm d, yyyy')
put(d, 'A6', 'Process / finish (as quoted)');   d.merge_cells('B6:L6'); put(d, 'B6', PROCESS, fill=INPUT_F)
put(d, 'A7', 'Sales tax on parts');             put(d, 'B7', STATED['tax_pct'], fill=INPUT_F, fmt='0.0%')
put(d, 'A8', 'Back-solve divisor (assumed margin gross-up)'); put(d, 'B8', 0.85, fill=INPUT_F, fmt='0.00')
put(d, 'A9', 'Quote-level feedback');           d.merge_cells('B9:L9'); put(d, 'B9', '; '.join(QUOTE_FEEDBACK), fill=INPUT_F)
for r in (5,6,7,8,9):
    for col in 'CDEFGHIJKL':
        if r not in (6,9): sty(d[f'{col}{r}'])
put(d, 'C5', 'Vendor name does not appear on the pages provided — confirm before citing as Fictiv.',
    size=9, i=True, color='FF595959')
put(d, 'C8', 'B8 = 0.85 is an inference, not a quoted figure: 13 of 15 Tier-1 prices land on a round-dollar base when divided by it.',
    size=9, i=True, color='FF595959')

section(d, 11, 'Tier structure & reconciliation against the as-quoted totals', DLAST)
TH = {'A':'Tier','B':'Qty / part','C':'Lead time (bus. days)','D':'Parts subtotal ($)','E':'Sales tax ($)',
      'F':'Freight, DAP ($)','G':'Landed total ($)','H':'As-quoted total ($)','I':'Check'}
d.row_dimensions[12].height = 34.5
for col, txt in TH.items(): hdr(d, f'{col}12', txt)
for col in 'JKL': sty(d[f'{col}12'], size=11, b=True, color='FFFFFFFF', fill=NAVY)
for r, (label, tier, fill) in enumerate([('Tier 1 — prototype','t1',ORNG_D),
                                         ('Tier 2 — production','t2',ORNG_D)], start=13):
    s = STATED[tier]
    put(d, f'A{r}', label, b=True, fill=fill)
    put(d, f'B{r}', s['qty'],      fill=INPUT_F, h='center', v='center')
    put(d, f'C{r}', s['lt_days'],  fill=INPUT_F, h='center', v='center')
    put(d, f'D{r}', f'=SUM({"F" if tier=="t1" else "H"}19:{"F" if tier=="t1" else "H"}33)', fmt=CUR, fill=fill)
    put(d, f'E{r}', f'=D{r}*$B$7', fmt=CUR, fill=fill)
    put(d, f'F{r}', s['ship'],     fmt=CUR, fill=INPUT_F)
    put(d, f'G{r}', f'=D{r}+E{r}+F{r}', b=True, fmt=CUR, fill=fill)
    put(d, f'H{r}', s['total'],    fmt=CUR, fill=INPUT_F)
    put(d, f'I{r}', f'=IF(ABS(G{r}-H{r})<0.02,"OK","MISMATCH")', b=True, h='center', v='center', fill=fill)
    for col in 'JKL': sty(d[f'{col}{r}'], fill=fill)

section(d, 16, 'As-quoted line items — all 15 parts, both tiers, with price-structure back-solve', DLAST)
d.row_dimensions[17].height = 21.75
for rng, txt, fill, col in [('A17:D17','Part',NAVY,'FFFFFFFF'),
                            ('E17:F17','Tier 1 — qty 2 @ 26 business days',ORNG_H,'FF000000'),
                            ('G17:H17','Tier 2 — qty 30 @ 45 business days',ORNG_H,'FF000000'),
                            ('I17:I17','Ratio',MIDBLUE,'FFFFFFFF'),
                            ('J17:K17','Price-structure back-solve',MIDBLUE,'FFFFFFFF'),
                            ('L17:L17','DFM',NAVY,'FFFFFFFF')]:
    a, b = rng.split(':')
    if a != b: d.merge_cells(rng)
    for c in d[rng][0]: sty(c, size=11, b=True, color=col, fill=fill, h='center', v='center')
    d[a].value = txt
d.row_dimensions[18].height = 34.5
LH = {'A':'#','B':'Part No.','C':'Description','D':'Quoted file','E':'Unit ($)','F':'Ext ($)',
      'G':'Unit ($)','H':'Ext ($)','I':'T2 ÷ T1','J':'Base @ divisor ($)','K':'Round-$ base?','L':'Remarks / DFM feedback'}
for col, txt in LH.items():
    hdr(d, f'{col}18', txt, ORNG_H if col in 'EFGH' else NAVY, 'FF000000' if col in 'EFGH' else 'FFFFFFFF')

DESCR = {prev[f'B{r}']: prev[f'C{r}'] for r in range(6, 21)}
REM = {'BI+TOL': '1. Brass inserts will be used equivalent instead.  2. Printing tolerance will be +/-0.80mm.',
       'TOL':    'Printing tolerance will be +/-0.80mm.'}
for n, (pn, fname, u1, e1, u2, e2, rem) in enumerate(Q5):
    r = 19 + n
    d.row_dimensions[r].height = 15
    put(d, f'A{r}', n + 1)
    put(d, f'B{r}', pn)
    put(d, f'C{r}', DESCR[pn])
    put(d, f'D{r}', fname, size=9)
    put(d, f'E{r}', u1, fmt=CUR, fill=INPUT_F)
    put(d, f'F{r}', f'=E{r}*$B$13', fmt=CUR, fill=ORNG_D)
    put(d, f'G{r}', u2, fmt=CUR, fill=INPUT_F)
    put(d, f'H{r}', f'=G{r}*$B$14', fmt=CUR, fill=ORNG_D)
    put(d, f'I{r}', f'=G{r}/E{r}', fmt=PCT2)
    put(d, f'J{r}', f'=E{r}*$B$8', fmt=CUR)
    put(d, f'K{r}', f'=IF(ABS(J{r}-ROUND(J{r},0))<0.02,"Yes — $"&TEXT(ROUND(J{r},0),"#,##0"),"NO — breaks pattern")')
    put(d, f'L{r}', REM[rem], size=9)
put(d, 'C34', 'TOTAL — as quoted', b=True, fill=YELLOW)
for col in 'ABDEGIJKL': put(d, f'{col}34', None, b=True, fill=YELLOW)
put(d, 'F34', '=SUM(F19:F33)', b=True, fmt=CUR, fill=YELLOW)
put(d, 'H34', '=SUM(H19:H33)', b=True, fmt=CUR, fill=YELLOW)

section(d, 36, 'Landed cost per part — sales tax + DAP freight allocated pro-rata by part value', DLAST)
d.row_dimensions[37].height = 34.5
LC = {'A':'#','B':'Part No.','C':'Description','D':'T1 Ext ($)','E':'T1 freight alloc. ($)',
      'F':'T1 landed ext ($)','G':'T1 landed / unit ($)','H':'T2 Ext ($)','I':'T2 freight alloc. ($)',
      'J':'T2 landed ext ($)','K':'T2 landed / unit ($)','L':'T2 landed unit as % of T1'}
for col, txt in LC.items(): hdr(d, f'{col}37', txt)
for n in range(15):
    r, src = 38 + n, 19 + n
    d.row_dimensions[r].height = 15
    put(d, f'A{r}', f'=A{src}')
    put(d, f'B{r}', f'=B{src}')
    put(d, f'C{r}', f'=C{src}')
    put(d, f'D{r}', f'=F{src}', fmt=CUR)
    put(d, f'E{r}', f'=$F$13*D{r}/$D$13', fmt=CUR)
    put(d, f'F{r}', f'=D{r}*(1+$B$7)+E{r}', fmt=CUR)
    put(d, f'G{r}', f'=F{r}/$B$13', b=True, fmt=CUR, fill=ORNG_D)
    put(d, f'H{r}', f'=H{src}', fmt=CUR)
    put(d, f'I{r}', f'=$F$14*H{r}/$D$14', fmt=CUR)
    put(d, f'J{r}', f'=H{r}*(1+$B$7)+I{r}', fmt=CUR)
    put(d, f'K{r}', f'=J{r}/$B$14', b=True, fmt=CUR, fill=ORNG_D)
    put(d, f'L{r}', f'=K{r}/G{r}', fmt=PCT2)
put(d, 'C53', 'TOTAL landed', b=True, fill=YELLOW)
for col in 'ABL': put(d, f'{col}53', None, b=True, fill=YELLOW)
for col in 'DEFHIJ': put(d, f'{col}53', f'=SUM({col}38:{col}52)', b=True, fmt=CUR, fill=YELLOW)
put(d, 'G53', '=F53/($B$13*15)', b=True, fmt=CUR, fill=YELLOW)
put(d, 'K53', '=J53/($B$14*15)', b=True, fmt=CUR, fill=YELLOW)
note(d, 54, 'G53 / K53 are the blended landed cost per piece across all 15 parts (landed total ÷ pieces), '
            'not a sum — the per-part figures above them are the ones to quote.', DLAST)

section(d, 56, 'Audit findings', DLAST)
AUD = [
 ('Volume-tier ratio — lowest across the 15 parts', '=MIN(I19:I33)', PCT2,
  'Tier 2 unit price as a share of Tier 1.'),
 ('Volume-tier ratio — highest across the 15 parts', '=MAX(I19:I33)', PCT2, ''),
 ('Spread (highest − lowest)', '=D58-D57', PCT2,
  'A spread this close to zero means one flat multiplier was applied to every part, not part-by-part volume economics.'),
 ('Implied volume discount, Tier 1 → Tier 2', '=1-AVERAGE(I19:I33)', PCT,
  'A flat cut applied uniformly — so it is a pricing policy, not a reflection of how these parts nest or batch.'),
 ('Parts whose Tier-1 price back-solves to a round-dollar base', '=SUMPRODUCT(--(ABS(J19:J33-ROUND(J19:J33,0))<0.02))', CNT,
  'At the B8 divisor. The two exceptions are 670832 / 670833 — the mirrored rim pair.'),
 ('Implied gross margin at that back-solved base', '=1-$B$8', PCT,
  'If the round-dollar bases are the underlying cost, this is the margin sitting on top.'),
 ('Tier 1 — blended landed cost per piece', '=G53', CUR, 'Landed total ÷ 30 pieces (15 parts × qty 2).'),
 ('Tier 2 — blended landed cost per piece', '=K53', CUR, 'Landed total ÷ 450 pieces (15 parts × qty 30).'),
 ('Freight as % of parts subtotal — Tier 1', '=F13/D13', PCT, ''),
 ('Freight as % of parts subtotal — Tier 2', '=F14/D14', PCT,
  'Freight scales close to linearly with part value, so there is little freight leverage in ordering more.'),
]
d.column_dimensions['A'].width = 4
for n, (label, formula, fmt, comment) in enumerate(AUD):
    r = 57 + n
    d.row_dimensions[r].height = 15
    d.merge_cells(f'A{r}:C{r}')
    put(d, f'A{r}', label)
    for col in 'BC': sty(d[f'{col}{r}'])
    put(d, f'D{r}', formula, b=True, fmt=fmt, fill=ORNG_D)
    d.merge_cells(f'E{r}:L{r}')
    put(d, f'E{r}', comment, size=9, i=True, color='FF595959')
    for col in 'FGHIJKL': sty(d[f'{col}{r}'], size=9)

for i, txt in enumerate([
 'Source: 3-page quote images supplied by the user, Sep 3, 2026. Vendor not named on those pages.',
 'Transcription proof: the 15 line items sum to the quote’s own stated part-price totals ($23,929.48 Tier 1 / '
 '$161,523.60 Tier 2), and parts × 1.075 + freight reproduces the stated grand totals ($26,772.51 / $182,193.19) to the cent.',
 'Assumption — tax base: the 7.5% is applied to the part price only, not to freight. This is what makes both stated '
 'grand totals reconcile exactly, so it is confirmed by the quote’s own arithmetic rather than assumed.',
 'Assumption — freight allocation: DAP freight is spread across parts in proportion to part value. The quote gives one '
 'freight figure per tier, not per part, so any split is a modelling choice; by weight or volume would shift it.',
 'Assumption — back-solve divisor (B8 = 0.85): inferred from the data, not quoted. Change B8 to test other margin '
 'assumptions; column K re-tests every part against it.',
 'Lead times exclude the 9.25–9.27 and 10.1–10.7 holiday periods — both fall inside a 26- or 45-business-day '
 'window starting Sep 3, so the calendar dates land later than the day counts suggest.',
]):
    note(d, 68 + i, txt, DLAST)
    d.row_dimensions[68+i].height = 24
    d[f'A{68+i}'].alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)

wb.save(OUT)
print('stage 1 saved:', OUT, '| sheets:', wb.sheetnames)
