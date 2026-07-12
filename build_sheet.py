import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker

wb = openpyxl.Workbook()

TITLE = Font(bold=True, size=14, color="FFFFFF")
HDR   = Font(bold=True, size=10, color="FFFFFF")
BOLD  = Font(bold=True)
input_fill = PatternFill("solid", fgColor="FFF2CC")
hdr_fill   = PatternFill("solid", fgColor="1F4E78")
prob_hdr   = PatternFill("solid", fgColor="375623")
title_fill = PatternFill("solid", fgColor="2E75B6")
you_fill   = PatternFill("solid", fgColor="FCE4D6")
her_fill   = PatternFill("solid", fgColor="E2EFDA")
prob_fill  = PatternFill("solid", fgColor="EDF7E9")
grp_fill   = PatternFill("solid", fgColor="DDEBF7")
param_fill = PatternFill("solid", fgColor="FFF2CC")
money='#,##0'; pct='0.0%'; pct0='0%'
thin=Side(style="thin",color="BFBFBF")
border=Border(left=thin,right=thin,top=thin,bottom=thin)

def cell(ws,coord,val,font=None,fill=None,fmt=None,align=None,bd=False):
    c=ws[coord]; c.value=val
    if font:c.font=font
    if fill:c.fill=fill
    if fmt:c.number_format=fmt
    if align:c.alignment=Alignment(horizontal=align,vertical="center",wrap_text=True)
    if bd:c.border=border
    return c

ws=wb.active; ws.title="Scenarios"
cell(ws,"A1","House Sale Scenarios + Probability of Sale — 537 Duchart Ln, Fuquay-Varina NC",TITLE,title_fill)
ws.merge_cells("A1:R1"); ws.row_dimensions[1].height=24

# ---- DEAL INPUTS (left) ----
cell(ws,"A3","DEAL INPUTS (edit yellow)",BOLD)
inputs=[("A4","Purchase price (sold Mar 18 2024)","B4",405700,money),
        ("A5","Down payment (her contribution)","B5",48000,money),
        ("A6","Current mortgage balance / payoff","B6",356149.76,money),
        ("A7","Excise/transfer tax rate (NC)","B7",0.002,pct),
        ("A8","Other closing costs (flat $)","B8",2000,money),
        ("A9","HER recovery share — Option A","B9",0.80,pct),
        ("A10","HER recovery share — Option B","B10",0.70,pct)]
for lc,lab,vc,val,fmt in inputs:
    cell(ws,lc,lab); cell(ws,vc,val,BOLD,input_fill,fmt,bd=True)

# ---- PROBABILITY MODEL PARAMETERS (right, above table) ----
cell(ws,"N3","PROBABILITY MODEL (edit yellow)",BOLD)
params=[("N4","Est. current market value","O4",410000,money),
        ("N5","Base monthly sale chance (h0)","O5",0.25,pct),
        ("N6","Price sensitivity (k)","O6",7,'0.0'),
        ("N7","Exposure: full-service agent","O7",1.10,'0.00'),
        ("N8","Exposure: typical agent","O8",1.05,'0.00'),
        ("N9","Exposure: FSBO + buyer agent 3%","O9",0.90,'0.00'),
        ("N10","Exposure: FSBO + buyer agent 2.5%","O10",0.85,'0.00'),
        ("N11","Season — Jul 2026","O11",1.05,'0.00'),
        ("N12","Season — Aug 2026","O12",0.95,'0.00'),
        ("N13","Season — Sep 2026","O13",0.85,'0.00'),
        ("N14","Season — Oct 2026","O14",0.75,'0.00')]
for lc,lab,vc,val,fmt in params:
    cell(ws,lc,lab); cell(ws,vc,val,BOLD,param_fill,fmt,bd=True)

# ---- TABLE ----
H=18
headers=["Sale price","Selling option","Comm %","Comm $","Other\nclosing $",
         "Total\nsell costs","NET proceeds\n(money back)","Shortfall\nvs $48k",
         "YOUR pay\n@80%","YOUR pay\n@70%","Her total\n@80%","Her total\n@70%",
         "Price\nfactor","Exposure\nfactor","P(sold by\nend Jul)","P(sold by\nend Aug)",
         "P(sold by\nend Sep)","P(sold by\nend Oct)","★ RECOMMENDATION"]
tag_hdr=PatternFill("solid",fgColor="BF8F00")
tag_fill=PatternFill("solid",fgColor="FFE699")
tags={(405000,"FSBO + buyer agent (2.5%)"):"★ BEST VALUE — low cost to Oscar + decent odds",
      (409900,"FSBO + buyer agent (2.5%)"):"★ CHEAPEST for Oscar — list high, cut later if needed",
      (399000,"Typical agent (5.5%)"):"★ BEST ODDS / hands-off if you want it sold"}
for i,h in enumerate(headers):
    fill = tag_hdr if i==18 else (prob_hdr if i>=12 else hdr_fill)
    cell(ws,f"{get_column_letter(i+1)}{H}",h,HDR,fill,align="center",bd=True)
ws.row_dimensions[H].height=42

sale_prices=[390000,395000,399000,405000,409900]
# option: (label, commission, exposure-param-cell)
options=[("Full-service agent (6%)",0.06,"$O$7"),
         ("Typical agent (5.5%)",0.055,"$O$8"),
         ("FSBO + buyer agent (3%)",0.03,"$O$9"),
         ("FSBO + buyer agent (2.5%)",0.025,"$O$10")]

r=H+1
for sp in sale_prices:
    for name,comm,exp_cell in options:
        a=f"A{r}";c=f"C{r}";d=f"D{r}";e=f"E{r}";f=f"F{r}";g=f"G{r}";h=f"H{r}"
        i=f"I{r}";j=f"J{r}";k=f"K{r}";l=f"L{r}";m=f"M{r}";n=f"N{r}"
        o=f"O{r}";p=f"P{r}";q=f"Q{r}";rr=f"R{r}"
        cell(ws,a,sp,None,grp_fill,money,bd=True)
        cell(ws,f"B{r}",name,None,None,None,"left",bd=True)
        cell(ws,c,comm,None,None,pct,bd=True)
        cell(ws,d,f"={a}*{c}",None,None,money,bd=True)
        cell(ws,e,f"={a}*$B$7+$B$8",None,None,money,bd=True)
        cell(ws,f,f"={d}+{e}",None,None,money,bd=True)
        cell(ws,g,f"={a}-$B$6-{f}",BOLD,None,money,bd=True)
        cell(ws,h,f"=MAX(0,$B$5-{g})",None,None,money,bd=True)
        cell(ws,i,f"={h}*$B$9",BOLD,you_fill,money,bd=True)
        cell(ws,j,f"={h}*$B$10",BOLD,you_fill,money,bd=True)
        cell(ws,k,f"={g}+{i}",None,her_fill,money,bd=True)
        cell(ws,l,f"={g}+{j}",None,her_fill,money,bd=True)
        # probability model
        cell(ws,m,f"=EXP(-$O$6*({a}-$O$4)/$O$4)",None,None,'0.00',bd=True)
        cell(ws,n,f"={exp_cell}",None,None,'0.00',bd=True)
        cell(ws,o,f"=MIN(0.95,$O$5*{m}*{n}*$O$11)",BOLD,prob_fill,pct0,bd=True)
        cell(ws,p,f"=1-(1-{o})*(1-MIN(0.95,$O$5*{m}*{n}*$O$12))",None,prob_fill,pct0,bd=True)
        cell(ws,q,f"=1-(1-{p})*(1-MIN(0.95,$O$5*{m}*{n}*$O$13))",None,prob_fill,pct0,bd=True)
        cell(ws,rr,f"=1-(1-{q})*(1-MIN(0.95,$O$5*{m}*{n}*$O$14))",BOLD,prob_fill,pct0,bd=True)
        tg=tags.get((sp,name))
        if tg:
            cell(ws,f"S{r}",tg,BOLD,tag_fill,None,"left",bd=True)
            for col in "ABIJOPQR":
                ws[f"{col}{r}"].fill=tag_fill
        else:
            cell(ws,f"S{r}","",None,None,None,None,bd=True)
        r+=1
LAST=r-1

widths=[12,28,8,11,10,11,13,11,11,11,11,11,9,9,10,10,10,10,46]
for i,w in enumerate(widths):
    ws.column_dimensions[get_column_letter(i+1)].width=w
ws.freeze_panes=f"A{H+1}"

# ---- AVERAGES / SUMMARY BLOCK ----
S=LAST+2
cell(ws,f"A{S}","AVERAGE ACROSS ALL 20 SCENARIOS (equal weight)",HDR,hdr_fill)
ws.merge_cells(f"A{S}:F{S}")
rng=f"{H+1}:{LAST}"
hcol=f"H{H+1}:H{LAST}"; icol=f"I{H+1}:I{LAST}"; jcol=f"J{H+1}:J{LAST}"
summ=[("Average loss of Francie's $48k capital (shortfall)",f"=AVERAGE({hcol})"),
      ("Average she keeps of the $48k",f"=$B$5-AVERAGE({hcol})"),
      ("Average COST TO OSCAR @ 80%",f"=AVERAGE({icol})"),
      ("Average COST TO OSCAR @ 70%",f"=AVERAGE({jcol})")]
rr=S+1
for lab,formula in summ:
    cell(ws,f"A{rr}",lab,BOLD); ws.merge_cells(f"A{rr}:F{rr}")
    cell(ws,f"G{rr}",formula,BOLD,you_fill if "OSCAR" in lab else her_fill,money,bd=True)
    rr+=1

rr+=1
cell(ws,f"A{rr}","AVERAGE COST TO OSCAR BY ROUTE (across the 5 prices)",HDR,hdr_fill)
ws.merge_cells(f"A{rr}:F{rr}"); rr+=1
routes=[("Full-service agent (6%)",[19,23,27,31,35]),
        ("Typical agent (5.5%)",[20,24,28,32,36]),
        ("FSBO + buyer agent (3%)",[21,25,29,33,37]),
        ("FSBO + buyer agent (2.5%)",[22,26,30,34,38])]
cell(ws,f"A{rr}","Route",BOLD,grp_fill,bd=True); ws.merge_cells(f"A{rr}:D{rr}")
cell(ws,f"E{rr}","Avg loss of $48k",BOLD,grp_fill,align="center",bd=True); ws.merge_cells(f"E{rr}:F{rr}")
cell(ws,f"G{rr}","Oscar @80%",BOLD,grp_fill,align="center",bd=True)
cell(ws,f"H{rr}","Oscar @70%",BOLD,grp_fill,align="center",bd=True); rr+=1
for label,rows_ in routes:
    cell(ws,f"A{rr}",label,None,None,None,"left",bd=True); ws.merge_cells(f"A{rr}:D{rr}")
    cell(ws,f"E{rr}","=AVERAGE("+",".join(f'H{x}' for x in rows_)+")",None,None,money,bd=True); ws.merge_cells(f"E{rr}:F{rr}")
    cell(ws,f"G{rr}","=AVERAGE("+",".join(f'I{x}' for x in rows_)+")",BOLD,you_fill,money,bd=True)
    cell(ws,f"H{rr}","=AVERAGE("+",".join(f'J{x}' for x in rows_)+")",BOLD,you_fill,money,bd=True)
    rr+=1

# ================= NOTES =================
ns=wb.create_sheet("Notes & Assumptions"); ns.column_dimensions["A"].width=104
notes=[
 ("House Sale Scenarios + Probability of Sale — Notes",TITLE,title_fill),
 ("",None,None),
 ("PROPERTY (verified online)",BOLD,None),
 ("• 537 Duchart Ln, Fuquay-Varina, NC 27526 — 3bd/2.5ba, 2,216 sqft, Lakestone Village.",None,None),
 ("• Recorded sale: Mar 18, 2024 for $405,700.  Currently listed $409,900.",None,None),
 ("",None,None),
 ("MORTGAGE (from your screenshot, loan *5990)",BOLD,None),
 ("• Payoff used: $356,149.76. Payment $3,050.77/mo. Get an official payoff quote — true payoff",None,None),
 ("  runs a bit higher (per-day interest). Only ~$1,550 principal paid since 2024 (normal this early).",None,None),
 ("",None,None),
 ("DEAL MATH (unchanged from v1)",BOLD,None),
 ("• NET proceeds = Sale price − payoff − selling costs = 'money back' toward her $48,000.",None,None),
 ("• Shortfall = $48,000 − NET proceeds.  YOUR payment = Shortfall × 80% (or 70%).",None,None),
 ("",None,None),
 ("CURRENT MARKET — Fuquay-Varina (June 2026)",BOLD,None),
 ("• Median days on market ≈ 85 days (flat vs a year ago) — a slow, normalizing market.",None,None),
 ("• Inventory elevated: ~700+ active listings. Prices down ~3–4% year-over-year.",None,None),
 ("• Net read: this is a BUYER-leaning market. Overpricing sits; competitive pricing moves.",None,None),
 ("• Your home has already sat ~2 weeks at $409,900 with no contract — consistent with the model's",None,None),
 ("  low near-term odds at the top price.",None,None),
 ("",None,None),
 ("SEASONALITY (Triangle / Raleigh area) — answers 'when do odds rise?'",BOLD,None),
 ("• Peak selling season is APRIL–JUNE (fastest sales, best prices). July still strong.",None,None),
 ("• Aug → Oct demand TAPERS; Nov–Jan is the slowest window; odds rebound Feb–Apr.",None,None),
 ("• So over your next 4 months (Jul→Oct 2026) the seasonal trend works AGAINST you — the monthly",None,None),
 ("  chance of selling DECLINES each month. The next natural uptick is ~Feb–Apr 2027.",None,None),
 ("• Bottom line: in the near term, PRICE and EXPOSURE (using an agent) are the levers that raise",None,None),
 ("  your odds — not waiting. Waiting into fall lowers them.",None,None),
 ("",None,None),
 ("HOW THE PROBABILITY COLUMNS ARE BUILT (all factors editable on the Scenarios tab)",BOLD,None),
 ("• Monthly chance of going UNDER CONTRACT = h0 × Price factor × Exposure factor × Season factor.",None,None),
 ("• h0 = 0.25/month: the baseline ~ derived from an ~85-day median time-to-contract.",None,None),
 ("• Price factor = EXP(−k × (price − market value)/market value), k=7. Pricing BELOW the ~$410k",None,None),
 ("  market value raises odds; pricing at/above lowers them. Adjust 'market value' and k to taste.",None,None),
 ("• Exposure factor: full-service 1.10, typical agent 1.05, FSBO+3% buyer agent 0.90, FSBO+2.5% 0.85.",None,None),
 ("  ASSUMPTION: the buyer always has a realtor, so you ALWAYS pay a buyer's-agent commission. 'FSBO'",None,None),
 ("  means you skip the LISTING agent only (saving ~3%) and still pay the buyer's agent — there is no",None,None),
 ("  0% / no-commission case. FSBO is penalized a little (amateur photos/marketing/negotiation), not a",None,None),
 ("  lot, because a competitive buyer-agent commission still gets the home shown.",None,None),
 ("• Season factor: Jul 1.05, Aug 0.95, Sep 0.85, Oct 0.75 (declining into fall, per above).",None,None),
 ("• The four 'P(sold by end of month)' columns are CUMULATIVE — the chance you're under contract",None,None),
 ("  by the end of that month. They rise across months because each month adds another chance,",None,None),
 ("  even though the per-month odds shrink seasonally.",None,None),
 ("• 'Sold' here = under contract / offer accepted, not the closing date (closing is ~30–45 days later).",None,None),
 ("",None,None),
 ("TAGGED 'BEST' OPTIONS (★ on the Scenarios tab) — why these",BOLD,None),
 ("• ★ CHEAPEST for Oscar: $409,900 FSBO + buyer agent 2.5% → you pay ~$5,900, ~57% sold by Oct.",None,None),
 ("  Logic: list high as FSBO; you can always cut the price later, but you can't un-cut it. Lowest cost.",None,None),
 ("• ★ BEST VALUE: $405,000 FSBO + buyer agent 2.5% → you pay ~$9,700, ~61% by Oct. A small price",None,None),
 ("  trim buys better odds while still skipping the listing commission — the best cost-vs-odds balance.",None,None),
 ("• ★ BEST ODDS / hands-off: $399,000 typical agent (5.5%) → you pay ~$23,900, ~74% by Oct. Pick",None,None),
 ("  this if a fast, certain, low-effort sale matters more than saving money. Costs you the most.",None,None),
 ("• Pattern: every $1 of commission or price cut comes straight out of Oscar's pocket via the 80/70",None,None),
 ("  split, because it shrinks the net proceeds that repay her $48k. FSBO is Oscar's biggest saver.",None,None),
 ("",None,None),
 ("AVERAGE LOSS OF FRANCIE'S $48k (summary block under the table)",BOLD,None),
 ("• Averaged equally across all 20 scenarios: ~$24,200 of her $48k is NOT returned by the sale",None,None),
 ("  (she gets back ~$23,800). Oscar's average cost is 80% of that ≈ $19,300 (or ~$16,900 at 70%).",None,None),
 ("• This equal-weight average spans cheap and expensive cases; your REAL number depends on the",None,None),
 ("  option you choose — see the per-route averages: FSBO routes cut Oscar's average cost to ~$14k,",None,None),
 ("  vs ~$25k with a full-service agent.",None,None),
 ("",None,None),
 ("THE CORE TENSION",BOLD,None),
 ("• Lower price / using an agent = HIGHER chance of selling soon, but LOWER net proceeds = you pay",None,None),
 ("  her MORE. Higher price / FSBO = you pay her less IF it sells, but lower odds it sells at all.",None,None),
 ("• Read the table as a trade-off: cross-reference 'YOUR pay' against 'P(sold by end Oct)' per row.",None,None),
 ("",None,None),
 ("These are estimates to compare scenarios, not a guarantee, appraisal, or legal/financial advice.",None,None),
 ("Probabilities are modeled, not market-quoted. Confirm payoff & costs with your servicer/attorney.",None,None),
 ("",None,None),
 ("SOURCES",BOLD,None),
 ("• Listing/sale history: Coldwell Banker & Apartments.com (537 Duchart Ln).",None,None),
 ("• Market: Redfin & Long&Foster Market Minute (Fuquay-Varina DOM/inventory/prices, 2026).",None,None),
 ("• Seasonality: Houzeo, HomeLight, ListWithClever (best time to sell Raleigh/NC, 2025–2026).",None,None),
 ("• Closing costs: ListWithClever, Houzeo, Bankrate, Redfin (NC seller costs & transfer tax).",None,None),
]
for idx,(txt,font,fill) in enumerate(notes,start=1):
    cell(ns,f"A{idx}",txt,font,fill)
ns.row_dimensions[1].height=24

# ================= CHART: COST vs ODDS =================
cs=wb.create_sheet("Chart — Cost vs Odds")
cell(cs,"A1","Trade-off: what Oscar pays (@80%) vs. odds the house is sold by end of October",TITLE,title_fill)
cs.merge_cells("A1:H1"); cs.row_dimensions[1].height=22

# helper table: ALL 20 points (live refs into Scenarios)
cell(cs,"A3","All options",BOLD)
cell(cs,"A4","P(sold by Oct)",BOLD); cell(cs,"B4","Oscar pay @80%",BOLD)
hr=5
for rsrc in range(H+1,LAST+1):
    cell(cs,f"A{hr}",f"=Scenarios!R{rsrc}",fmt=pct0)
    cell(cs,f"B{hr}",f"=Scenarios!I{rsrc}",fmt=money)
    hr+=1
all_end=hr-1

# helper table: the 3 tagged picks (separate highlighted series)
cell(cs,"D3","★ Recommended picks",BOLD)
cell(cs,"D4","P(sold by Oct)",BOLD); cell(cs,"E4","Oscar pay @80%",BOLD); cell(cs,"F4","Label",BOLD)
pick_rows={34:"BEST VALUE ($405k FSBO 2.5%)",38:"CHEAPEST ($409.9k FSBO 2.5%)",28:"BEST ODDS ($399k agent 5.5%)"}
pr=5
for srcrow,lab in pick_rows.items():
    cell(cs,f"D{pr}",f"=Scenarios!R{srcrow}",fmt=pct0)
    cell(cs,f"E{pr}",f"=Scenarios!I{srcrow}",fmt=money)
    cell(cs,f"F{pr}",lab)
    pr+=1
pick_end=pr-1

chart=ScatterChart()
chart.title="Oscar's cost vs. probability of sale by October"
chart.x_axis.title="Probability sold by end of October"
chart.y_axis.title="Oscar's payment to Francie (@80%)"
chart.x_axis.numFmt='0%'; chart.y_axis.numFmt='$#,##0'
chart.x_axis.delete=False; chart.y_axis.delete=False
chart.height=12; chart.width=22

xall=Reference(cs,min_col=1,min_row=5,max_row=all_end)
yall=Reference(cs,min_col=2,min_row=5,max_row=all_end)
s_all=Series(yall,xall,title="All 20 options")
s_all.marker=Marker(symbol="circle",size=6)
s_all.graphicalProperties.line.noFill=True
chart.series.append(s_all)

xpk=Reference(cs,min_col=4,min_row=5,max_row=pick_end)
ypk=Reference(cs,min_col=5,min_row=5,max_row=pick_end)
s_pk=Series(ypk,xpk,title="★ Recommended")
s_pk.marker=Marker(symbol="star",size=12)
s_pk.graphicalProperties.line.noFill=True
chart.series.append(s_pk)

cs.add_chart(chart,"H3")
cell(cs,"A22","Read it like this: down-and-to-the-right is ideal (cheap for Oscar AND likely to sell).",BOLD)
cell(cs,"A23","The ★ stars are the tagged picks. Points up high cost Oscar the most; points to the left are least likely to sell.")
cs.column_dimensions["A"].width=15; cs.column_dimensions["B"].width=15
cs.column_dimensions["D"].width=15; cs.column_dimensions["E"].width=15; cs.column_dimensions["F"].width=30

# ================= DIVORCE SETTLEMENT TAB =================
ds=wb.create_sheet("Divorce Settlement")
cell(ds,"A1","Divorce Settlement Scenarios (NC) — what it costs Oscar",TITLE,title_fill)
ds.merge_cells("A1:F1"); ds.row_dimensions[1].height=22
cell(ds,"A2","General information, NOT legal advice. NC equitable distribution is fact-specific — confirm with a NC family-law attorney.",Font(italic=True,size=9))

cell(ds,"A4","INPUTS (edit yellow)",BOLD)
inp=[("Mortgage payoff","B5",356149.76,money,False),
     ("Her separate down payment (traces to house)","B6",48000,money,False),
     ("Your marital 401k (Aug23–May26)","B7",20000,money,False),
     ("Net 401k she can claim (her ~$12k offsets yours)","B8","=MAX(0,(B7-B15)/2)",money,True),
     ("Joint account balance (split 50/50)","B9",0,money,False),
     ("Your post-separation mortgage credit (optional)","B10",0,money,False),
     ("Representative sale price","B11",399000,money,False),
     ("Commission %","B12",0.055,pct,False),
     ("Transfer tax %","B13",0.002,pct,False),
     ("Other closing flat $","B14",2000,money,False),
     ("Her marital 401k (Aug23–May26)","B15",12000,money,False)]
r=5
for lab,cref,val,fmt,isf in inp:
    cell(ds,f"A{r}",lab); cell(ds,cref,val,BOLD,None if isf else param_fill,fmt,bd=True); r+=1

cell(ds,"A16","Selling costs"); cell(ds,"B16","=B11*B12+B11*B13+B14",None,None,money,bd=True)
cell(ds,"A17","Net proceeds (she takes as her credit)",BOLD); cell(ds,"B17","=B11-B5-B16",BOLD,her_fill,money,bd=True)
cell(ds,"A18","Shortfall vs her $48k (market loss on HER money)"); cell(ds,"B18","=MAX(0,B6-B17)",None,None,money,bd=True)

T=20
heads=["Settlement option","Cash you pay her","Assets you give up","YOUR TOTAL COST","Her total recovery","vs her ask, you SAVE"]
for i,h in enumerate(heads):
    cell(ds,f"{get_column_letter(i+1)}{T}",h,HDR,hdr_fill,align="center",bd=True)
ds.row_dimensions[T].height=30
# rows 21..24
cell(ds,"A21","Her full ask — 80% of shortfall"); cell(ds,"B21","=B18*0.8",None,None,money,bd=True); cell(ds,"C21",0,None,None,money,bd=True)
cell(ds,"A22","Her ask — 70% of shortfall"); cell(ds,"B22","=B18*0.7",None,None,money,bd=True); cell(ds,"C22",0,None,None,money,bd=True)
cell(ds,"A23","Negotiated middle (give your 401k share)"); cell(ds,"B23",0,None,None,money,bd=True); cell(ds,"C23","=B8",None,None,money,bd=True)
cell(ds,"A24","★ Clean trade / legal baseline (recommended)",BOLD,tag_fill); cell(ds,"B24",0,None,tag_fill,money,bd=True); cell(ds,"C24","=MAX(0,B8-B10)",None,tag_fill,money,bd=True)
for rw in (21,22,23,24):
    fillv = tag_fill if rw==24 else None
    cell(ds,f"D{rw}",f"=B{rw}+C{rw}",BOLD,fillv if rw==24 else you_fill,money,bd=True)
    cell(ds,f"E{rw}",f"=B17+D{rw}",None,fillv if rw==24 else her_fill,money,bd=True)
    cell(ds,f"F{rw}",f"=$D$21-D{rw}",BOLD,fillv,money,bd=True)
    if rw==24:
        cell(ds,f"A{rw}","★ Clean trade / legal baseline (recommended)",BOLD,tag_fill,None,"left",bd=True)

# savings-by-price mini matrix
M=27
cell(ds,f"A{M}","WHAT HER 80% ASK COSTS YOU vs A CLEAN TRADE — by sale price",HDR,hdr_fill)
ds.merge_cells(f"A{M}:F{M}")
mh=["Sale price","Her 80% ask costs you","Clean trade costs you","You SAVE"]
for i,h in enumerate(mh):
    cell(ds,f"{get_column_letter(i+1)}{M+1}",h,BOLD,grp_fill,align="center",bd=True)
mr=M+2
for p in [390000,395000,399000,405000,409900]:
    cell(ds,f"A{mr}",p,None,None,money,bd=True)
    cell(ds,f"B{mr}",f"=MAX(0,$B$6-({p}-$B$5-({p}*$B$12+{p}*$B$13+$B$14)))*0.8",None,you_fill,money,bd=True)
    cell(ds,f"C{mr}","=MAX(0,$B$8-$B$10)",None,None,money,bd=True)
    cell(ds,f"D{mr}",f"=B{mr}-C{mr}",BOLD,tag_fill,money,bd=True)
    mr+=1

# house options notes
hn=mr+2
optnotes=[
 ("FOUR OPTIONS FOR THE HOUSE (fastest first)",BOLD),
 ("1. SELL now, she takes 100% of net proceeds, both sign a full mutual release. Fast; crystallizes the",None),
 ("   loss — which she bears (it's her separate money that the market moved), not you. You walk from the debt.",None),
 ("2. LUMP-SUM BUYOUT: you keep the house, pay her a fixed sum (~current net equity) to release ALL claims,",None),
 ("   she quitclaims. Fastest + most certain cap on your cost; you keep future appreciation. Needs you to",None),
 ("   refinance the mortgage into your name alone (likely doable on $140k income).",None),
 ("3. RENT & DELAY to spring 2027 (better season + possible price recovery) → smaller loss → less owed. But",None),
 ("   you'd still settle her $48k credit in the divorce now; only do this if you keep the asset yourself.",None),
 ("4. HER DEAL (pay 70–80% of the shortfall): the most expensive for you and NOT required by NC law. Avoid",None),
 ("   signing this as written — it makes you personally insure her investment against the market.",None),
 ("",None),
 ("WHY THE CLEAN TRADE IS YOUR BASELINE",BOLD),
 ("• Her $48k down payment is her SEPARATE property — she can trace it into the house and take the net",None),
 ("  proceeds, but NC caps that at the equity that actually EXISTS. The market shortfall is a loss on HER",None),
 ("  investment; the law does not make you cover it with cash.",None),
 ("• Alimony is NOT barred (her affair was in her PRIOR marriage — legally irrelevant here). But it's likely",None),
 ("  modest/short: ~2.75-yr marriage + her ~$30k savings, car & remaining $90k (separate assets that cut",None),
 ("  against 'need') + her earning capacity. A court might award a small short-term amount — or none. No kids.",None),
 ("• The mutual ALIMONY WAIVER in the agreement is a key part of what your settlement money buys.",None),
 ("• 401k nearly washes out: your ~$20k and her ~$12k marital contributions offset, so a strict 50/50 split",None),
 ("  nets her only ~$4k from you — not the ~$7.5k it looked like before. Trade even that for her dropping the",None),
 ("  shortfall guarantee and you both walk clean. (401k set aside during the marriage is marital for BOTH of",None),
 ("  you regardless of the account — it's funded by wages, not by her savings.)",None),
 ("• You've paid the mortgage from your salary since separation → ask for a post-separation reimbursement",None),
 ("  credit, which offsets anything you'd owe her.",None),
 ("",None),
 ("FAST + CHEAP PATH",BOLD),
 ("• One separation agreement settling everything (property + mutual alimony waiver + house). ~$1.5k–$4k total,",None),
 ("  signed in weeks. Then file for absolute divorce after the 1-yr separation (~May 2027): ~$225 + minimal.",None),
 ("• Do NOT file equitable-distribution / alimony claims in court — that triggers the $15k–$40k litigation path.",None),
]
for i,(txt,font) in enumerate(optnotes):
    cell(ds,f"A{hn+i}",txt,font)
ds.column_dimensions["A"].width=52
for c in "BCDEF": ds.column_dimensions[c].width=17

# ================= WHY $14K WORKS TAB =================
w=wb.create_sheet("Why $14k Works")
cell(w,"A1","Why a $14,000 guaranteed payment is a generous deal",TITLE,title_fill)
w.merge_cells("A1:H1"); w.row_dimensions[1].height=22

cell(w,"A3","INPUTS (edit yellow)",BOLD)
wi=[("Her capital (down payment)","B4",48000,money,False),
    ("Your guaranteed cash to her","B5",14000,money,False),
    ("= that as % of her $48k","B6","=B5/B4",pct,True),
    ("Mortgage payoff","B7",356149.76,money,False),
    ("Commission % (5.5% agent / 2.5% FSBO)","B8",0.055,pct,False),
    ("Transfer tax %","B9",0.002,pct,False),
    ("Other closing flat $","B10",2000,money,False)]
rr=4
for lab,cref,val,fmt,isf in wi:
    cell(w,f"A{rr}",lab); cell(w,cref,val,BOLD,None if isf else param_fill,fmt,bd=True); rr+=1

T=12
heads=["Sale price","Net proceeds\n(her house return)","House return\n% of $48k",
       "Your $14k\n(% of $48k)","Her total\n(flat $14k)","Total\n% of $48k",
       "Capped payment\n(no overpay)","Her total\n(capped)"]
for i,h in enumerate(heads):
    cell(w,f"{get_column_letter(i+1)}{T}",h,HDR,hdr_fill,align="center",bd=True)
w.row_dimensions[T].height=42
r=T+1
for p in [390000,395000,399000,405000,409900]:
    a=f"A{r}";b=f"B{r}";c=f"C{r}";d=f"D{r}";e=f"E{r}";ff=f"F{r}";g=f"G{r}";hh=f"H{r}"
    cell(w,a,p,None,grp_fill,money,bd=True)
    cell(w,b,f"={p}-$B$7-({p}*$B$8+{p}*$B$9+$B$10)",BOLD,her_fill,money,bd=True)
    cell(w,c,f"={b}/$B$4",None,None,pct0,bd=True)
    cell(w,d,"=$B$6",None,you_fill,pct0,bd=True)
    cell(w,e,f"={b}+$B$5",BOLD,None,money,bd=True)
    cell(w,ff,f"={e}/$B$4",BOLD,None,pct0,bd=True)
    cell(w,g,f"=MIN($B$5,MAX(0,$B$4-{b}))",None,you_fill,money,bd=True)
    cell(w,hh,f"={b}+{g}",None,None,money,bd=True)
    r+=1

notes2=[
 ("",None),
 ("HOW TO READ THIS",BOLD),
 ("• Your $14,000 = 29.2% of her $48k, handed to her GUARANTEED — regardless of what the house sells for,",None),
 ("  when it sells, or whether it sits unsold while you keep paying the mortgage.",None),
 ("• Her house return (net proceeds) stacks ON TOP of that 29%. So her total = 29% floor + house return.",None),
 ("• 'Capped payment' column = pay MIN($14k, whatever gets her to exactly $48k). Use this so a strong FSBO",None),
 ("  sale doesn't push her PAST her full $48k (she'd otherwise land at 104–114%). It only ever helps you.",None),
 ("",None),
 ("WHY IT'S GENEROUS (three yardsticks)",BOLD),
 ("• vs. the law (property): NC likely requires only ~$4k (her $48k is separate property capped at actual",None),
 ("  equity; and her ~$12k 401k offsets your ~$20k, so the marital split nets her ~$4k). Alimony is NOT barred",None),
 ("  (her affair was pre-marriage) but is likely modest/short given the ~2.75-yr marriage and her own assets.",None),
 ("  $14k covers property + a full alimony waiver in one number — well above what she'd realistically get.",None),
 ("• vs. her real recovery: guaranteed cash beats an uncertain sale. The house is only ~57–82% likely to be",None),
 ("  under contract by end of October, at an unknown price in a soft market. $14k removes all that risk for her.",None),
 ("• vs. her alternative: fighting it in court likely nets her ~$0–$7,500 minus $10k+ in her own legal fees —",None),
 ("  a negative result. $14k beats that decisively.",None),
 ("",None),
 ("Not legal advice — confirm with a NC family-law attorney before signing anything.",Font(italic=True,size=9)),
]
nr=r+1
for txt,font in notes2:
    cell(w,f"A{nr}",txt,font); nr+=1
w.column_dimensions["A"].width=40
for c in "BCDEFGH": w.column_dimensions[c].width=15

# ================= HER $35K-TOTAL OFFER TAB =================
z=wb.create_sheet("Her $35k-Total Offer")
cell(z,"A1","Francie's $35k-TOTAL offer — modeled & compared to your $14k plan",TITLE,title_fill)
z.merge_cells("A1:H1"); z.row_dimensions[1].height=22
cell(z,"A2","Her logic: (money she gets from the house) + (cash you give her) = $35,000 total, capped. You cover any shortfall and KEEP any house proceeds above $35k.",Font(italic=True,size=9))

cell(z,"A4","INPUTS (edit yellow)",BOLD)
zi=[("Her fixed TOTAL (her offer)","B5",35000,money,False),
    ("Mortgage payoff","B6",356149.76,money,False),
    ("Transfer tax %","B7",0.002,pct,False),
    ("Other closing flat $","B8",2000,money,False),
    ("(compare) Your $14k top-up","B9",14000,money,False),
    ("(compare) Her $48k cap in that plan","B10",48000,money,False)]
rr=5
for lab,cref,val,fmt,isf in zi:
    cell(z,f"A{rr}",lab); cell(z,cref,val,BOLD,param_fill,fmt,bd=True); rr+=1

T=12
heads=["Sale price","Route","Net proceeds\n(house money)","Her TOTAL\n($35k deal)",
       "YOUR cost\n($35k deal)","YOUR cost\n($14k deal)","$35k deal\nsaves you","Cheaper\nfor you"]
for i,h in enumerate(heads):
    cell(z,f"{get_column_letter(i+1)}{T}",h,HDR,hdr_fill,align="center",bd=True)
z.row_dimensions[T].height=42
rows=[(390000,"Agent 5.5%",0.055),(395000,"Agent 5.5%",0.055),(399000,"Agent 5.5%",0.055),
      (405000,"Agent 5.5%",0.055),(409900,"Agent 5.5%",0.055),
      (390000,"FSBO 2.5%",0.025),(395000,"FSBO 2.5%",0.025),(399000,"FSBO 2.5%",0.025),
      (405000,"FSBO 2.5%",0.025),(409900,"FSBO 2.5%",0.025)]
r=T+1
for p,route,c in rows:
    a=f"A{r}";b=f"B{r}";d=f"D{r}";e=f"E{r}";f=f"F{r}";g=f"G{r}";hh=f"H{r}";i=f"I{r}"
    cell(z,a,p,None,grp_fill,money,bd=True)
    cell(z,b,route,None,None,None,"left",bd=True)
    cell(z,f"C{r}",f"={p}-$B$6-({p}*{c}+{p}*$B$7+$B$8)",BOLD,None,money,bd=True)
    cell(z,d,"=$B$5",None,her_fill,money,bd=True)                     # her total always 35k
    cell(z,e,f"=$B$5-C{r}",BOLD,you_fill,money,bd=True)               # your cost = 35k - net
    cell(z,f,f"=MIN($B$9,MAX(0,$B$10-C{r}))",None,None,money,bd=True) # your 14k-deal cost
    cell(z,g,f"={f}-{e}",BOLD,None,money,bd=True)                     # positive = 35k cheaper
    cell(z,hh,f'=IF({e}<{f},"$35k deal","$14k deal")',None,None,None,"center",bd=True)
    r+=1
LASTZ=r-1
# averages
cell(z,f"A{r+1}","AVG — Agent route",BOLD)
cell(z,f"E{r+1}",f"=AVERAGE(E{T+1}:E{T+5})",BOLD,you_fill,money,bd=True)
cell(z,f"F{r+1}",f"=AVERAGE(F{T+1}:F{T+5})",None,None,money,bd=True)
cell(z,f"A{r+2}","AVG — FSBO route (your likely path)",BOLD)
cell(z,f"E{r+2}",f"=AVERAGE(E{T+6}:E{T+10})",BOLD,tag_fill,money,bd=True)
cell(z,f"F{r+2}",f"=AVERAGE(F{T+6}:F{T+10})",None,None,money,bd=True)

vn=r+4
vnotes=[
 ("THE VERDICT — is $35k fair?",BOLD),
 ("• The crossover: her $35k-total is CHEAPER for you than your own $14k plan whenever the house nets more",None),
 ("  than $21,000. On the FSBO route (your likely path) it nets $21k–$41k — so $35k-total beats your $14k",None),
 ("  plan in almost every FSBO case, and you even KEEP proceeds above $35k at the high end.",None),
 ("• On FSBO your average cost under $35k-total is ~$4k — which is about what NC law would make you pay",None),
 ("  anyway (~$4k for the 401k). So on your likely route it's genuinely FAIR to both: she gets a guaranteed",None),
 ("  $35k, you pay roughly your legal share, and her recovery is capped BELOW her $48k.",None),
 ("• The catch: $35k is GUARANTEED, so YOU carry the downside. If the house sells badly (agent + low price,",None),
 ("  big concessions, or a long delay), your top-up climbs — up to ~$25k at a $390k agent sale.",None),
 ("• It is NOT 'generous of her' in the strict sense — she'd legally bear the whole market loss on her own",None),
 ("  $48k — but she HAS moved a long way down (from ~$25k+ to a capped $35k), so it's a reasonable landing.",None),
 ("",None),
 ("HOW TO MAKE IT FAIR TO YOU (if you accept $35k)",BOLD),
 ("• Since you're guaranteeing the number, YOU control the sale: you pick FSBO vs agent, set the list price,",None),
 ("  and approve any offer. That lets you steer net proceeds up and shrink your top-up.",None),
 ("• Keep it a FIXED $35k total (cleaner than percentages) and put in writing that proceeds above $35k are",None),
 ("  yours. Francie herself said 'it's a signed agreement' that matters — so get it signed and notarized.",None),
 ("• Hold firm on the MUTUAL non-disparagement clause (the Alessia/your-daughter issue). She resisted it;",None),
 ("  it belongs in the signed document, binding both of you.",None),
 ("• Optional counter: given her ~$4k legal floor, a fixed total of ~$28k–$30k is defensible — but $35k is",None),
 ("  not unfair to you on the FSBO math, and may be worth it to close cleanly.",None),
]
for i,(txt,font) in enumerate(vnotes):
    cell(z,f"A{vn+i}",txt,font)
z.column_dimensions["A"].width=40
for c in "BCDEFGH": z.column_dimensions[c].width=15
cell(z,"H12","",None)

wb.save("/home/user/Claude-Works/House_Sale_Scenarios_537_Duchart.xlsx")
print("saved")
