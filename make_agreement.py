from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc=Document()
st=doc.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(11)

def H(text,size=13,after=6,before=10):
    p=doc.add_paragraph(); r=p.add_run(text); r.bold=True; r.font.size=Pt(size)
    p.paragraph_format.space_after=Pt(after); p.paragraph_format.space_before=Pt(before); return p
def P(text,after=6,italic=False,bold=False):
    p=doc.add_paragraph(); r=p.add_run(text); r.italic=italic; r.bold=bold
    p.paragraph_format.space_after=Pt(after); return p
def NUM(text,after=4):
    p=doc.add_paragraph(text,style='List Number'); p.paragraph_format.space_after=Pt(after); return p

# Title
t=doc.add_paragraph(); t.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=t.add_run('SEPARATION AND PROPERTY SETTLEMENT AGREEMENT'); r.bold=True; r.font.size=Pt(15)

d=doc.add_paragraph(); d.alignment=WD_ALIGN_PARAGRAPH.CENTER
rd=d.add_run('DRAFT for the parties’ review — not legal advice. To be effective in North Carolina this must be signed by both parties and acknowledged before a notary.')
rd.italic=True; rd.font.size=Pt(9); rd.font.color.rgb=RGBColor(0x80,0x80,0x80)

P('This Separation and Property Settlement Agreement (the “Agreement”) is entered into by and between '
  'OSCAR A. PENNY (“Husband”) and FRANCIE R. PENNY (“Wife”), collectively the “parties.”',after=8)

H('Recitals',12)
P('A. The parties were lawfully married on August [__], 2023.')
P('B. The parties separated on May [__], 2026, and have lived continuously separate and apart since that date '
  'with the intent to end the marriage.')
P('C. No children were born of or adopted during this marriage, and Wife is not now pregnant.')
P('D. For reference only, Wife entered the marriage with approximately $90,000 received from her prior divorce, '
  'which was applied approximately as follows: $30,000 to the parties’ savings, $48,000 to the down payment on '
  'the residence, and $12,000 to Wife’s retirement account.')
P('E. The parties desire to settle, fully and finally, all rights and obligations arising from their marriage, '
  'including all rights to property (equitable distribution), spousal support, and related claims.',after=8)
P('NOW, THEREFORE, in consideration of the mutual promises below and other good and valuable consideration, '
  'the parties agree as follows:',after=8)

H('1. Separation',12)
NUM('The parties shall continue to live separate and apart. Each may reside where and with whom they choose and '
    'conduct their affairs as if unmarried. Neither shall molest, harass, or interfere with the other.')

H('2. Separate Property Retained by Each Party',12)
NUM('Wife shall retain as her sole and separate property the approximately $30,000 in savings previously '
    'delivered to her by Husband, together with any account holding her separate funds. Wife’s motor vehicle and '
    'retirement account are addressed in Sections 6 and 4 below.')
NUM('Husband shall retain as his sole and separate property all property owned by him before the marriage or '
    'otherwise titled in his name, except as specifically provided in this Agreement.')
NUM('Each party waives any claim to the separate property of the other.')

H('3. Marital Residence — 537 Duchart Lane, Fuquay-Varina, NC 27526',12)
NUM('The residence is titled in Husband’s name and is encumbered by a mortgage (loan ending 5990) in '
    'Husband’s name, with an approximate payoff of $356,149.76.')
NUM('The parties shall list and sell the residence on terms they mutually approve. Until closing, Husband shall '
    'be solely responsible for the mortgage payment, property taxes, and insurance. The parties acknowledge '
    'Husband has paid, and shall continue to pay, these amounts from his own income since the date of separation.')
NUM('“Net Proceeds” means the gross sale price less: the mortgage payoff; real estate commissions; '
    'excise/transfer taxes; attorney, title, and closing costs; and customary seller costs and prorations at closing.')
NUM('Distribution of Net Proceeds and Husband’s Contribution:')
sub=[
 '(a) Wife shall receive from the Net Proceeds an amount equal to the lesser of (i) $48,000 or (ii) the Net '
 'Proceeds, but not less than $0.',
 '(b) Any Net Proceeds remaining after payment to Wife under (a) — that is, any Net Proceeds above $48,000 '
 '— shall be paid to Husband.',
 '(c) If the amount Wife receives under (a) is less than $48,000, Husband shall additionally pay Wife, from his '
 'own funds, an amount equal to the lesser of (i) $14,000 or (ii) the difference between $48,000 and the amount '
 'Wife received under (a). In no event shall Wife’s total recovery under this Section 3 exceed $48,000.',
 '(d) If the Net Proceeds are insufficient to pay the mortgage and costs of sale, Husband shall be solely '
 'responsible for any deficiency, and Wife shall owe nothing.',
]
for s in sub:
    p=doc.add_paragraph(s); p.paragraph_format.left_indent=Inches(0.5); p.paragraph_format.space_after=Pt(4)
NUM('The parties acknowledge that the $48,000 down payment originated from Wife’s separate property. This '
    'Section fully and finally resolves any claim by Wife to the residence, its equity, or the return of her down '
    'payment. Wife waives any further interest in the residence and shall execute any deed or documents needed to '
    'complete the sale or convey her interest, if any.')

H('4. Retirement Accounts',12)
NUM('Each party shall retain, as their sole and separate property, all retirement, 401(k), pension, and similar '
    'accounts held in their own name, free of any claim by the other. Each party waives any claim to the other’s '
    'retirement accounts.')

H('5. Bank Accounts',12)
NUM('Any joint account (account ending [____]) shall be closed, and its balance as of the date of separation '
    '(approximately $[____]) shall be divided equally. Each party shall retain all accounts held in their own name.')

H('6. Motor Vehicles and Personal Property',12)
NUM('Each party shall keep, free of any claim by the other, the motor vehicle they currently possess and title, '
    'and all household goods, furnishings, and personal effects in their respective possession.')

H('7. Debts',12)
NUM('Each party shall be solely responsible for, and shall indemnify and hold the other harmless from, any debt '
    'in their own name. Husband shall be solely responsible for the mortgage on the residence and any deficiency.')

H('8. Spousal Support / Alimony — Mutual Waiver',12)
NUM('Each party is self-supporting. Husband and Wife each fully, finally, and forever waive and release any claim '
    'to alimony, post-separation support, or spousal support of any kind from the other, whether past, present, or '
    'future. This waiver is a material part of the consideration for this Agreement and is not modifiable.')

H('9. Non-Disparagement — Mutual',12)
NUM('Neither party shall make disparaging or derogatory statements about the other to family members (including '
    'the parties’ respective children), friends, employers, or any third party. Neither party shall involve any '
    'child or family member in disputes between the parties.')

H('10. Full and Final Settlement; Release',12)
NUM('This Agreement settles all property rights between the parties, including all rights to equitable '
    'distribution under N.C. Gen. Stat. § 50-20. Except for the obligations created by this Agreement, each '
    'party releases the other from all claims arising out of the marriage, and each waives any right to share in '
    'the other’s property or estate.')

H('11. Absolute Divorce',12)
NUM('After the parties have lived separate and apart for one year and one day, either party may seek an absolute '
    'divorce, and the other shall reasonably cooperate. This Agreement shall survive the divorce as an independent '
    'contract and shall not be merged into any divorce judgment.')

H('12. General Provisions',12)
NUM('Governing Law. This Agreement is governed by the laws of the State of North Carolina.')
NUM('Voluntary Execution; Disclosure. Each party enters this Agreement freely and voluntarily, after full and fair '
    'disclosure of the other’s assets and debts, and has had the opportunity to consult independent legal counsel.')
NUM('Entire Agreement; Amendment. This is the entire agreement of the parties and supersedes all prior '
    'discussions. It may be amended only by a writing signed by both parties and acknowledged before a notary.')
NUM('Severability. If any provision is held invalid, the remaining provisions shall remain in effect.')
NUM('Binding Effect. This Agreement binds the parties and their heirs, executors, and assigns, and may be signed '
    'in counterparts.')

# Signatures
doc.add_paragraph()
P('IN WITNESS WHEREOF, the parties have signed this Agreement as of the dates written below.',after=14)

for name in ['OSCAR A. PENNY (Husband)','FRANCIE R. PENNY (Wife)']:
    doc.add_paragraph('______________________________________     Date: __________')
    P(name,after=16,bold=True)

# Notary blocks
H('Acknowledgment (North Carolina)',11,before=8)
for who in ['Oscar A. Penny','Francie R. Penny']:
    P('State of North Carolina, County of ____________________.',after=2)
    P('I certify that ' + who + ' personally appeared before me this day and acknowledged the due execution of the '
      'foregoing Agreement. Witness my hand and official seal this ____ day of ____________, 20____.',after=2)
    P('____________________________________   Notary Public.   My commission expires: ____________', after=12)

# footer disclaimer
d2=doc.add_paragraph(); rd2=d2.add_run(
 'This document is a draft template prepared to reflect the parties’ stated terms. It is not legal advice. '
 'The parties are each encouraged to have it reviewed before signing. It must be signed by both parties and '
 'notarized to be enforceable in North Carolina.')
rd2.italic=True; rd2.font.size=Pt(8.5); rd2.font.color.rgb=RGBColor(0x80,0x80,0x80)

doc.save('/home/user/Claude-Works/Penny_Separation_Agreement_DRAFT.docx')
print('saved docx')
